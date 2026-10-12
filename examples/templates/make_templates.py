"""Rebuild the example templates and the sample Word file.

    python templates/make_templates.py

It needs pandoc on the PATH (for its reference documents), python-docx, python-pptx and Pillow.
The files it writes are already in the folder, so you only run this to change them.
"""

from __future__ import annotations

import io
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt, RGBColor
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxColour
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

from py_doc.pandoc.reference import add_styles

HERE = Path(__file__).resolve().parent
BLUE = RGBColor(0x0B, 0x4F, 0x8A)


VERSION_BOX = (
    f'<p:sp {nsdecls("p", "a")}><p:nvSpPr><p:cNvPr id="90" name="Version"/><p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr>'
    '<p:spPr><a:xfrm><a:off x="457200" y="6400800"/><a:ext cx="4572000" cy="304800"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
    '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-GB" sz="1000"/><a:t>Version {{VERSION}}</a:t></a:r></a:p></p:txBody></p:sp>'
)


def pandoc_file(name: str) -> bytes:
    return subprocess.run(["pandoc", "--print-default-data-file", name], capture_output=True, check=True).stdout


def find(styles, name: str):
    """A style by its name as written in the file (python-docx looks up Heading names in lower case)."""
    return next(style for style in styles if style.name == name)


def company_docx() -> None:
    """A Word template: your fonts and colours, a title page with placeholders, and the marker for the body."""
    document = Document(io.BytesIO(pandoc_file("reference.docx")))
    add_styles(document)  # the boxes for callouts
    styles = document.styles
    find(styles, "Normal").font.name = "Calibri"
    find(styles, "Body Text").font.name = "Calibri"
    find(styles, "Body Text").font.size = Pt(11)
    for level, size in ((1, 20), (2, 15), (3, 13)):
        style = find(styles, f"Heading {level}")
        style.font.name, style.font.size, style.font.bold, style.font.color.rgb = "Calibri", Pt(size), True, BLUE
    for name in ("Caption", "Image Caption", "Table Caption"):
        if name in [s.name for s in styles]:
            find(styles, name).font.bold = False
            find(styles, name).font.italic = True
    if "Source Code" not in [s.name for s in styles]:
        styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH).base_style = find(styles, "Normal")
    find(styles, "Source Code").font.name = "Consolas"
    find(styles, "Source Code").font.size = Pt(9)

    for paragraph in list(document.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)  # pandoc's sample text
    title = document.add_paragraph("Replace-with-Title", style=find(styles, "Title"))
    title.alignment = 1
    for label, placeholder in (
        ("Document number", "Replace-with-Document-Number"),
        ("Version", "Replace-with-Version"),
        ("Author", "Replace-with-Author"),
        ("Date", "Replace-with-Date"),
        ("Classification", "Replace-with-Classification"),
    ):
        document.add_paragraph(f"{label}: {placeholder}", style="Normal")
    document.add_paragraph("Replace-with-Description", style="Normal")
    document.add_page_break()
    document.add_paragraph("[[INSERTCONTENT]]", style="Normal")  # the document body goes here
    section = document.sections[0]
    section.header.paragraphs[0].text = "Example Ltd"
    section.footer.paragraphs[0].text = "Replace-with-Title, version Replace-with-Version"
    document.save(HERE / "company.docx")


def company_pptx() -> None:
    """A PowerPoint template: pandoc's layouts with a coloured title and a footer shape."""
    presentation = Presentation(io.BytesIO(pandoc_file("reference.pptx")))
    for master in presentation.slide_masters:
        for shape in master.shapes:
            if shape.has_text_frame and shape.is_placeholder and shape.placeholder_format.type in (1, 3):  # title
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = PptxColour(0x0B, 0x4F, 0x8A)
        master.shapes._spTree.append(parse_xml(VERSION_BOX))  # a master has no add_textbox, so the shape is written out
    ids = presentation.slides._sldIdLst  # pandoc's reference deck has sample slides: a template needs none
    for slide_id in list(ids):
        presentation.part.drop_rel(slide_id.rId)
        ids.remove(slide_id)
    presentation.save(HERE / "company.pptx")


def logo() -> None:
    image = Image.new("RGB", (240, 80), (11, 79, 138))
    ImageDraw.Draw(image).text((20, 30), "Example Ltd", fill="white")
    image.save(HERE.parent / "docs" / "images" / "logo.png")


def meeting_notes() -> None:
    """A plain Word file to try ``py-doc import`` on."""
    document = Document()
    document.core_properties.title = "Design meeting"
    document.core_properties.author = "A. Writer"
    document.add_heading("Design meeting", level=1)
    document.add_paragraph("We agreed the buffer design. The decision is ").add_run("final").bold = True
    document.add_heading("Actions", level=2)
    for text in ("Confirm the buffer size", "Write the test plan"):
        document.add_paragraph(text, style="List Bullet")
    document.add_paragraph("Ask the vendor", style="List Number")
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    for row, (name, owner) in enumerate((("Action", "Owner"), ("Buffer size", "Firmware"), ("Test plan", "Test"))):
        table.cell(row, 0).text, table.cell(row, 1).text = name, owner
    data = io.BytesIO()
    Image.new("RGB", (120, 60), (46, 125, 50)).save(data, format="PNG")
    data.seek(0)
    document.add_picture(data)
    document.add_paragraph("Figure 1: The agreed layout", style="Caption")
    (HERE.parent / "word").mkdir(exist_ok=True)
    document.save(HERE.parent / "word" / "design_meeting.docx")


if __name__ == "__main__":
    company_docx()
    company_pptx()
    logo()
    meeting_notes()
    print("Wrote company.docx, company.pptx, docs/images/logo.png and word/design_meeting.docx")
