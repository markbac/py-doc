"""Regression tests for known defects, written as the behaviour we want.

Each test is marked ``xfail(strict=True)`` and names its tracking issue. When
the issue is fixed the test starts passing, strict mode turns that into a
failure, and the fix PR must remove the marker. This keeps every known defect
visible and stops fixes landing without a test.

Fixed defects move to the regular test modules (for example #6 directory
mode in test_cli.py, #49 and #53 in test_converter.py, #9 in test_converter.py
and test_cli.py, #10 import-time logging in test_logging.py).
"""

from __future__ import annotations

import pytest
from docx import Document
from docx_helpers import styles, texts


def defect(issue: str):
    return pytest.mark.xfail(strict=True, reason=issue)


class TestMarkdownSemantics:
    """Constructs the line-based converter does not understand (#8)."""

    @defect("#8")
    def test_inline_strong_and_emphasis_become_run_formatting(self, convert_text):
        out = convert_text("A **strong** and *emphasised* word.\n")
        runs = Document(str(out)).paragraphs[0].runs
        assert any(r.bold and r.text == "strong" for r in runs)
        assert any(r.italic and r.text == "emphasised" for r in runs)

    @defect("#8")
    def test_inline_code_uses_monospace_run(self, convert_text):
        out = convert_text("Run `make test` now.\n")
        runs = Document(str(out)).paragraphs[0].runs
        assert any(r.font.name == "Consolas" and r.text == "make test" for r in runs)

    @defect("#8")
    def test_links_render_as_hyperlinks_not_raw_markdown(self, convert_text):
        out = convert_text("See [the site](https://example.com).\n")
        assert "](" not in " ".join(texts(out))

    @defect("#8")
    def test_tables_become_word_tables(self, convert_text):
        out = convert_text("| A | B |\n| - | - |\n| 1 | 2 |\n")
        table = Document(str(out)).tables[0]
        assert [c.text for c in table.rows[0].cells] == ["A", "B"]
        assert [c.text for c in table.rows[1].cells] == ["1", "2"]

    @defect("#8")
    def test_block_quotes_use_a_quote_style(self, convert_text):
        out = convert_text("> Quoted words\n")
        assert styles(out) == ["Quote"]
        assert texts(out) == ["Quoted words"]

    @defect("#8")
    def test_thematic_break_is_not_rendered_as_literal_dashes(self, convert_text):
        out = convert_text("Above\n\n---\n\nBelow\n")
        assert "---" not in texts(out)

    @defect("#8")
    def test_nested_lists_use_deeper_list_styles(self, convert_text):
        out = convert_text("- Parent\n  - Child\n")
        assert styles(out) == ["List Bullet", "List Bullet 2"]

    @defect("#8")
    def test_soft_wrapped_lines_join_into_one_paragraph(self, convert_text):
        out = convert_text("Second paragraph\nwrapped over two lines.\n")
        assert texts(out) == ["Second paragraph wrapped over two lines."]
