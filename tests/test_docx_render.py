"""How each Markdown construct is rendered in Word (#8)."""

from __future__ import annotations

import base64
import struct
import zlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.run import Run
from docx_helpers import hyperlink_targets, outline, styles, table_texts, texts

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


def runs_of(path, index=0):
    """Every run of a paragraph, including those inside hyperlinks."""
    paragraph = Document(str(path)).paragraphs[index]
    return [Run(element, paragraph) for element in paragraph._p.iter(qn("w:r"))]


def wide_png(width: int = 2000, height: int = 10) -> bytes:
    """A valid grey PNG that is much wider than a page at the default 72 dpi."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        body = kind + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

    rows = b"".join(b"\x00" + b"\x80" * width for _ in range(height))
    header = struct.pack(">IIBBBBB", width, height, 8, 0, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b"")


def numbering_of(path):
    """(numId, ilvl) of every numbered paragraph, in order."""
    result = []
    for p in Document(str(path)).paragraphs:
        num_pr = p._p.pPr.find(qn("w:numPr")) if p._p.pPr is not None else None
        if num_pr is not None:
            result.append((int(num_pr.find(qn("w:numId")).get(qn("w:val"))), int(num_pr.find(qn("w:ilvl")).get(qn("w:val")))))
    return result


def start_of(path, num_id):
    definitions = Document(str(path)).part.numbering_part.numbering_definitions._numbering
    num = next(n for n in definitions.num_lst if n.numId == num_id)
    return int(num.find(qn("w:lvlOverride")).find(qn("w:startOverride")).get(qn("w:val")))


class TestInlineFormatting:
    def test_strong_and_emphasis_become_run_formatting(self, convert_text):
        out = convert_text("A **strong** and *emphasised* word.\n")
        runs = runs_of(out)
        assert any(r.bold and r.text == "strong" for r in runs)
        assert any(r.italic and r.text == "emphasised" for r in runs)

    def test_plain_text_around_them_is_not_formatted(self, convert_text):
        runs = runs_of(convert_text("A **strong** word.\n"))
        assert [(r.text, bool(r.bold)) for r in runs] == [("A ", False), ("strong", True), (" word.", False)]

    def test_nested_formatting_combines(self, convert_text):
        runs = runs_of(convert_text("***both***\n"))
        assert any(r.bold and r.italic and r.text == "both" for r in runs)

    def test_inline_code_uses_monospace_run(self, convert_text):
        runs = runs_of(convert_text("Run `make test` now.\n"))
        assert any(r.font.name == "Consolas" and r.text == "make test" for r in runs)
        assert not any(r.font.name == "Consolas" and "Run" in r.text for r in runs)

    def test_formatting_applies_inside_headings(self, convert_text):
        out = convert_text("# A **bold** title\n")
        assert texts(out) == ["A bold title"]
        assert any(r.bold and r.text == "bold" for r in runs_of(out))

    def test_soft_wrapped_lines_join_into_one_paragraph(self, convert_text):
        out = convert_text("Second paragraph\nwrapped over two lines.\n")
        assert texts(out) == ["Second paragraph wrapped over two lines."]

    def test_hard_break_becomes_a_line_break(self, convert_text):
        out = convert_text("one  \ntwo\n")
        assert texts(out) == ["one\ntwo"]

    def test_escaped_characters_are_literal(self, convert_text):
        assert texts(convert_text("a \\*literal\\* star\n")) == ["a *literal* star"]


class TestLinks:
    def test_links_render_as_hyperlinks_not_raw_markdown(self, convert_text):
        out = convert_text("See [the site](https://example.com).\n")
        assert texts(out) == ["See the site."]
        assert hyperlink_targets(out) == ["https://example.com"]

    def test_only_the_link_text_is_inside_the_hyperlink(self, convert_text):
        out = convert_text("See [the site](https://example.com) now.\n")
        (link,) = Document(str(out)).element.body.iter(qn("w:hyperlink"))
        assert "".join(t.text for t in link.iter(qn("w:t"))) == "the site"

    def test_formatting_inside_link_text_is_kept(self, convert_text):
        out = convert_text("[**bold** link](https://example.com)\n")
        assert any(r.bold and r.text == "bold" for r in runs_of(out))
        assert hyperlink_targets(out) == ["https://example.com"]

    def test_autolinks_and_reference_links(self, convert_text):
        out = convert_text("<https://example.com/a> and [ref][r]\n\n[r]: https://example.com/b\n")
        assert hyperlink_targets(out) == ["https://example.com/a", "https://example.com/b"]

    def test_link_without_text_shows_its_address(self, convert_text):
        out = convert_text("[](https://example.com)\n")
        assert texts(out) == ["https://example.com"]

    def test_link_title_becomes_a_tooltip(self, convert_text):
        out = convert_text('[x](https://example.com "Hello")\n')
        (link,) = Document(str(out)).element.body.iter(qn("w:hyperlink"))
        assert link.get(qn("w:tooltip")) == "Hello"


class TestImages:
    def test_local_picture_is_embedded_with_alt_text(self, tmp_path, convert_text):
        (tmp_path / "pic.png").write_bytes(PNG)
        out = convert_text("![A tiny picture](pic.png)\n")
        doc = Document(str(out))
        assert len(doc.inline_shapes) == 1
        assert doc.inline_shapes[0]._inline.docPr.get("descr") == "A tiny picture"
        assert texts(out) == [""]

    def test_picture_wider_than_the_page_is_scaled_down_keeping_its_shape(self, tmp_path, convert_text):
        (tmp_path / "wide.png").write_bytes(wide_png(2000, 10))
        doc = Document(str(convert_text("![x](wide.png)\n")))
        section = doc.sections[0]
        shape = doc.inline_shapes[0]
        assert shape.width == section.page_width - section.left_margin - section.right_margin
        assert abs(shape.width / shape.height - 200) < 1

    def test_missing_picture_keeps_its_alt_text_and_warns(self, convert_text, caplog):
        out = convert_text("![Architecture overview](images/missing.png)\n")
        assert texts(out) == ["[Image: Architecture overview]"]
        assert hyperlink_targets(out) == ["images/missing.png"]
        assert "images/missing.png" in caplog.text

    def test_remote_picture_is_not_downloaded(self, convert_text, caplog):
        out = convert_text("![Badge](https://example.com/badge.svg)\n")
        assert texts(out) == ["[Image: Badge]"]
        assert len(Document(str(out)).inline_shapes) == 0
        assert "not a local file" in caplog.text

    def test_unsupported_local_file_falls_back_to_alt_text(self, tmp_path, convert_text, caplog):
        (tmp_path / "logo.svg").write_text("<svg/>", encoding="utf-8")
        out = convert_text("![Logo](logo.svg)\n")
        assert texts(out) == ["[Image: Logo]"]
        assert "not a supported picture" in caplog.text


class TestTables:
    def test_tables_become_word_tables(self, convert_text):
        out = convert_text("| A | B |\n| - | - |\n| 1 | 2 |\n")
        assert table_texts(out) == [["A", "B"], ["1", "2"]]

    def test_header_row_is_bold_and_body_is_not(self, convert_text):
        table = Document(str(convert_text("| A | B |\n| - | - |\n| 1 | 2 |\n"))).tables[0]
        assert all(r.bold for c in table.rows[0].cells for r in c.paragraphs[0].runs)
        assert not any(r.bold for c in table.rows[1].cells for r in c.paragraphs[0].runs)

    def test_column_alignment_is_kept(self, convert_text):
        table = Document(str(convert_text("| L | C | R |\n| :- | :-: | -: |\n| 1 | 2 | 3 |\n"))).tables[0]
        alignments = [c.paragraphs[0].alignment for c in table.rows[1].cells]
        assert alignments == [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT]

    def test_cells_keep_inline_formatting_and_links(self, convert_text):
        out = convert_text("| A |\n| - |\n| `code` and [x](https://example.com) |\n")
        assert hyperlink_targets(out) == ["https://example.com"]
        cell_runs = Document(str(out)).tables[0].rows[1].cells[0].paragraphs[0].runs
        assert any(r.font.name == "Consolas" for r in cell_runs)

    def test_text_before_and_after_a_table_is_kept(self, convert_text):
        out = convert_text("Before\n\n| A |\n| - |\n| 1 |\n\nAfter\n")
        assert texts(out) == ["Before", "After"]


class TestQuotesAndBreaks:
    def test_block_quotes_use_a_quote_style(self, convert_text):
        out = convert_text("> Quoted words\n")
        assert styles(out) == ["Quote"]
        assert texts(out) == ["Quoted words"]

    def test_quote_lines_join_and_paragraphs_stay_separate(self, convert_text):
        out = convert_text("> one\n> two\n>\n> three\n")
        assert texts(out) == ["one two", "three"]
        assert styles(out) == ["Quote", "Quote"]

    def test_nested_quotes_are_indented_further(self, convert_text):
        out = convert_text("> outer\n>\n> > inner\n")
        assert styles(out) == ["Quote", "Quote"]
        assert outline(out)[0]["left_indent_in"] is None
        assert outline(out)[1]["left_indent_in"] == 0.8

    def test_quote_can_hold_a_list_and_code(self, convert_text):
        out = convert_text("> Intro\n>\n> - item\n>\n> ```\n> code\n> ```\n")
        assert texts(out) == ["Intro", "item", "code"]
        assert styles(out) == ["Quote", "List Bullet", "Code"]

    def test_thematic_break_is_not_rendered_as_literal_dashes(self, convert_text):
        out = convert_text("Above\n\n---\n\nBelow\n")
        assert "---" not in texts(out)
        assert texts(out) == ["Above", "", "Below"]

    def test_thematic_break_is_a_bordered_paragraph(self, convert_text):
        paragraph = Document(str(convert_text("***\n"))).paragraphs[0]
        assert paragraph._p.pPr.find(qn("w:pBdr")).find(qn("w:bottom")) is not None


class TestLists:
    def test_nested_lists_use_deeper_list_styles(self, convert_text):
        out = convert_text("- Parent\n  - Child\n    - Grandchild\n")
        assert styles(out) == ["List Bullet", "List Bullet 2", "List Bullet 3"]
        assert texts(out) == ["Parent", "Child", "Grandchild"]

    def test_levels_beyond_three_stay_at_level_three(self, convert_text):
        out = convert_text("- a\n  - b\n    - c\n      - d\n")
        assert styles(out) == ["List Bullet", "List Bullet 2", "List Bullet 3", "List Bullet 3"]

    def test_nested_ordered_lists(self, convert_text):
        out = convert_text("1. Step\n   1. Sub\n2. Next\n")
        assert styles(out) == ["List Number", "List Number 2", "List Number"]

    def test_bullets_can_nest_under_numbers(self, convert_text):
        out = convert_text("1. Step\n   - note\n")
        assert styles(out) == ["List Number", "List Bullet 2"]

    def test_separate_ordered_lists_each_restart_at_one(self, convert_text):
        out = convert_text("1. a\n2. b\n\nBreak\n\n1. c\n2. d\n")
        first, second = [num_id for num_id, _ in numbering_of(out)][::2]
        assert first != second
        assert start_of(out, first) == start_of(out, second) == 1

    def test_items_of_one_list_share_a_numbering_instance(self, convert_text):
        out = convert_text("1. a\n2. b\n3. c\n")
        assert len({num_id for num_id, _ in numbering_of(out)}) == 1

    def test_list_start_number_is_kept(self, convert_text):
        out = convert_text("5. five\n6. six\n")
        ((num_id, _), _) = numbering_of(out)
        assert start_of(out, num_id) == 5

    def test_second_paragraph_of_an_item_is_a_continuation(self, convert_text):
        out = convert_text("- first\n\n  more text\n")
        assert styles(out) == ["List Bullet", "List Continue"]
        assert texts(out) == ["first", "more text"]

    def test_code_inside_an_item_follows_it(self, convert_text):
        out = convert_text("- step\n\n  ```\n  run\n  ```\n")
        assert styles(out) == ["List Bullet", "Code"]
        assert texts(out) == ["step", "run"]

    def test_empty_items_keep_their_bullet(self, convert_text):
        out = convert_text("- \n- text\n")
        assert styles(out) == ["List Bullet", "List Bullet"]
        assert texts(out) == ["", "text"]

    def test_list_text_keeps_inline_formatting(self, convert_text):
        out = convert_text("- a **bold** item\n")
        assert any(r.bold and r.text == "bold" for r in runs_of(out))


class TestOtherBlocks:
    def test_raw_html_is_shown_as_written(self, convert_text):
        out = convert_text('<div class="note">\nRaw html\n</div>\n')
        assert texts(out) == ['<div class="note">\nRaw html\n</div>']

    def test_inline_html_is_shown_as_written(self, convert_text):
        assert texts(convert_text("Inline <span>html</span> text\n")) == ["Inline <span>html</span> text"]

    def test_front_matter_is_not_rendered(self, convert_text):
        out = convert_text("---\ntitle: Notes\n---\n\n# Heading\n")
        assert texts(out) == ["Heading"]

    def test_front_matter_sets_the_document_properties(self, convert_text):
        out = convert_text(
            "---\ntitle: Design Notes\nsubtitle: Draft\nauthor: Mark\nversion: 1.2\nstatus: Review\n"
            "id: DN-7\ntags: [a, b]\n---\n\n# Heading\n"
        )
        properties = Document(str(out)).core_properties
        assert (properties.title, properties.subject, properties.author) == ("Design Notes", "Draft", "Mark")
        assert (properties.version, properties.content_status, properties.identifier) == ("1.2", "Review", "DN-7")
        assert properties.keywords == "a, b"

    def test_a_template_property_is_kept_when_the_front_matter_does_not_set_it(self, tmp_path, convert_text):
        template = tmp_path / "t.docx"
        doc = Document()
        doc.core_properties.title = "Template title"
        doc.core_properties.author = "Template author"
        doc.save(str(template))
        out = convert_text("---\nauthor: Mark\n---\n\nText\n", template=template)
        properties = Document(str(out)).core_properties
        assert (properties.title, properties.author) == ("Template title", "Mark")

    def test_invalid_front_matter_is_a_conversion_error(self, tmp_path, convert_text):
        import pytest
        from py_doc import ConversionError

        with pytest.raises(ConversionError, match=r"Invalid front matter at line 1"):
            convert_text("---\ntitle: [unclosed\n---\n\n# A\n")
        assert not (tmp_path / "doc.docx").exists()

    def test_unsupported_markdown_raises_a_conversion_error(self, tmp_path, monkeypatch):
        import pytest
        from py_doc import ConversionError, DocxConverter
        from py_doc.markdown import ParseError
        from py_doc.render import docx as renderer

        def fail(_text):
            raise ParseError("Unsupported Markdown block 'x' at line 3")

        monkeypatch.setattr(renderer, "parse", fail)
        src = tmp_path / "a.md"
        src.write_text("x\n", encoding="utf-8")
        with pytest.raises(ConversionError, match=r"a\.md.*line 3"):
            DocxConverter().convert_file(src, tmp_path / "a.docx")
        assert not (tmp_path / "a.docx").exists()
