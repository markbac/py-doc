"""py-doc template check and show: one failing fixture per rule."""

from __future__ import annotations

import logging
import shutil
import subprocess

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from pptx import Presentation

from py_doc.app import main
from py_doc.template_check import check_template, inventory


def checks(report, level):
    return {f.check for f in report.findings if f.level == level}


@pytest.fixture
def good_docx(tmp_path):
    """A template with everything the pipeline wants."""
    document = Document()
    for kind in ("note", "tip", "important", "warning", "caution"):
        document.styles.add_style(f"Div.{kind}", WD_STYLE_TYPE.PARAGRAPH)
    for name in ("Caption", "Body Text", "First Paragraph", "Compact", "Block Text", "Source Code", "Image Caption", "Table Caption",
                 "Footnote Text", "Subtitle", "Author", "Date", "Abstract"):  # fmt: skip
        if name not in [s.name for s in document.styles]:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    for name in ("Verbatim Char", "Hyperlink", "Footnote Reference"):
        document.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
    document.styles["Caption"].font.bold = False
    document.add_paragraph("[[INSERTCONTENT]]")
    path = tmp_path / "good.docx"
    document.save(str(path))
    return path


def edit(path, change):
    document = Document(str(path))
    change(document)
    out = path.with_name("bad.docx")
    document.save(str(out))
    return out


def test_a_complete_template_passes(good_docx):
    report = check_template(good_docx)
    assert not report.failures
    assert checks(report, "WARN") <= {"Table style", "Headings"}


def test_missing_callout_styles_are_a_warning(good_docx):
    bad = edit(good_docx, lambda d: d.styles["Div.note"].element.getparent().remove(d.styles["Div.note"].element))
    assert "Callout styles" in checks(check_template(bad), "WARN")


def test_a_callout_style_of_the_wrong_type_fails(good_docx):
    def change(d):
        element = d.styles["Div.tip"].element
        element.getparent().remove(element)
        d.styles.add_style("Div.tip", WD_STYLE_TYPE.CHARACTER)

    assert "Callout styles" in checks(check_template(edit(good_docx, change)), "FAIL")


def test_inline_code_style_of_the_wrong_type_fails(good_docx):
    def change(d):
        element = d.styles["Verbatim Char"].element
        element.getparent().remove(element)
        d.styles.add_style("Verbatim Char", WD_STYLE_TYPE.PARAGRAPH)

    assert "Pandoc styles" in checks(check_template(edit(good_docx, change)), "FAIL")


def test_bold_body_text_fails(good_docx):
    assert "Bold leak" in checks(
        check_template(edit(good_docx, lambda d: setattr(d.styles["Normal"].font, "bold", True))), "FAIL"
    )


def test_a_normal_style_based_on_a_heading_fails(good_docx):
    def change(d):
        d.styles["Normal"].base_style = d.styles["Heading 1"]

    assert "Normal style" in checks(check_template(edit(good_docx, change)), "FAIL")


def test_a_missing_caption_and_marker_are_warnings(good_docx):
    def change(d):
        element = d.styles["Caption"].element
        element.getparent().remove(element)
        for p in d.paragraphs:
            if "INSERTCONTENT" in p.text:
                p._element.getparent().remove(p._element)

    assert {"Caption style", "Content marker"} <= checks(check_template(edit(good_docx, change)), "WARN")


def test_a_dotx_is_checked_like_a_docx(good_docx, tmp_path):
    from test_convert import make_dotx

    make_dotx(tmp_path / "t.dotx")
    assert check_template(tmp_path / "t.dotx").findings


def test_powerpoint_needs_pandoc_s_layouts(tmp_path):
    pptx = tmp_path / "plain.pptx"
    deck = Presentation()
    for layout in deck.slide_layouts:
        layout._element.cSld.set("name", "Custom " + layout.name)
    deck.save(str(pptx))
    assert "Layouts" in checks(check_template(pptx), "FAIL")


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc is not installed")
def test_pandoc_s_own_reference_deck_passes(tmp_path):
    pptx = tmp_path / "ref.pptx"
    pptx.write_bytes(
        subprocess.run(
            ["pandoc", "--print-default-data-file", "reference.pptx"], capture_output=True, check=True
        ).stdout
    )
    assert not check_template(pptx).failures


def test_a_deck_with_an_empty_placeholder_and_many_fonts_is_audited(tmp_path):
    pptx = tmp_path / "deck.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "T"  # the body placeholder stays empty
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    for font in ("A", "B", "C"):
        run = box.text_frame.add_paragraph().add_run()
        run.text = "x"
        run.font.name = font
    deck.save(str(pptx))
    report = check_template(pptx)
    assert "Empty placeholder" in checks(report, "WARN") and "Fonts" in checks(report, "WARN")


def test_inventory_lists_styles_and_layouts(good_docx, tmp_path):
    rows = inventory(good_docx)
    assert ("paragraph", "Div.note", "", "Div.note") in rows or any(r[1] == "Div.note" for r in rows)
    Presentation().save(str(tmp_path / "p.pptx"))
    assert any(r[0] == "Title Slide" for r in inventory(tmp_path / "p.pptx"))


class TestCommand:
    def test_exit_status_follows_failures(self, good_docx, capsys):
        assert main(["template", "check", str(good_docx)]) == 0
        bad = edit(good_docx, lambda d: setattr(d.styles["Normal"].font, "bold", True))
        assert main(["template", "check", str(bad)]) == 1
        assert "FAIL  Bold leak" in capsys.readouterr().out

    def test_strict_fails_on_warnings(self, good_docx):
        bad = edit(good_docx, lambda d: d.styles["Div.note"].element.getparent().remove(d.styles["Div.note"].element))
        assert main(["template", "check", str(bad)]) == 0
        assert main(["template", "check", str(bad), "--strict"]) == 1

    def test_show_prints_a_table(self, good_docx, capsys):
        assert main(["template", "show", str(good_docx)]) == 0
        assert "Div.note" in capsys.readouterr().out

    def test_not_a_template_is_status_2(self, tmp_path, capsys):
        text = tmp_path / "a.txt"
        text.write_text("x", encoding="utf-8")
        assert main(["template", "check", str(text)]) == 2
        assert main(["template", "check", str(tmp_path / "gone.docx")]) == 2
        broken = tmp_path / "broken.docx"
        broken.write_text("not a zip", encoding="utf-8")
        assert main(["template", "check", str(broken)]) == 2


def test_a_build_reports_template_problems_once_and_goes_on(tmp_path, pandoc_runtime, caplog):
    from py_doc.convert import DocxConverter

    template = tmp_path / "t.docx"
    document = Document()
    document.styles["Normal"].font.bold = True
    document.save(str(template))
    (tmp_path / "a.md").write_text("# A\n\ntext\n", encoding="utf-8")
    with caplog.at_level(logging.WARNING):
        for name in ("one", "two"):
            DocxConverter(template_path=template).convert_file(tmp_path / "a.md", tmp_path / f"{name}.docx")
    assert caplog.text.count("'Normal' is bold") == 1
    assert (tmp_path / "two.docx").is_file()


def test_diff_names_the_styles_that_differ(good_docx, tmp_path, capsys):
    other = tmp_path / "other.docx"
    document = Document(good_docx)
    document.styles.add_style("Extra", WD_STYLE_TYPE.PARAGRAPH)
    document.styles["Normal"].font.size = None
    document.styles["Caption"].font.bold = True
    document.save(other)
    assert main(["template", "diff", str(good_docx), str(other)]) == 1
    out = capsys.readouterr().out
    assert "only in second" in out and "Extra" in out and "differs" in out and "Caption" in out
    assert main(["template", "diff", str(good_docx), str(good_docx)]) == 0


def test_diff_refuses_a_word_file_against_a_deck(good_docx, tmp_path, capsys):
    good_pptx = tmp_path / "deck.pptx"
    Presentation().save(good_pptx)
    assert main(["template", "diff", str(good_docx), str(good_pptx)]) == 2


def test_usage_counts_paragraph_and_table_styles(tmp_path, capsys):
    document = Document()
    document.add_paragraph("a", style="Heading 1")
    document.add_paragraph("b")
    document.add_paragraph("c")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "cell"
    path = tmp_path / "used.docx"
    document.save(path)
    assert main(["template", "usage", str(path)]) == 0
    lines = capsys.readouterr().out.splitlines()
    assert lines[1].split() == ["Normal", "3"]
    assert any(line.split()[:1] == ["Heading"] and line.split()[-1] == "1" for line in lines)
