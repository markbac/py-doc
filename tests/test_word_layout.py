"""Word layout through pandoc: contents list, captions, callouts, landscape pages, code, bold."""

from __future__ import annotations

import logging
import re
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fake_engine import install
from test_pandoc import png

from py_doc.build import build_document
from py_doc.config import load_config
from py_doc.runtime import Runtime

pytestmark = pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc is not installed")

SAMPLE = """# Guide

## Section one

Text with **bold words** and `code`.

> [!NOTE]
> A note.

> [!TIP]
> A tip.

> [!IMPORTANT]
> Something important.

> **Warning:** Be careful.

> [!CAUTION]
> Mind the gap.

[[CAPTION:Figure]] First picture
![](pic.png)

[[CAPTION:Figure]] Second picture
![](pic.png)

Table: Sizes

| a | b |
|---|---|
| 1 | 2 |

```python
def f(x):
    return "text"
```

[[LANDSCAPE]]

## Wide section

![](pic.png)

[[PORTRAIT]]

## Back to portrait

Last paragraph.
"""


@pytest.fixture
def built(tmp_path, monkeypatch):
    engine = install(tmp_path, monkeypatch, mode="host")
    runtime = lambda: Runtime("tools:test", docker=engine.docker, platform_name="linux")  # noqa: E731

    def build(outputs="docx: {toc: true}", meta="toc_depth: 2\n", sample=SAMPLE):
        (tmp_path / "g.md").write_text(sample, encoding="utf-8")
        png(tmp_path / "pic.png")
        (tmp_path / "py-doc.yml").write_text(
            f"documents:\n  g:\n    title: Guide\n{meta and '    ' + meta}    sources: [g.md]\n    outputs:\n      {outputs}\n", encoding="utf-8"
        )
        config = load_config(tmp_path / "py-doc.yml")
        (out,) = build_document(config, config.documents["g"], runtime=runtime)
        return out

    return build


def xml(path: Path, member: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read(member).decode("utf-8")


def test_the_contents_list_follows_toc_depth_and_updates_on_opening(built):
    out = built()
    body = xml(out, "word/document.xml")
    assert re.search(r'TOC \\o "1-2"', body)
    assert "<w:updateFields" in body


def test_no_contents_list_unless_asked(built):
    assert "TOC \\o" not in xml(built("docx: {}"), "word/document.xml")


def test_figure_captions_are_numbered_by_fields_across_the_document(built):
    out = built()
    body = xml(out, "word/document.xml")
    assert len(re.findall(r"SEQ Figure", body)) == 2
    captions = [p for p in Document(str(out)).paragraphs if p.style.name == "Caption"]
    assert [c.text for c in captions[:2]] == ["Figure 1: First picture", "Figure 2: Second picture"]
    assert "[[CAPTION" not in body


def test_a_missing_explicit_style_is_one_warning(built, caplog):
    sample = "# T\n\n[[STYLE:Nope]] one\n\n[[STYLE:Nope]] two\n"
    with caplog.at_level(logging.WARNING):
        built(sample=sample)
    assert len([r for r in caplog.records if "Nope" in r.getMessage()]) == 1


def test_table_captions_are_numbered(built):
    assert "SEQ Table" in xml(built(), "word/document.xml")


def test_each_callout_kind_is_a_box_in_its_own_style(built):
    out = built()
    document = Document(str(out))
    styles = xml(out, "word/styles.xml")
    used = {p.style.name: p.text for p in document.paragraphs if p.style.name.startswith("Div.")}
    for kind, fill in (("note", "E7F1F8"), ("tip", "E8F5E9"), ("important", "F1ECF9"), ("warning", "FFF4E0"), ("caution", "FCE8E8")):
        assert f"Div.{kind}" in used or any(s.startswith(f"Div.{kind}") for s in used), kind
        assert fill in styles, f"{kind} box has no background"
    assert re.search(r'<w:left [^>]*w:val="single"', styles)


def test_a_template_without_the_callout_styles_is_told_once(built, tmp_path, caplog):
    Document().save(tmp_path / "plain.docx")
    with caplog.at_level(logging.WARNING):
        built("docx: {template: plain.docx}")
    warnings = [r.getMessage() for r in caplog.records if "has no style" in r.getMessage()]
    assert len(warnings) == 1 and "Div.note" in warnings[0] and "Div.caution" in warnings[0]


def test_a_template_is_only_told_about_the_callouts_used(built, tmp_path, caplog):
    Document().save(tmp_path / "plain.docx")
    with caplog.at_level(logging.WARNING):
        built("docx: {template: plain.docx}", sample="# T\n\n> [!NOTE]\n> n\n")
    (message,) = [r.getMessage() for r in caplog.records if "has no style" in r.getMessage()]
    assert "Div.note" in message and "Div.tip" not in message


def test_landscape_section_between_the_markers(built):
    out = built()
    body = xml(out, "word/document.xml")
    assert "[[LANDSCAPE]]" not in body and "[[PORTRAIT]]" not in body
    orientations = re.findall(r'<w:pgSz[^>]*w:orient="(\w+)"', body)
    assert "landscape" in orientations
    sizes = [(s.page_width, s.page_height) for s in Document(str(out)).sections if s.page_width and s.page_height]
    assert any(w > h for w, h in sizes) and sizes[-1][1] > sizes[-1][0]  # landscape in the middle, portrait again at the end


def test_code_uses_the_highlight_theme(built):
    assert "4f6bed" in xml(built(), "word/styles.xml").lower()


def test_bold_in_the_text_stays_bold(built):
    """normalize_body_bold removed direct bold from body paragraphs, which also removed **bold** text."""
    document = Document(str(built()))
    bold = [r.text for p in document.paragraphs for r in p.runs if r.bold]
    assert "bold words" in bold


def test_the_file_is_well_formed(built):
    from lxml import etree

    out = built()
    with zipfile.ZipFile(out) as archive:
        assert archive.testzip() is None
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                etree.fromstring(archive.read(name))


def test_a_picture_in_the_last_paragraph_is_kept(built):
    out = built(sample="# T\n\nText.\n\n![](pic.png)\n")
    assert len(Document(str(out)).inline_shapes) == 1


def test_a_caption_below_a_table_stays_on_its_page(built):
    sample = "# T\n\n| a | b |\n|---|---|\n| 1 | 2 |\n| 3 | 4 |\n\n[[CAPTION:Table]] Sizes\n\nAfter.\n"
    document = Document(str(built(sample=sample)))
    (table,) = document.tables
    caption = table._tbl.getnext()
    assert caption.tag.endswith("}p") and "Sizes" in "".join(caption.itertext())
    assert all(p.paragraph_format.keep_with_next for cell in table.rows[-1].cells for p in cell.paragraphs)
    assert not any(p.paragraph_format.keep_with_next for cell in table.rows[1].cells for p in cell.paragraphs)
