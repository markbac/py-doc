"""PowerPoint through pandoc: footer, version token, notes, font repair, long slides."""

from __future__ import annotations

import logging
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from fake_engine import install
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from py_doc.build import build_document
from py_doc.config import load_config
from py_doc.pandoc.pptx_finish import ensure_fntdata_content_type, notes_to_divs
from py_doc.runtime import Runtime

pytestmark = pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc is not installed")

SAMPLE = """---
title: Deck
version: 2.1
---

# Deck

## First

- one
- two

<!-- notes: Say hello. -->

## Second

Plain text.
"""


@pytest.fixture
def make(tmp_path, monkeypatch):
    engine = install(tmp_path, monkeypatch, mode="host")
    runtime = lambda: Runtime("tools:test", docker=engine.docker, platform_name="linux")  # noqa: E731

    def build(options="", sample=SAMPLE, template=None):
        (tmp_path / "d.md").write_text(sample, encoding="utf-8")
        tpl = f"template: {template}, " if template else ""
        (tmp_path / "py-doc.yml").write_text(
            f"documents:\n  deck:\n    sources: [d.md]\n    outputs:\n      pptx: {{{tpl}{options}}}\n",
            encoding="utf-8",
        )
        config = load_config(tmp_path / "py-doc.yml")
        (out,) = build_document(config, config.documents["deck"], ["pptx"], runtime=runtime)
        return out

    return build


def reference(tmp_path: Path) -> Path:
    path = tmp_path / "ref.pptx"
    path.write_bytes(
        subprocess.run(
            ["pandoc", "--print-default-data-file", "reference.pptx"], capture_output=True, check=True
        ).stdout
    )
    return path


def texts(out: Path) -> list[str]:
    return [s.text_frame.text for slide in Presentation(str(out)).slides for s in slide.shapes if s.has_text_frame]


def test_notes_comments_become_speaker_notes(make):
    out = make()
    notes = [s.notes_slide.notes_text_frame.text for s in Presentation(str(out)).slides if s.has_notes_slide]
    assert "Say hello." in notes
    assert not any("notes:" in t for t in texts(out))


def test_notes_conversion_leaves_code_alone():
    text = "```\n<!-- notes: keep -->\n```\n\n<!-- notes: multi\nline -->\n"
    out = notes_to_divs(text)
    assert "<!-- notes: keep -->" in out
    assert "::: notes" in out and "multi" in out


def test_the_version_token_in_a_template_shape_is_replaced(make, tmp_path):
    from pptx.util import Inches

    from py_doc.markdown import Metadata
    from py_doc.pandoc.pptx_finish import finish_pptx

    out = make()
    presentation = Presentation(str(out))
    box = presentation.slides[0].shapes.add_textbox(0, 0, Inches(2), Inches(1))
    box.text_frame.text = "Issue {{VERSION}}"
    presentation.save(str(out))
    finish_pptx(out, Metadata(version="2.1"), {})
    assert "Issue 2.1" in texts(out)


def test_the_version_token_on_the_master_and_layouts_is_replaced(make):
    from py_doc.markdown import Metadata
    from py_doc.pandoc.pptx_finish import finish_pptx

    out = make()
    presentation = Presentation(str(out))
    next(x for x in presentation.slide_masters[0].shapes if x.has_text_frame).text_frame.text = "Master {{VERSION}}"
    next(x for x in presentation.slide_layouts[0].shapes if x.has_text_frame).text_frame.text = "Layout {{VERSION}}"
    presentation.save(str(out))
    finish_pptx(out, Metadata(version="3.4"), {})
    with zipfile.ZipFile(out) as archive:
        xml = "".join(archive.read(n).decode("utf-8") for n in archive.namelist() if n.endswith(".xml"))
    assert "Master 3.4" in xml and "Layout 3.4" in xml and "{{VERSION}}" not in xml


def test_footer_is_set_from_title_and_version(make, tmp_path):
    out = make("footer: '{title} v{version}'", template=reference(tmp_path).name)
    footers = [
        s.text_frame.text
        for slide in Presentation(str(out)).slides
        for s in slide.placeholders
        if s.placeholder_format.type is not None and "FOOTER" in str(s.placeholder_format.type)
    ]
    assert footers and set(footers) == {"Deck v2.1"}


def test_a_layout_without_a_footer_is_reported(make, tmp_path, caplog):
    """pandoc's own layouts have footers; a layout without one is reported (checked on the finishing step)."""

    from py_doc.markdown import Metadata
    from py_doc.pandoc.pptx_finish import finish_pptx

    out = make()
    presentation = Presentation(str(out))
    for layout in presentation.slide_layouts:
        for shape in list(layout.placeholders):
            if "FOOTER" in str(shape.placeholder_format.type):
                shape._element.getparent().remove(shape._element)
    presentation.save(str(out))
    with caplog.at_level(logging.WARNING):
        finish_pptx(out, Metadata(title="T"), {"footer": "x"})
    assert "has no footer placeholder" in caplog.text


def test_a_crowded_slide_is_reported(make, caplog):
    sample = "# D\n\n## Crowded\n\n" + "\n".join(f"- item {i}" for i in range(20)) + "\n"
    with caplog.at_level(logging.WARNING):
        make(sample=sample)
    assert "Slide" in caplog.text and "too many" in caplog.text


def test_embedded_font_content_type_is_registered(tmp_path):
    path = tmp_path / "f.pptx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="x"></Types>')
        z.writestr("ppt/fonts/font1.fntdata", b"data")
    ensure_fntdata_content_type(path)
    with zipfile.ZipFile(path) as z:
        assert 'Extension="fntdata"' in z.read("[Content_Types].xml").decode()
        assert z.read("ppt/fonts/font1.fntdata") == b"data"
    ensure_fntdata_content_type(path)  # a second run changes nothing
    with zipfile.ZipFile(path) as z:
        assert z.read("[Content_Types].xml").decode().count("fntdata") == 1


def test_split_level_controls_the_slides(make):
    one = len(Presentation(str(make("split_level: 1"))).slides)
    two = len(Presentation(str(make("split_level: 2"))).slides)
    assert one != two


@pytest.mark.skipif(shutil.which("dot") is None, reason="Graphviz is not installed")
def test_a_diagram_on_a_slide_is_a_picture_with_its_marker_caption(make):
    out = make(sample="## Flow\n\n[[CAPTION:Figure]] From sensor to cloud\n```dot\ndigraph { a -> b }\n```\n")
    (slide,) = Presentation(str(out)).slides
    assert any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
    assert "From sensor to cloud" in texts(out) and "dot diagram" not in texts(out)
