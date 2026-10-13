"""Section numbers written in headings are left out where the output numbers its headings itself."""

from __future__ import annotations

import shutil
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fake_engine import install

from py_doc.build import _numbers_headings_itself, build_document
from py_doc.config import ConfigError, load_config
from py_doc.pandoc.engine import PandocJob, _command
from py_doc.pandoc.reference import numbers_headings
from py_doc.pandoc.source import strip_heading_numbers
from py_doc.markdown import Metadata
from py_doc.runtime import Runtime


@pytest.mark.parametrize(
    ("before", "after"),
    [
        ("## 2.1 Design\n", "## Design\n"),
        ("### 2.1.3. Buffer size\n", "### Buffer size\n"),
        ("# 3. Results\n", "# Results\n"),
        ("## A.1 Glossary\n", "## Glossary\n"),
        ("## 2024 roadmap\n", "## 2024 roadmap\n"),  # a bare number that is not a count of the sections
        ("## 2.1\n", "## 2.1\n"),  # nothing would be left
        ("# Appendix A\n", "# Appendix A\n"),
        ("```\n# 2.1 code\n```\n", "```\n# 2.1 code\n```\n"),
    ],
)
def test_section_numbers_come_off_and_titles_stay(before, after):
    assert strip_heading_numbers(before) == after


def test_bare_numbers_count_only_when_the_level_counts_one_two_three():
    counted = "# 1 Introduction\n\ntext\n\n# 2 Design\n\n# 3 Tests\n"
    assert strip_heading_numbers(counted) == "# Introduction\n\ntext\n\n# Design\n\n# Tests\n"
    not_counted = "# 1 Introduction\n\n# 2024 roadmap\n"
    assert strip_heading_numbers(not_counted) == not_counted


def _template(path: Path, numbered: bool) -> Path:
    document = Document()
    if numbered:
        style = document.styles["Heading 1"]
        style.element.get_or_add_pPr().append(parse_xml(f'<w:numPr {nsdecls("w")}><w:numId w:val="1"/></w:numPr>'))
    document.save(str(path))
    return path


def test_a_template_with_numbered_heading_styles_is_found(tmp_path):
    assert numbers_headings(_template(tmp_path / "numbered.docx", True))
    assert not numbers_headings(_template(tmp_path / "plain.docx", False))
    assert not numbers_headings(tmp_path / "missing.docx")


def test_the_choice_per_output(tmp_path):
    numbered, plain = _template(tmp_path / "n.docx", True), _template(tmp_path / "p.docx", False)
    assert _numbers_headings_itself("docx", numbered, {})
    assert not _numbers_headings_itself("docx", plain, {})
    assert not _numbers_headings_itself("docx", None, {})
    assert _numbers_headings_itself("docx", plain, {"heading_numbers": "strip"})
    assert not _numbers_headings_itself("docx", numbered, {"heading_numbers": "keep"})
    assert _numbers_headings_itself("pdf", None, {"number_sections": True})
    assert not _numbers_headings_itself("html", None, {})
    assert not _numbers_headings_itself("pptx", None, {})


def test_number_sections_asks_pandoc_to_number(tmp_path):
    class Box:
        def container_path(self, path):
            return Path(path).as_posix()

        def run(self, *args, **kwargs):
            raise AssertionError

    job = PandocJob("html", "# A\n", Metadata(), tmp_path / "a.html", options={"number_sections": True})
    argv, _ = _command(job, Box(), tmp_path, None)
    assert "--number-sections" in argv


def test_the_options_are_checked(tmp_path):
    (tmp_path / "py-doc.yml").write_text(
        "documents:\n  d:\n    sources: [a.md]\n    outputs:\n      docx: {heading_numbers: strip}\n      pdf: {number_sections: true}\n",
        encoding="utf-8",
    )
    config = load_config(tmp_path / "py-doc.yml")
    assert config.documents["d"].outputs["docx"].options["heading_numbers"] == "strip"
    (tmp_path / "py-doc.yml").write_text("documents:\n  d:\n    sources: [a.md]\n    outputs:\n      docx: {heading_numbers: maybe}\n", encoding="utf-8")
    with pytest.raises(ConfigError, match="heading_numbers must be one of auto, strip, keep"):
        load_config(tmp_path / "py-doc.yml")


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc is not installed")
def test_a_numbered_word_template_gets_headings_without_their_numbers(tmp_path, monkeypatch):
    engine = install(tmp_path, monkeypatch, mode="host")
    runtime = lambda: Runtime("tools:test", docker=engine.docker, platform_name="linux")  # noqa: E731
    _template(tmp_path / "numbered.docx", True)
    (tmp_path / "a.md").write_text("# 1. Introduction\n\nText.\n\n## 1.1 Scope\n\nMore.\n", encoding="utf-8")
    (tmp_path / "py-doc.yml").write_text(
        "documents:\n  a:\n    sources: [a.md]\n    outputs:\n      docx: {template: numbered.docx}\n", encoding="utf-8"
    )
    config = load_config(tmp_path / "py-doc.yml")
    (out,) = build_document(config, config.documents["a"], ["docx"], runtime=runtime)
    headings = [p.text for p in Document(str(out)).paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["Introduction", "Scope"]
