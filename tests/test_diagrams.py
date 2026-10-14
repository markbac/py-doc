"""Diagram fences: found outside code, drawn through the container runner, cached, loud when broken."""

from __future__ import annotations

import json
import logging
import os
import re
import shutil
import stat
import subprocess
import sys
from pathlib import Path

import pytest
from fake_engine import install
from test_pandoc import png

from py_doc import diagrams
from py_doc.build import build_document
from py_doc.config import load_config
from py_doc.diagrams import DiagramError, render_parts
from py_doc.diagrams.registry import BY_FENCE, RENDERERS, TOOLS, Job
from py_doc.diagrams.render import _finish_png
from py_doc.diagrams.scan import find_blocks
from py_doc.runtime import Runtime

needs_dot = pytest.mark.skipif(shutil.which("dot") is None, reason="graphviz is not installed")
needs_pandoc = pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc is not installed")

# Every fence createdocs understood.
CREATEDOCS_FENCES = {
    "mermaid", "plantuml", "puml", "dot", "ditaa", "packetdiag", "nwdiag", "blockdiag", "seqdiag", "actdiag",
    "wavedrom", "bytefield-svg", "pikchr", "railroad", "wireviz", "vega", "vega-lite",
}  # fmt: skip


@pytest.fixture(autouse=True)
def cache(tmp_path, monkeypatch):
    monkeypatch.setenv("PY_DOC_CACHE", str(tmp_path / "cache"))
    return tmp_path / "cache" / "diagrams"


@pytest.fixture
def host(tmp_path, monkeypatch):
    engine = install(tmp_path, monkeypatch, mode="host")
    return engine, (lambda: Runtime("tools:test", docker=engine.docker, platform_name="linux"))


def stub(tmp_path, monkeypatch, name: str, body: str) -> Path:
    """A program called ``name`` on the PATH of the 'container' (the host, in these tests)."""
    folder = tmp_path / "stubs"
    folder.mkdir(exist_ok=True)
    path = folder / name
    code = f"import json, os, sys\n{body}\n"
    if sys.platform == "win32":  # a script cannot start by its bare name there: name.bat runs name.py
        (folder / f"{name}.py").write_text(code, encoding="utf-8")
        path = folder / f"{name}.bat"
        path.write_text(f'@"{sys.executable}" "%~dp0{name}.py" %*\n', encoding="utf-8")
    else:
        path.write_text(f"#!{sys.executable}\n{code}", encoding="utf-8")
        path.chmod(path.stat().st_mode | stat.S_IEXEC)
    monkeypatch.setenv("PATH", f"{folder}{os.pathsep}{os.environ['PATH']}")
    return path


PNG_STUB = """
log = open(os.environ["STUB_LOG"], "a")
log.write(json.dumps({"argv": sys.argv[1:], "java": os.environ.get("JAVA_TOOL_OPTIONS")}) + "\\n")
"""


# --- finding the fences ------------------------------------------------------------------------------


class TestScan:
    def test_finds_diagram_fences_in_all_spellings(self):
        text = "a\n```plantuml title=\"Ctx\"\nA -> B\n```\n\n~~~puml\nX\n~~~\n\n:::mermaid\ngraph TD\n:::\n\n::: mermaid\nflowchart\n:::\n"
        blocks = find_blocks(text)
        assert [(b.fence, b.line) for b in blocks] == [("plantuml", 2), ("puml", 6), ("mermaid", 10), ("mermaid", 14)]
        assert blocks[0].body == "A -> B" and blocks[0].options == {"title": "Ctx"}

    def test_diagram_text_inside_other_code_is_left_alone(self):
        text = "````markdown\n```plantuml\nA\n```\n````\n\n```python\n:::mermaid\n:::\n```\n"
        assert find_blocks(text) == []

    def test_other_languages_and_plain_fences_are_ignored(self):
        assert find_blocks("```python\nx\n```\n\n```\ny\n```\n") == []

    def test_a_longer_fence_closes_only_with_a_fence_as_long(self):
        (block,) = find_blocks("````dot\n```\ndigraph{}\n```\n````\n")
        assert block.body == "```\ndigraph{}\n```"

    def test_every_createdocs_fence_has_a_renderer(self):
        assert CREATEDOCS_FENCES <= set(BY_FENCE)


# --- the registry -------------------------------------------------------------------------------------


class TestRegistry:
    def job(self, kind="png", landscape=False):
        return Job(Path("/w/diagram.x"), Path(f"/w/diagram.{kind}"), kind, landscape, Path("/src"))

    @pytest.mark.parametrize("renderer", RENDERERS, ids=lambda r: r.name)
    def test_every_renderer_builds_commands_for_both_kinds(self, renderer):
        for kind in ("png", "svg"):
            commands = renderer.commands(self.job(kind))
            assert commands and all(isinstance(word, (str, Path)) for command in commands for word in command)

    def test_a_tool_that_cannot_write_svg_gives_png(self):
        (command,) = BY_FENCE["ditaa"].commands(self.job("svg"))
        assert str(command[-1]).endswith("diagram.png")

    def test_the_tool_scripts_exist(self):
        for name in ("pikchr_render.py", "railroad_render.py", "mermaid.json"):
            assert (TOOLS / name).is_file()

    def test_svg_tools_are_followed_by_inkscape_for_png_only(self):
        wavedrom = BY_FENCE["wavedrom"]
        assert [c[0] for c in wavedrom.commands(self.job("png"))] == ["wavedrom-cli", "sh"]
        assert [c[0] for c in wavedrom.commands(self.job("svg"))] == ["wavedrom-cli"]

    def test_mermaid_is_wider_on_a_landscape_page(self):
        portrait = BY_FENCE["mermaid"].commands(self.job())[0]
        landscape = BY_FENCE["mermaid"].commands(self.job(landscape=True))[0]
        assert portrait[portrait.index("--width") + 1] == "1200" and landscape[landscape.index("--width") + 1] == "1800"
        assert "/etc/py-doc/puppeteer.json" in portrait

    def test_plantuml_only_caches_diagrams_without_includes(self):
        assert BY_FENCE["plantuml"].cacheable("A -> B") and not BY_FENCE["plantuml"].cacheable("!include x.puml")


# --- drawing -----------------------------------------------------------------------------------------


def part(tmp_path, text: str, name="doc.md"):
    return [(tmp_path / name, text)]


@needs_dot
class TestDot:
    DOT = "digraph { a -> b }"

    def test_png_with_density_and_alt_text(self, tmp_path, host):
        _, runtime = host
        rendered = render_parts(part(tmp_path, f"before\n\n```dot title=\"Flow\"\n{self.DOT}\n```\n\nafter\n"), "docx", runtime)
        (name, path), = rendered.pictures.items()
        assert name.startswith("diagrams/") and path.read_bytes().startswith(b"\x89PNG")
        assert '![Flow](' + name + ')' in rendered.markdown and "digraph" not in rendered.markdown
        from PIL import Image

        assert Image.open(path).info["dpi"][0] == pytest.approx(192, abs=1)

    def test_a_deck_gets_the_picture_without_a_div(self, tmp_path, host):
        """pandoc's PowerPoint writer prints the alt text of a picture inside a div, instead of the picture."""
        _, runtime = host
        rendered = render_parts(part(tmp_path, f"```dot\n{self.DOT}\n```\n"), "pptx", runtime)
        assert ":::" not in rendered.markdown and "![dot diagram](diagrams/" in rendered.markdown

    def test_html_gets_svg(self, tmp_path, host):
        _, runtime = host
        rendered = render_parts(part(tmp_path, f"```dot\n{self.DOT}\n```\n"), "html", runtime)
        (path,) = rendered.pictures.values()
        assert path.suffix == ".svg" and b"<svg" in path.read_bytes()

    def test_an_unchanged_diagram_is_not_drawn_again(self, tmp_path, host, caplog):
        _, runtime = host
        text = f"```dot\n{self.DOT}\n```\n"
        with caplog.at_level(logging.INFO):
            render_parts(part(tmp_path, text), "docx", runtime)
            assert "drawn" in caplog.text and "reused" not in caplog.text
            caplog.clear()
            render_parts(part(tmp_path, text), "docx", runtime)
        assert "reused from the cache" in caplog.text and "drawn" not in caplog.text

    def test_a_changed_diagram_or_format_is_drawn_anew(self, tmp_path, host, cache):
        _, runtime = host
        render_parts(part(tmp_path, f"```dot\n{self.DOT}\n```\n"), "docx", runtime)
        render_parts(part(tmp_path, "```dot\ndigraph { a -> c }\n```\n"), "docx", runtime)
        render_parts(part(tmp_path, f"```dot\n{self.DOT}\n```\n"), "html", runtime)
        assert len(list(cache.iterdir())) == 3

    def test_a_broken_diagram_names_file_line_and_the_tool_message(self, tmp_path, host):
        _, runtime = host
        text = f"```dot\n{self.DOT}\n```\n\ntext\n\n```dot\ndigraph {{ a -> \n```\n"
        with pytest.raises(DiagramError) as exc:
            render_parts(part(tmp_path, text), "docx", runtime)
        message = str(exc.value)
        assert f"{tmp_path / 'doc.md'}:7: dot diagram failed" in message and "syntax error" in message

    def test_keep_going_lists_every_failure(self, tmp_path, host):
        _, runtime = host
        bad = "```dot\ndigraph { a -> \n```\n\n"
        with pytest.raises(DiagramError) as exc:
            render_parts(part(tmp_path, bad + bad), "docx", runtime, keep_going=True)
        assert "2 diagram(s) failed" in str(exc.value) and "doc.md:1:" in str(exc.value) and "doc.md:5:" in str(exc.value)

    @needs_pandoc
    def test_a_broken_diagram_stops_the_build_and_writes_nothing(self, tmp_path, host):
        _, runtime = host
        (tmp_path / "g.md").write_text("# G\n\n```dot\ndigraph { a -> b }\n```\n\n```dot\ndigraph { a -> \n```\n", encoding="utf-8")
        (tmp_path / "py-doc.yml").write_text("documents:\n  g:\n    sources: [g.md]\n    outputs:\n      docx: {}\n", encoding="utf-8")
        config = load_config(tmp_path / "py-doc.yml")
        with pytest.raises(DiagramError, match=r"g\.md:7"):
            build_document(config, config.documents["g"], runtime=runtime)
        assert not (tmp_path / "build").exists()

    @needs_pandoc
    def test_the_picture_reaches_the_document(self, tmp_path, host):
        from docx import Document

        _, runtime = host
        (tmp_path / "g.md").write_text(f"# G\n\n```dot title=\"Flow\"\n{self.DOT}\n```\n", encoding="utf-8")
        (tmp_path / "py-doc.yml").write_text("documents:\n  g:\n    sources: [g.md]\n    outputs:\n      docx: {}\n      html: {}\n", encoding="utf-8")
        config = load_config(tmp_path / "py-doc.yml")
        docx, html = build_document(config, config.documents["g"], runtime=runtime)
        assert len(Document(str(docx)).inline_shapes) == 1
        assert "data:image/svg+xml" in html.read_text(encoding="utf-8")


DOT_FENCE = "```dot\ndigraph { a -> b }\n```\n"


@needs_pandoc
@pytest.mark.parametrize(
    ("text", "expected"),
    [
        (f"# G\n\n{DOT_FENCE}\ntext\n", "dot diagram"),
        (f"# G\n\n[[CAPTION:Figure]] Mine\n\n{DOT_FENCE}\ntext\n", "Mine"),
        (f"# G\n\n{DOT_FENCE}\n[[CAPTION:Figure]] Mine\n\ntext\n", "Mine"),
    ],
    ids=["none", "above", "below"],
)
def test_a_diagram_has_one_caption_its_own_or_else_the_tools(tmp_path, pandoc_runtime, monkeypatch, text, expected):
    from docx import Document
    from pptx import Presentation

    picture = png(tmp_path / "pic.png")
    monkeypatch.setattr(diagrams.render, "draw_cached", lambda *args, **kwargs: picture)
    (tmp_path / "g.md").write_text(text, encoding="utf-8")
    (tmp_path / "py-doc.yml").write_text("documents:\n  g:\n    sources: [g.md]\n    outputs: {docx: {}, html: {}, pptx: {}}\n", encoding="utf-8")
    config = load_config(tmp_path / "py-doc.yml")
    docx, html, pptx = sorted(build_document(config, config.documents["g"], runtime=lambda: pandoc_runtime), key=lambda p: p.suffix)
    captions = [p.text for p in Document(str(docx)).paragraphs if p.style.name == "Caption" or p.text.startswith("Figure")]
    assert len(captions) == 1 and captions[0].endswith(expected)
    body = html.read_text(encoding="utf-8")
    assert re.findall(r"<figcaption[^>]*>(.*?)</figcaption>", body, re.S) == [expected] and "[[CAPTION" not in body
    shown = [shape.text_frame.text for slide in Presentation(str(pptx)).slides for shape in slide.shapes if shape.has_text_frame]
    assert shown.count(expected) == 1


class TestOtherTools:
    """The tools themselves live only in the image, so a stand-in checks the command that is run."""

    def test_plantuml_gets_its_include_folder_and_the_text(self, tmp_path, host, monkeypatch):
        _, runtime = host
        monkeypatch.setenv("STUB_LOG", str(tmp_path / "stub.log"))
        png_bytes = png(tmp_path / "x.png").read_bytes()
        stub(tmp_path, monkeypatch, "plantuml", PNG_STUB + f"open(sys.argv[-1].rsplit('.', 1)[0] + '.png', 'wb').write({png_bytes!r})")
        text = "```plantuml\n@startuml\nA -> B\n@enduml\n```\n"
        rendered = render_parts(part(tmp_path, text), "pdf", runtime)
        call = json.loads((tmp_path / "stub.log").read_text(encoding="utf-8").splitlines()[0])
        assert call["argv"][:3] == ["-charset", "UTF-8", "-tpng"]
        assert call["java"].replace("\\", "/") == f"-Dplantuml.include.path={tmp_path.as_posix()}"
        assert len(rendered.pictures) == 1

    def test_plantuml_with_an_include_is_drawn_every_time(self, tmp_path, host, monkeypatch):
        _, runtime = host
        monkeypatch.setenv("STUB_LOG", str(tmp_path / "stub.log"))
        png_bytes = png(tmp_path / "x.png").read_bytes()
        stub(tmp_path, monkeypatch, "plantuml", PNG_STUB + f"open(sys.argv[-1].rsplit('.', 1)[0] + '.png', 'wb').write({png_bytes!r})")
        text = "```plantuml\n!include part.puml\nA -> B\n```\n"
        for _ in range(2):
            render_parts(part(tmp_path, text), "docx", runtime)
        assert len((tmp_path / "stub.log").read_text(encoding="utf-8").splitlines()) == 2

    def test_a_tool_that_writes_nothing_is_an_error(self, tmp_path, host, monkeypatch):
        _, runtime = host
        stub(tmp_path, monkeypatch, "ditaa", "pass")
        with pytest.raises(DiagramError, match="ditaa diagram failed: the tool finished but produced no picture"):
            render_parts(part(tmp_path, "```ditaa\n+--+\n```\n"), "docx", runtime)

    def test_a_file_that_is_not_a_picture_is_an_error(self, tmp_path, host, monkeypatch):
        _, runtime = host
        stub(tmp_path, monkeypatch, "ditaa", "open(sys.argv[2], 'w').write('oops')")
        with pytest.raises(DiagramError, match="not a PNG picture"):
            render_parts(part(tmp_path, "```ditaa\n+--+\n```\n"), "docx", runtime)

    def test_the_orientation_marker_chooses_the_mermaid_width(self, tmp_path, host, monkeypatch):
        _, runtime = host
        monkeypatch.setenv("STUB_LOG", str(tmp_path / "stub.log"))
        png_bytes = png(tmp_path / "x.png").read_bytes()
        stub(tmp_path, monkeypatch, "mmdc", PNG_STUB + f"open(sys.argv[sys.argv.index('-o') + 1], 'wb').write({png_bytes!r})")
        text = "```mermaid\ngraph TD;A-->B\n```\n\n[[LANDSCAPE]]\n\n```mermaid\ngraph TD;A-->C\n```\n\n[[PORTRAIT]]\n\n```mermaid\ngraph TD;A-->D\n```\n"
        render_parts(part(tmp_path, text), "docx", runtime)
        widths = [c["argv"][c["argv"].index("--width") + 1] for c in map(json.loads, (tmp_path / "stub.log").read_text(encoding="utf-8").splitlines())]
        assert widths == ["1200", "1800", "1200"]

    def test_the_container_runs_diagram_tools_without_a_network(self, tmp_path, monkeypatch):
        engine = install(tmp_path, monkeypatch)
        runtime = Runtime("tools:test", docker=engine.docker, platform_name="linux")
        with pytest.raises(DiagramError):
            render_parts(part(tmp_path, "```dot\ndigraph{}\n```\n"), "docx", lambda: runtime)
        (run,) = engine.runs()
        assert run[run.index("--network") + 1] == "none"
        assert any(w.endswith("readonly") and "diagrams/tools" in w for w in run if w.startswith("type=bind"))


def test_png_is_shrunk_to_the_page_and_given_a_density(tmp_path):
    from PIL import Image

    wide = tmp_path / "wide.png"
    Image.new("RGB", (3000, 1500), "white").save(wide)
    _finish_png(wide, landscape=False)
    with Image.open(wide) as image:
        assert image.size == (1200, 600) and round(image.info["dpi"][0]) == 192


@pytest.mark.parametrize("script", ["pikchr_render.py", "railroad_render.py"])
def test_the_python_tools_write_svg(tmp_path, script):
    module = {"pikchr_render.py": "pypikchr", "railroad_render.py": "railroad"}[script]
    pytest.importorskip(module)
    source = tmp_path / "in.txt"
    source.write_text('box "hi"' if module == "pypikchr" else 'Sequence\n  Terminal "SELECT"\n  Optional\n    Terminal "*"\n', encoding="utf-8")
    out = tmp_path / "out.svg"
    done = subprocess.run([sys.executable, str(TOOLS / script), str(source), str(out)], capture_output=True, text=True)
    assert done.returncode == 0, done.stderr
    assert "<svg" in out.read_text(encoding="utf-8")


def test_a_pikchr_syntax_error_fails_loudly(tmp_path):
    pytest.importorskip("pypikchr")
    source = tmp_path / "in.txt"
    source.write_text("this is not pikchr (", encoding="utf-8")
    done = subprocess.run([sys.executable, str(TOOLS / "pikchr_render.py"), str(source), str(tmp_path / "o.svg")], capture_output=True, text=True)
    assert done.returncode != 0 and not (tmp_path / "o.svg").exists()


def test_the_package_exports():
    assert diagrams.render_parts is render_parts


def test_the_svg_to_png_step_falls_back_to_librsvg_when_inkscape_writes_nothing(tmp_path):
    import os
    import subprocess

    from py_doc.diagrams.registry import _SVG_TO_PNG

    bin_dir = tmp_path / "bin"
    bin_dir.mkdir()
    (bin_dir / "inkscape").write_text("#!/bin/sh\nexit 0\n")  # finishes, writes nothing
    (bin_dir / "rsvg-convert").write_text('#!/bin/sh\nwhile [ $# -gt 0 ]; do if [ "$1" = -o ]; then echo png > "$2"; fi; shift; done\n')
    for tool in bin_dir.iterdir():
        tool.chmod(0o755)
    target = tmp_path / "diagram.png"
    done = subprocess.run(
        ["sh", "-c", _SVG_TO_PNG, "x", str(tmp_path / "diagram.svg"), str(target)],
        env={**os.environ, "PATH": f"{bin_dir}:{os.environ['PATH']}"}, capture_output=True, text=True,
    )
    assert done.returncode == 0 and target.read_text() == "png\n" and "librsvg" in done.stderr
