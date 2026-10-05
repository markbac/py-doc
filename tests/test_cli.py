"""Integration tests for the py-doc2docx command line interface."""

from __future__ import annotations

import subprocess
import sys

import pytest
from docx import Document
from docx_helpers import texts
from py_doc.cli import main


def run_cli(monkeypatch, *args: str) -> None:
    monkeypatch.setattr(sys, "argv", ["py-doc2docx", *args])
    main()


@pytest.fixture
def md_tree(tmp_path):
    """docs/ with a root file, a nested file and a non-Markdown file."""
    root = tmp_path / "docs"
    (root / "guide").mkdir(parents=True)
    (root / "index.md").write_text("# Index\n", encoding="utf-8")
    (root / "guide" / "setup.md").write_text("# Setup\n\n- Step\n", encoding="utf-8")
    (root / "notes.txt").write_text("not markdown", encoding="utf-8")
    return root


class TestFileMode:
    def test_default_output_is_next_to_input(self, tmp_path, monkeypatch):
        src = tmp_path / "spec.md"
        src.write_text("# Spec\n", encoding="utf-8")
        run_cli(monkeypatch, str(src))
        assert texts(tmp_path / "spec.docx") == ["Spec"]

    def test_explicit_output_path(self, tmp_path, monkeypatch):
        src = tmp_path / "spec.md"
        src.write_text("# Spec\n", encoding="utf-8")
        run_cli(monkeypatch, str(src), "-o", str(tmp_path / "out" / "custom.docx"))
        assert texts(tmp_path / "out" / "custom.docx") == ["Spec"]
        assert not (tmp_path / "spec.docx").exists()

    def test_template_option_is_applied(self, tmp_path, monkeypatch):
        template = tmp_path / "template.docx"
        doc = Document()
        doc.core_properties.title = "Corporate template"
        doc.save(str(template))
        src = tmp_path / "spec.md"
        src.write_text("# Spec\n", encoding="utf-8")
        run_cli(monkeypatch, str(src), "-t", str(template))
        assert Document(str(tmp_path / "spec.docx")).core_properties.title == "Corporate template"

    def test_template_body_is_removed_unless_keep_template_body_is_given(self, tmp_path, monkeypatch):
        template = tmp_path / "template.docx"
        doc = Document()
        doc.add_paragraph("Template boilerplate")
        doc.save(str(template))
        src = tmp_path / "spec.md"
        src.write_text("# Spec\n", encoding="utf-8")
        run_cli(monkeypatch, str(src), "-t", str(template))
        assert texts(tmp_path / "spec.docx") == ["Spec"]
        run_cli(monkeypatch, str(src), "-t", str(template), "--keep-template-body", "-o", str(tmp_path / "kept.docx"))
        assert texts(tmp_path / "kept.docx") == ["Template boilerplate", "Spec"]

    def test_invalid_utf8_input_exits_with_status_1_and_no_traceback(self, tmp_path, monkeypatch, capsys):
        src = tmp_path / "spec.md"
        src.write_bytes(b"caf\xe9\n")
        with pytest.raises(SystemExit) as exc:
            run_cli(monkeypatch, str(src))
        assert exc.value.code == 1
        assert "not valid UTF-8" in capsys.readouterr().err
        assert not (tmp_path / "spec.docx").exists()

    def test_missing_template_exits_with_status_1_and_writes_nothing(self, tmp_path, monkeypatch, capsys):
        src = tmp_path / "spec.md"
        src.write_text("# Spec\n", encoding="utf-8")
        with pytest.raises(SystemExit) as exc:
            run_cli(monkeypatch, str(src), "-t", str(tmp_path / "missing.docx"))
        assert exc.value.code == 1
        assert "template not found" in capsys.readouterr().err.lower()
        assert not (tmp_path / "spec.docx").exists()

    def test_allow_missing_template_flag_falls_back(self, tmp_path, monkeypatch):
        src = tmp_path / "spec.md"
        src.write_text("# Spec\n", encoding="utf-8")
        run_cli(monkeypatch, str(src), "-t", str(tmp_path / "missing.docx"), "--allow-missing-template")
        assert texts(tmp_path / "spec.docx") == ["Spec"]

    def test_missing_input_exits_with_status_1(self, tmp_path, monkeypatch, capsys):
        with pytest.raises(SystemExit) as exc:
            run_cli(monkeypatch, str(tmp_path / "nope.md"))
        assert exc.value.code == 1
        assert "not found" in capsys.readouterr().err.lower()

    def test_missing_argument_exits_with_usage_error(self, monkeypatch):
        with pytest.raises(SystemExit) as exc:
            run_cli(monkeypatch)
        assert exc.value.code == 2

    def test_module_entry_point_prints_help(self):
        result = subprocess.run(
            [sys.executable, "-m", "py_doc.cli", "--help"],
            capture_output=True,
            text=True,
            check=False,
        )
        assert result.returncode == 0
        assert "Markdown" in result.stdout


class TestDirectoryMode:
    def test_converts_markdown_files_in_place_by_default(self, md_tree, monkeypatch):
        run_cli(monkeypatch, str(md_tree))
        assert texts(md_tree / "index.docx") == ["Index"]
        assert texts(md_tree / "guide" / "setup.docx") == ["Setup", "Step"]

    def test_output_directory_mirrors_input_structure(self, md_tree, tmp_path, monkeypatch):
        out = tmp_path / "dist"
        run_cli(monkeypatch, str(md_tree), "-o", str(out))
        assert texts(out / "index.docx") == ["Index"]
        assert texts(out / "guide" / "setup.docx") == ["Setup", "Step"]

    def test_non_markdown_files_are_ignored(self, md_tree, tmp_path, monkeypatch):
        out = tmp_path / "dist"
        run_cli(monkeypatch, str(md_tree), "-o", str(out))
        assert not list(out.rglob("notes*"))


class TestOutputPaths:
    def test_output_naming_an_existing_directory_gets_the_file_inside(self, tmp_path, monkeypatch):
        src = tmp_path / "spec.md"
        src.write_text("# Spec\n", encoding="utf-8")
        out = tmp_path / "out"
        out.mkdir()
        run_cli(monkeypatch, str(src), "-o", str(out))
        assert texts(out / "spec.docx") == ["Spec"]

    def test_output_with_trailing_separator_is_a_new_directory(self, tmp_path, monkeypatch):
        src = tmp_path / "spec.md"
        src.write_text("# Spec\n", encoding="utf-8")
        run_cli(monkeypatch, str(src), "-o", str(tmp_path / "new") + "/")
        assert texts(tmp_path / "new" / "spec.docx") == ["Spec"]

    def test_markdown_extension_in_any_case_is_converted_in_a_directory(self, tmp_path, monkeypatch):
        root = tmp_path / "docs"
        root.mkdir()
        (root / "A.MD").write_text("# A\n", encoding="utf-8")
        (root / "b.markdown").write_text("# B\n", encoding="utf-8")
        run_cli(monkeypatch, str(root))
        assert texts(root / "A.docx") == ["A"]
        assert texts(root / "b.docx") == ["B"]

    def test_hidden_and_vendor_directories_are_skipped(self, md_tree, monkeypatch):
        for name in (".git", "node_modules"):
            (md_tree / name).mkdir()
            (md_tree / name / "skip.md").write_text("# Skip\n", encoding="utf-8")
        run_cli(monkeypatch, str(md_tree))
        assert not (md_tree / ".git" / "skip.docx").exists()
        assert not (md_tree / "node_modules" / "skip.docx").exists()

    def test_output_directory_inside_the_input_is_not_walked(self, md_tree, monkeypatch):
        (md_tree / "dist").mkdir()
        (md_tree / "dist" / "old.md").write_text("# Old\n", encoding="utf-8")
        (md_tree / "out").mkdir()
        (md_tree / "out" / "copied.md").write_text("# Copied\n", encoding="utf-8")
        run_cli(monkeypatch, str(md_tree), "-o", str(md_tree / "out"))
        assert (md_tree / "out" / "index.docx").exists()
        assert not (md_tree / "out" / "out").exists()

    def test_directory_output_that_is_a_file_is_an_error(self, md_tree, tmp_path, monkeypatch, capsys):
        target = tmp_path / "file.docx"
        target.write_text("x", encoding="utf-8")
        with pytest.raises(SystemExit) as exc:
            run_cli(monkeypatch, str(md_tree), "-o", str(target))
        assert exc.value.code == 1
        assert "is a file" in capsys.readouterr().err

    def test_two_sources_with_one_destination_are_refused_before_converting(self, tmp_path, monkeypatch, capsys):
        root = tmp_path / "docs"
        root.mkdir()
        (root / "a.md").write_text("# One\n", encoding="utf-8")
        (root / "a.markdown").write_text("# Two\n", encoding="utf-8")
        with pytest.raises(SystemExit) as exc:
            run_cli(monkeypatch, str(root))
        assert exc.value.code == 1
        assert "would both be written" in capsys.readouterr().err
        assert not (root / "a.docx").exists()

    def test_empty_directory_warns_and_succeeds(self, tmp_path, monkeypatch, capsys):
        run_cli(monkeypatch, str(tmp_path))
        assert "No Markdown files found" in capsys.readouterr().err


class TestFailures:
    def test_missing_input_creates_no_output_directory(self, tmp_path, monkeypatch):
        with pytest.raises(SystemExit):
            run_cli(monkeypatch, str(tmp_path / "nope.md"), "-o", str(tmp_path / "out" / "x.docx"))
        assert not (tmp_path / "out").exists()

    def test_one_bad_file_does_not_stop_the_batch(self, md_tree, tmp_path, monkeypatch, capsys):
        (md_tree / "bad.md").write_bytes(b"caf\xe9\n")
        with pytest.raises(SystemExit) as exc:
            run_cli(monkeypatch, str(md_tree), "-o", str(tmp_path / "dist"))
        assert exc.value.code == 1
        err = capsys.readouterr().err
        assert "bad.md" in err
        assert "2 of 3 file(s) converted to Word, 1 failed" in err
        assert texts(tmp_path / "dist" / "index.docx") == ["Index"]
        assert texts(tmp_path / "dist" / "guide" / "setup.docx") == ["Setup", "Step"]

    def test_unexpected_error_in_one_file_is_reported_and_the_batch_goes_on(
        self, md_tree, tmp_path, monkeypatch, capsys
    ):
        from py_doc.render.docx import DocxConverter

        real = DocxConverter.convert_file

        def flaky(self, md_path, output_path):
            if md_path.name == "index.md":
                raise RuntimeError("boom")
            return real(self, md_path, output_path)

        monkeypatch.setattr(DocxConverter, "convert_file", flaky)
        with pytest.raises(SystemExit) as exc:
            run_cli(monkeypatch, str(md_tree), "-o", str(tmp_path / "dist"))
        assert exc.value.code == 1
        assert "RuntimeError: boom" in capsys.readouterr().err
        assert (tmp_path / "dist" / "guide" / "setup.docx").exists()

    def test_unwritable_destination_is_a_clean_error(self, tmp_path, monkeypatch, capsys):
        src = tmp_path / "spec.md"
        src.write_text("# Spec\n", encoding="utf-8")
        blocker = tmp_path / "blocker"
        blocker.write_text("a file, not a directory", encoding="utf-8")
        with pytest.raises(SystemExit) as exc:
            run_cli(monkeypatch, str(src), "-o", str(blocker / "out.docx"))
        assert exc.value.code == 1
        assert "Cannot write" in capsys.readouterr().err

    def test_template_that_is_not_a_word_file_is_one_clean_error(self, md_tree, tmp_path, monkeypatch, capsys):
        template = tmp_path / "template.docx"
        template.write_text("not a zip", encoding="utf-8")
        with pytest.raises(SystemExit) as exc:
            run_cli(monkeypatch, str(md_tree), "-t", str(template))
        assert exc.value.code == 1
        err = capsys.readouterr().err
        assert err.count("not a valid .docx or .dotx file") == 1
        assert not (md_tree / "index.docx").exists()


class TestOptions:
    def test_version_comes_from_the_package(self, monkeypatch, capsys):
        import py_doc

        with pytest.raises(SystemExit) as exc:
            run_cli(monkeypatch, "--version")
        assert exc.value.code == 0
        assert capsys.readouterr().out.strip() == f"py-doc2docx {py_doc.__version__}"

    def test_quiet_hides_progress_but_keeps_warnings(self, tmp_path, monkeypatch, capsys):
        src = tmp_path / "spec.md"
        src.write_text("# Spec\n\n![x](missing.png)\n", encoding="utf-8")
        run_cli(monkeypatch, str(src), "-q")
        err = capsys.readouterr().err
        assert "Converting" not in err
        assert "missing.png" in err

    def test_progress_is_shown_by_default(self, tmp_path, monkeypatch, capsys):
        src = tmp_path / "spec.md"
        src.write_text("# Spec\n", encoding="utf-8")
        run_cli(monkeypatch, str(src))
        assert "Converting" in capsys.readouterr().err

    def test_quiet_and_verbose_cannot_be_combined(self, monkeypatch):
        with pytest.raises(SystemExit) as exc:
            run_cli(monkeypatch, "x.md", "-q", "-v")
        assert exc.value.code == 2

    def test_compat_module_entry_point_runs(self, tmp_path):
        result = subprocess.run(
            [sys.executable, "-m", "pydoc2docx", "--version"], capture_output=True, text=True, check=False
        )
        assert result.returncode == 0
        assert "py-doc2docx" in result.stdout
