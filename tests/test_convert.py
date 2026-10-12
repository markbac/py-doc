"""The single-file converters, and the options of the retired native renderers."""

from __future__ import annotations

import subprocess
import zipfile

import pytest
from docx import Document
from pptx import Presentation

from py_doc.config import ConfigError, parse_config
from py_doc.convert import DocxConverter, SlidesConverter
from py_doc.errors import ConversionError


@pytest.mark.parametrize("option", ["engine", "keep_template_body", "keep_template_slides", "title_layout", "content_layout"])
def test_retired_options_are_refused_with_what_to_do(tmp_path, option):
    fmt = "pptx" if option in ("keep_template_slides", "title_layout", "content_layout") else "docx"
    data = {"documents": {"g": {"sources": ["a.md"], "outputs": {fmt: {option: "x"}}}}}
    with pytest.raises(ConfigError, match=f"{option} is no longer an option"):
        parse_config(data, tmp_path)


def make_dotx(path):
    Document().save(str(path))
    data = {}
    with zipfile.ZipFile(path) as z:
        data = {i.filename: z.read(i.filename) for i in z.infolist()}
    data["[Content_Types].xml"] = data["[Content_Types].xml"].replace(b"document.main+xml", b"template.main+xml")
    with zipfile.ZipFile(path, "w") as z:
        for name, content in data.items():
            z.writestr(name, content)


def test_a_dotx_template_is_accepted(tmp_path, pandoc_runtime):
    template = tmp_path / "t.dotx"
    make_dotx(template)
    (tmp_path / "a.md").write_text("# A\n\ntext\n", encoding="utf-8")
    out = DocxConverter(template_path=template).convert_file(tmp_path / "a.md", tmp_path / "a.docx")
    assert "text" in [p.text for p in Document(str(out)).paragraphs]


def test_a_potx_template_is_accepted(tmp_path, pandoc_runtime):
    template = tmp_path / "t.pptx"
    template.write_bytes(subprocess.run(["pandoc", "--print-default-data-file", "reference.pptx"], capture_output=True, check=True).stdout)
    with zipfile.ZipFile(template) as z:
        data = {i.filename: z.read(i.filename) for i in z.infolist()}
    data["[Content_Types].xml"] = data["[Content_Types].xml"].replace(b"presentation.main+xml", b"template.main+xml")
    potx = tmp_path / "t.potx"
    with zipfile.ZipFile(potx, "w") as z:
        for name, content in data.items():
            z.writestr(name, content)
    (tmp_path / "a.md").write_text("# A\n\n## B\n\n- x\n", encoding="utf-8")
    out = SlidesConverter(template_path=potx).convert_file(tmp_path / "a.md", tmp_path / "a.pptx")
    assert len(Presentation(str(out)).slides) >= 1


def test_template_checks(tmp_path):
    with pytest.raises(FileNotFoundError, match="Word template not found"):
        DocxConverter(template_path=tmp_path / "gone.docx").validate_template()
    DocxConverter(template_path=tmp_path / "gone.docx", allow_missing_template=True).validate_template()
    wrong = tmp_path / "t.txt"
    wrong.write_text("x", encoding="utf-8")
    with pytest.raises(ConversionError, match="must be one of"):
        SlidesConverter(template_path=wrong).validate_template()


def test_the_title_comes_from_the_first_heading_without_repeating_it(tmp_path, pandoc_runtime):
    (tmp_path / "a.md").write_text("# Only heading\n\ntext\n", encoding="utf-8")
    out = DocxConverter().convert_file(tmp_path / "a.md", tmp_path / "a.docx")
    document = Document(str(out))
    assert document.core_properties.title == "Only heading"
    assert [p.text for p in document.paragraphs].count("Only heading") == 1


def test_split_level_is_passed_to_pandoc(tmp_path, pandoc_runtime):
    (tmp_path / "a.md").write_text("# A\n\n## B\n\ntext\n\n## C\n\ntext\n", encoding="utf-8")
    one = SlidesConverter(split_level=1).convert_file(tmp_path / "a.md", tmp_path / "one.pptx")
    two = SlidesConverter(split_level=2).convert_file(tmp_path / "a.md", tmp_path / "two.pptx")
    assert len(Presentation(str(two)).slides) > len(Presentation(str(one)).slides)


def test_a_template_without_the_layouts_pandoc_needs_is_explained(tmp_path, pandoc_runtime):
    version = subprocess.run(["pandoc", "--version"], capture_output=True, text=True).stdout.split()[1]
    if tuple(int(n) for n in version.split(".")[:2]) >= (3, 10):
        pytest.skip(f"pandoc {version} falls back to its own layouts. The image's pandoc writes no slides")
    template = tmp_path / "plain.pptx"
    Presentation().save(str(template))  # python-pptx's blank template has other layout names
    (tmp_path / "a.md").write_text("# A\n\n## B\n\n- x\n", encoding="utf-8")
    with pytest.raises(ConversionError, match="no slides"):
        SlidesConverter(template_path=template).convert_file(tmp_path / "a.md", tmp_path / "a.pptx")
