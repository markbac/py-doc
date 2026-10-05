"""Word and PowerPoint template handling: validation, template file types and missing styles."""

from __future__ import annotations

import logging
import zipfile

import pytest
from docx import Document
from docx_helpers import texts
from py_doc import ConversionError, DocxConverter, SlidesConverter
from pptx import Presentation


def _retype(source, target, old: bytes, new: bytes) -> None:
    """Copy an Office file, changing its main content type (a template becomes a document, or back)."""
    with zipfile.ZipFile(source) as zin, zipfile.ZipFile(target, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                assert old in data
                data = data.replace(old, new)
            zout.writestr(item, data)


DOCX = b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
DOTX = b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
PPTX = b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
POTX = b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"


@pytest.fixture
def md(tmp_path):
    path = tmp_path / "spec.md"
    path.write_text("# Spec\n\n- one\n\n```\ncode\n```\n", encoding="utf-8")
    return path


class TestWordTemplates:
    def test_dotx_template_is_accepted(self, tmp_path, md):
        plain = tmp_path / "t.docx"
        doc = Document()
        doc.core_properties.title = "From dotx"
        doc.save(str(plain))
        dotx = tmp_path / "t.dotx"
        _retype(plain, dotx, DOCX, DOTX)
        out = DocxConverter(template_path=dotx).convert_file(md, tmp_path / "out.docx")
        assert Document(str(out)).core_properties.title == "From dotx"
        assert texts(out)[0] == "Spec"

    def test_file_that_is_not_a_zip_is_a_conversion_error(self, tmp_path, md):
        bad = tmp_path / "bad.docx"
        bad.write_text("plain text", encoding="utf-8")
        with pytest.raises(ConversionError, match="not a valid .docx or .dotx"):
            DocxConverter(template_path=bad).convert_file(md, tmp_path / "out.docx")
        assert not (tmp_path / "out.docx").exists()

    def test_zip_that_is_not_word_is_a_conversion_error(self, tmp_path, md):
        bad = tmp_path / "bad.docx"
        with zipfile.ZipFile(bad, "w") as z:
            z.writestr("hello.txt", "hi")
        with pytest.raises(ConversionError, match="not a valid .docx or .dotx"):
            DocxConverter(template_path=bad).convert_file(md, tmp_path / "out.docx")

    def test_validate_template_is_quiet_when_the_template_is_fine_or_absent(self, tmp_path):
        good = tmp_path / "t.docx"
        Document().save(str(good))
        DocxConverter(template_path=good).validate_template()
        DocxConverter().validate_template()
        DocxConverter(template_path=tmp_path / "gone.docx", allow_missing_template=True).validate_template()

    def test_validate_template_rejects_a_missing_template(self, tmp_path):
        with pytest.raises(FileNotFoundError, match="template not found"):
            DocxConverter(template_path=tmp_path / "gone.docx").validate_template()

    def test_missing_styles_are_reported_once_in_one_warning(self, tmp_path, md, caplog):
        template = tmp_path / "t.docx"
        doc = Document()
        for name in ("List Bullet", "Quote"):
            if name in [s.name for s in doc.styles]:
                doc.styles[name].delete()
        doc.save(str(template))
        md.write_text("# Spec\n\n> quoted\n\n- one\n", encoding="utf-8")
        with caplog.at_level(logging.WARNING, logger="py_doc"):
            DocxConverter(template_path=template).convert_file(md, tmp_path / "out.docx")
        warnings = [r.getMessage() for r in caplog.records if r.levelno == logging.WARNING]
        assert warnings == [
            "Template 't.docx' lacks these styles, so plain defaults were used: 'Quote', 'List Bullet'. "
            "The list styles it creates have no bullet or number"
        ]

    def test_code_style_missing_from_the_template_is_not_reported(self, tmp_path, md, caplog):
        template = tmp_path / "t.docx"
        Document().save(str(template))
        with caplog.at_level(logging.WARNING, logger="py_doc"):
            DocxConverter(template_path=template).convert_file(md, tmp_path / "out.docx")
        assert not [r for r in caplog.records if "lacks" in r.getMessage()]

    def test_no_style_report_without_a_template(self, tmp_path, md, caplog):
        with caplog.at_level(logging.WARNING, logger="py_doc"):
            DocxConverter().convert_file(md, tmp_path / "out.docx")
        assert not caplog.records


class TestSlideTemplates:
    def test_potx_template_is_accepted(self, tmp_path, md):
        plain = tmp_path / "t.pptx"
        prs = Presentation()
        prs.core_properties.title = "unused"
        prs.save(str(plain))
        potx = tmp_path / "t.potx"
        _retype(plain, potx, PPTX, POTX)
        out = SlidesConverter(template_path=potx).convert_file(md, tmp_path / "out.pptx")
        assert len(Presentation(str(out)).slides) >= 1

    def test_file_that_is_not_a_zip_is_a_conversion_error(self, tmp_path, md):
        bad = tmp_path / "bad.pptx"
        bad.write_text("plain text", encoding="utf-8")
        with pytest.raises(ConversionError, match="Cannot read PowerPoint template"):
            SlidesConverter(template_path=bad).convert_file(md, tmp_path / "out.pptx")
        assert not (tmp_path / "out.pptx").exists()

    def test_missing_input_creates_no_directory(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            SlidesConverter().convert_file(tmp_path / "nope.md", tmp_path / "out" / "x.pptx")
        assert not (tmp_path / "out").exists()

    def test_validate_template_rejects_a_missing_template(self, tmp_path):
        with pytest.raises(FileNotFoundError, match="template not found"):
            SlidesConverter(template_path=tmp_path / "gone.pptx").validate_template()
