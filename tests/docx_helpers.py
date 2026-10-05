"""Helpers for inspecting generated DOCX files semantically."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def outline(docx_path: Path) -> list[dict]:
    """Return a stable, comparable description of the body paragraphs.

    Font, size and indent are only the direct formatting. What a style provides is checked by name.
    """
    items = []
    for p in Document(str(docx_path)).paragraphs:
        runs = p.runs
        items.append(
            {
                "style": p.style.name,
                "text": p.text,
                "font": runs[0].font.name if runs else None,
                "size_pt": runs[0].font.size.pt if runs and runs[0].font.size else None,
                "left_indent_in": (
                    round(p.paragraph_format.left_indent.inches, 2)
                    if p.paragraph_format.left_indent is not None
                    else None
                ),
            }
        )
    return items


def texts(docx_path: Path) -> list[str]:
    return [i["text"] for i in outline(docx_path)]


def styles(docx_path: Path) -> list[str]:
    return [i["style"] for i in outline(docx_path)]


def hyperlink_targets(docx_path: Path) -> list[str]:
    """The address of every hyperlink in the body, in order."""
    doc = Document(str(docx_path))
    targets = []
    for element in doc.element.body.iter(qn("w:hyperlink")):
        targets.append(doc.part.rels[element.get(qn("r:id"))].target_ref)
    return targets


def table_texts(docx_path: Path) -> list[list[str]]:
    """The text of every cell of the first table, row by row."""
    table = Document(str(docx_path)).tables[0]
    return [[cell.text for cell in row.cells] for row in table.rows]


def all_text(docx_path: Path) -> str:
    """Everything a reader could see or follow: paragraphs, table cells and link addresses."""
    doc = Document(str(docx_path))
    parts = [p.text for p in doc.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    parts += hyperlink_targets(docx_path)
    return "\n".join(parts)
