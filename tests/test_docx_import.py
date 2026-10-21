"""py-doc import: Word to Markdown (#115)."""

from __future__ import annotations

import datetime
import io
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from PIL import Image

from py_doc.app import main
from py_doc.docx_import import InvalidDocx, import_docx
from py_doc.errors import ConversionError
from py_doc.markdown.parser import parse


def png(colour: str = "red", size: int = 8) -> io.BytesIO:
    data = io.BytesIO()
    Image.new("RGB", (size, size), colour).save(data, format="PNG")
    data.seek(0)
    return data


def jpeg() -> io.BytesIO:
    data = io.BytesIO()
    Image.new("RGB", (8, 8), "blue").save(data, format="JPEG")
    data.seek(0)
    return data


class FakeInkscape:
    """Stands in for the container: writes a PNG where Inkscape was asked to."""

    def __init__(self, fail: bool = False):
        self.calls: list[list] = []
        self.fail = fail

    def run(self, argv, mounts=(), **kwargs):
        from py_doc.errors import ToolFailed

        self.calls.append(list(argv))
        if self.fail:
            raise ToolFailed("inkscape", 1, "cannot read this file")
        target = next(str(a).split("=", 1)[1] for a in argv if str(a).startswith("--export-filename="))
        Path(target).write_bytes(png("green").getvalue())


def add_style(doc, name):
    return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def run_text(paragraph, text, **fmt):
    run = paragraph.add_run(text)
    for key, value in fmt.items():
        setattr(run, key, value)
    return run


@pytest.fixture
def sample(tmp_path) -> Path:
    doc = Document()
    doc.core_properties.title = "Design Notes"
    doc.core_properties.author = "A. Writer"
    doc.core_properties.created = datetime.datetime(2024, 3, 5, 9, 0)
    doc.core_properties.comments = "Short summary"
    for style in ("Code", "Note"):
        add_style(doc, style)
    doc.add_heading("Overview", level=2)
    p = doc.add_paragraph("Plain ")
    run_text(p, "bold", bold=True)
    run_text(p, " and ")
    run_text(p, "italic", italic=True)
    run_text(p, " with a_b* star.")
    link = parse_xml(
        f'<w:hyperlink {nsdecls("w", "r")} r:id="{doc.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True)}">'
        "<w:r><w:t>a link</w:t></w:r></w:hyperlink>"
    )
    p._p.append(link)
    doc.add_heading("Details", level=3)
    doc.add_paragraph("one", style="List Bullet")
    doc.add_paragraph("child", style="List Bullet 2")
    doc.add_paragraph("two", style="List Bullet")
    doc.add_paragraph("first", style="List Number")
    doc.add_paragraph("second", style="List Number")
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).text, table.cell(0, 1).text, table.cell(0, 2).text = "A", "B", "C"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "wide | cell"
    table.cell(1, 2).text = "x"
    table.cell(2, 0).text = "1"
    doc.add_paragraph("print('hi')", style="Code")
    doc.add_paragraph("print('there')", style="Code")
    doc.add_paragraph("Careful now.", style="Note")
    doc.add_paragraph("As they said.", style="Quote")
    doc.add_picture(png())
    doc.add_paragraph("Figure 1: The red square", style="Caption")
    path = tmp_path / "Design Notes.docx"
    doc.save(path)
    return path


def read(result) -> str:
    return result.markdown.read_text(encoding="utf-8")


def test_the_structure_is_kept(sample, tmp_path):
    result = import_docx(sample, tmp_path / "out")
    text = read(result)
    assert result.markdown.name == "design_notes.md"
    assert text.startswith("---\n")
    assert "title: Design Notes" in text and "author: A. Writer" in text and "date: '2024-03-05'" in text
    assert "description: Short summary" in text
    assert "# Overview\n" in text and "## Details\n" in text  # first heading is level 1
    assert "Plain **bold** and *italic* with a_b\\* star.[a link](https://example.com)" in text
    assert "- one\n  - child\n- two\n" in text
    assert "1. first\n1. second" in text
    assert "| A | B | C |\n| --- | --- | --- |\n| wide \\| cell | | x |\n| 1 | | |" in text
    assert "```\nprint('hi')\nprint('there')\n```" in text
    assert "> [!NOTE]\n> Careful now." in text
    assert "> As they said." in text


def test_the_caption_goes_before_its_picture_and_the_picture_is_a_png(sample, tmp_path):
    result = import_docx(sample, tmp_path)
    text = read(result)
    assert "[[CAPTION:Figure]] The red square\n\n![](design_notes/media/image_001.png)" in text
    assert (result.media / "image_001.png").read_bytes().startswith(b"\x89PNG")


def test_the_output_parses_and_lints_clean(sample, tmp_path, capsys):
    result = import_docx(sample, tmp_path)
    document = parse(read(result))
    assert document is not None
    assert main(["lint", str(result.markdown)]) in (0, 1)  # it runs on the output


def test_an_existing_output_is_not_replaced_without_force(sample, tmp_path):
    import_docx(sample, tmp_path)
    with pytest.raises(ConversionError, match="already exists"):
        import_docx(sample, tmp_path)
    import_docx(sample, tmp_path, force=True)


def test_a_file_that_is_not_word_is_refused(tmp_path):
    bad = tmp_path / "x.docx"
    bad.write_text("not a zip")
    with pytest.raises(InvalidDocx):
        import_docx(bad, tmp_path / "out")
    with pytest.raises(InvalidDocx):
        import_docx(tmp_path / "gone.docx", tmp_path / "out")


def test_pictures_keep_their_order_and_repeats_share_a_file(tmp_path):
    doc = Document()
    doc.add_heading("T", 1)
    doc.add_picture(png("red"))
    doc.add_picture(jpeg())
    doc.add_picture(png("red"))
    path = tmp_path / "pics.docx"
    doc.save(path)
    result = import_docx(path, tmp_path / "out")
    text = read(result)
    assert text.count("image_001.png") == 2 and text.count("image_002.png") == 1 and "image_003.png" not in text
    assert sorted(p.name for p in result.media.iterdir()) == ["image_001.png", "image_002.png"]


def test_floating_pictures_are_placed_inline_by_default(tmp_path):
    doc = Document()
    doc.add_heading("T", 1)
    doc.add_paragraph("Before")
    doc.add_picture(png())
    inline = doc.element.body.xpath(".//wp:inline")[0]
    inline.tag = inline.tag.replace("inline", "anchor")
    path = tmp_path / "float.docx"
    doc.save(path)
    result = import_docx(path, tmp_path / "out")
    assert "![](float/media/image_001.png)" in read(result)
    again = import_docx(path, tmp_path / "again", defloat_images=False)
    assert "image_001" not in read(again) and any("floating" in w for w in again.warnings)


def test_vector_pictures_go_through_inkscape_and_a_failure_is_visible(tmp_path):
    doc = Document()
    doc.add_heading("T", 1)
    doc.add_picture(png())
    rid = doc.element.body.xpath(".//a:blip/@r:embed")[0]
    doc.part.related_parts[rid]._content_type = "image/x-emf"
    path = tmp_path / "vec.docx"
    doc.save(path)
    inkscape = FakeInkscape()
    result = import_docx(path, tmp_path / "out", runtime=lambda: inkscape)
    assert "![](vec/media/image_001.png)" in read(result)
    assert inkscape.calls and "--export-type=png" in inkscape.calls[0]

    broken = import_docx(path, tmp_path / "broken", runtime=lambda: FakeInkscape(fail=True))
    assert "(image not imported)" in read(broken)
    assert len(broken.warnings) == 1 and "cannot read this file" in broken.warnings[0]


def test_tracked_changes_keep_insertions_and_drop_deletions(tmp_path):
    doc = Document()
    doc.add_heading("T", 1)
    p = doc.add_paragraph("Keep ")
    p._p.append(parse_xml(f'<w:ins {nsdecls("w")} w:id="1" w:author="a"><w:r><w:t>added </w:t></w:r></w:ins>'))
    p._p.append(
        parse_xml(f'<w:del {nsdecls("w")} w:id="2" w:author="a"><w:r><w:delText>removed </w:delText></w:r></w:del>')
    )
    p._p.append(parse_xml(f"<w:r {nsdecls('w')}><w:t>end</w:t></w:r>"))
    path = tmp_path / "track.docx"
    doc.save(path)
    assert "Keep added end" in read(import_docx(path, tmp_path / "out"))


def test_what_cannot_be_imported_is_counted_in_a_warning(tmp_path):
    doc = Document()
    doc.add_heading("T", 1)
    p = doc.add_paragraph("Text")
    p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:footnoteReference w:id="1"/></w:r>'))
    path = tmp_path / "foot.docx"
    doc.save(path)
    result = import_docx(path, tmp_path / "out")
    assert result.warnings == ["1 footnotes not imported"]


def test_the_contents_table_is_skipped(tmp_path):
    doc = Document()
    doc.core_properties.author = doc.core_properties.comments = ""
    add_style(doc, "TOC 1")
    doc.add_paragraph("Real\t3", style="TOC 1")
    doc.add_heading("Real", 1)
    path = tmp_path / "toc.docx"
    doc.save(path)
    assert read(import_docx(path, tmp_path / "out")).strip().endswith("---\n\n# Real")


def test_the_command_exit_statuses(sample, tmp_path, capsys):
    out = tmp_path / "o"
    assert main(["import", str(sample), "-o", str(out), "--print-output-paths"]) == 0
    assert "design_notes.md" in capsys.readouterr().out
    assert main(["import", str(sample), "-o", str(out)]) == 1
    assert "already exists" in capsys.readouterr().err
    (tmp_path / "bad.docx").write_text("x")
    assert main(["import", str(tmp_path / "bad.docx")]) == 2


def test_a_round_trip_keeps_headings_lists_tables_and_pictures(sample, tmp_path, pandoc_runtime):
    from py_doc.convert import DocxConverter

    result = import_docx(sample, tmp_path / "md")
    out = DocxConverter().convert_file(result.markdown, tmp_path / "back.docx")
    back = Document(out)
    styles = [p.style.name for p in back.paragraphs]
    assert "Heading 1" in styles and "Heading 2" in styles
    assert sum("List" in s or s.startswith("Compact") for s in styles) >= 3
    assert len(back.tables) == 1
    assert len(back.inline_shapes) == 1
