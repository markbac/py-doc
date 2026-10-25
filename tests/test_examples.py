"""The examples folder works: the config loads, the templates pass, the documents build and the Word sample imports."""

from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from py_doc.app import main
from py_doc.config import load_config

from conftest import tex_packages_found

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"
# The host has Graphviz and pandoc, but not the other diagram tools, so those blocks are taken out
NO_TOOLS = re.compile(
    r"```(?:mermaid|plantuml|puml|ditaa|blockdiag|seqdiag|actdiag|nwdiag|packetdiag|rackdiag|wavedrom|bytefield-svg|pikchr|railroad|wireviz|vega-lite|vega|chart|cards|circuit|music)\b[^\n]*\n.*?```\n",
    re.S,
)


needs_xelatex = pytest.mark.skipif(
    not tex_packages_found("ulem.sty", "tcolorbox.sty"), reason="xelatex or one of its packages is not installed"
)
# The examples keep their Graphviz diagrams (the fixture strips the other kinds), so building them runs dot
needs_dot = pytest.mark.skipif(shutil.which("dot") is None, reason="Graphviz dot is not installed")


@pytest.fixture
def examples(tmp_path, monkeypatch):
    """A copy of the examples without the diagrams that need tools the host does not have."""
    target = tmp_path / "examples"
    shutil.copytree(
        EXAMPLES, target, ignore=shutil.ignore_patterns("build", ".py-doc", ".structurizr")
    )  # not what a local build left
    for path in (target / "docs").rglob("*.md"):
        path.write_text(NO_TOOLS.sub("", path.read_text(encoding="utf-8")), encoding="utf-8")
    monkeypatch.chdir(target)
    return target


def test_the_config_loads_and_names_every_part():
    config = load_config(EXAMPLES / "py-doc.yml")
    assert sorted(config.documents) == [
        "api-reference",
        "architecture",
        "datasheet",
        "deck",
        "guide",
        "interfaces",
        "release-notes",
    ]
    assert config.documents["architecture"].sources[1].structurizr is not None
    assert config.lint.vale.rules["PyDoc.SentenceLength"].options == {"max": 35}
    assert config.lint.structure
    assert config.report is not None


@pytest.mark.parametrize("name", ["sensor.docx", "sensor_Draft.docx", "sensor.pptx", "sensor_Draft.pptx"])
def test_the_templates_pass_the_check(name, capsys):
    assert main(["template", "check", str(EXAMPLES / "templates" / name), "--strict"]) == 0


def test_the_datasheet_template_passes_with_its_own_prefix_and_marker(capsys):
    path = str(EXAMPLES / "templates" / "datasheet.docx")
    assert main(["template", "check", path, "--strict", "--prefix", "@@", "--marker", "<<BODY>>"]) == 0


def test_the_draft_template_differs_from_the_base_one(capsys):
    assert (
        main(
            [
                "template",
                "diff",
                str(EXAMPLES / "templates" / "sensor.docx"),
                str(EXAMPLES / "templates" / "sensor_Draft.docx"),
            ]
        )
        == 0
    )


def test_the_documents_lint_clean(monkeypatch, capsys):
    monkeypatch.chdir(EXAMPLES)
    assert main(["lint", "docs", "--engine", "builtin", "--strict"]) == 0
    assert "0 issue(s) reported" in capsys.readouterr().out


def test_the_model_checks_without_java(examples, capsys):
    assert main(["c4", "check", "architecture/workspace.json", "--strict"]) == 0
    assert "16 element(s), 10 relationship(s), 5 view(s), 0 error(s), 0 warning(s)" in capsys.readouterr().out


def test_the_dsl_and_the_json_describe_the_same_model():
    dsl = (EXAMPLES / "architecture" / "workspace.dsl").read_text(encoding="utf-8")
    from_dsl = set(re.findall(r'=\s*(?:person|softwareSystem|container|component)\s+"([^"]+)"', dsl))
    data = json.loads((EXAMPLES / "architecture" / "workspace.json").read_text(encoding="utf-8"))
    from_json: set[str] = set()

    def walk(node):
        if isinstance(node, dict):
            if "name" in node and "tags" in node:
                from_json.add(node["name"])
            for value in node.values():
                walk(value)
        elif isinstance(node, list):
            for value in node:
                walk(value)

    walk(data["model"])
    assert from_dsl == from_json


@needs_dot
def test_the_guide_builds_as_word_and_html(examples, pandoc_runtime):
    assert main(["build", "guide", "-f", "docx", "-f", "html", "--condition", "internal=true"]) == 0
    docx = next((examples / "build").glob("*.docx"))
    assert docx.name == "guide-v1.0.docx"
    document = Document(docx)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Sensor Node Guide" in text and "version 1.0" in text and "internal use" in text
    assert "Customers see only the public build." not in text
    assert [p.text for p in document.paragraphs if p.style.name == "Title"] == ["Sensor Node Guide"], (
        "the title page is not repeated"
    )
    assert any(p.style.name == "Lead Paragraph" and p.text.startswith("The SN-100") for p in document.paragraphs)
    assert "DRAFT" in document.sections[0].header.paragraphs[0].text, "status: Draft chooses sensor_Draft.docx"
    assert "TOC" in document.element.xml, "toc: true adds a contents list"
    assert "[[" not in text, "no marker is left in the text"
    html = (examples / "build" / "guide.html").read_text(encoding="utf-8")
    assert "#0b4f8a" in html, "the stylesheet from templates/doc.css is used"
    assert "<figcaption" in html and "<caption>Node requirements</caption>" in html
    assert "[[" not in html, "no marker is printed in HTML"
    assert 'src="data:image/png' in html and 'src="images/logo.png"' not in html, "the logo is embedded"
    assert "<figcaption>Data flow through the node</figcaption>" in html


def test_the_interface_samples_include_every_specification(examples, pandoc_runtime):
    assert main(["build", "interfaces", "-f", "html"]) == 0
    html = (examples / "build" / "interfaces.html").read_text(encoding="utf-8")
    text = re.sub(r"<[^>]+>", "", html)  # highlighting wraps each token in its own span
    for spec in ("openapi: 3.1.0", "asyncapi: 2.6.0", "service LibraryService", "type Query"):
        assert spec in text, f"{spec!r} is included from docs/includes/apis/"
    assert "[[" not in html, "no marker is printed in HTML"


@needs_dot
def test_the_conditions_choose_the_text(examples, pandoc_runtime):
    assert main(["build", "guide", "-f", "html", "--condition", "audience=staff"]) == 0
    html = (examples / "build" / "guide.html").read_text(encoding="utf-8")
    assert "Staff notes" in html and "Customers see only the public build." in html


@needs_xelatex
def test_the_guide_builds_as_pdf_with_task_lists_and_captions(examples, pandoc_runtime):
    assert main(["build", "guide", "-f", "pdf"]) == 0
    pdf = examples / "build" / "guide.pdf"
    assert pdf.read_bytes().startswith(b"%PDF")
    if shutil.which("pdftotext"):
        text = subprocess.run(
            ["pdftotext", "-enc", "UTF-8", str(pdf), "-"], capture_output=True, encoding="utf-8", check=True
        ).stdout
        assert "[[" not in text, "no marker is printed in the PDF"
        assert "Data flow through the node" in text and "Figure 1" in text


def test_the_datasheet_builds_with_its_own_placeholders_and_a_version_in_the_name(examples, pandoc_runtime):
    assert main(["build", "datasheet", "-f", "docx"]) == 0
    docx = examples / "build" / "datasheet-1_0.docx"
    document = Document(docx)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "@@" not in text and "<<BODY>>" not in text
    assert "Block diagram" in text and "Packet layout" in text, "the included diagram chapter is in the datasheet"
    assert "@@" not in "".join(s.footer.paragraphs[0].text for s in document.sections)


def test_the_release_notes_take_their_version_from_git(examples, pandoc_runtime):
    def git(*args):
        subprocess.run(
            ["git", "-c", "user.name=t", "-c", "user.email=t@example.com", *args],
            cwd=examples,
            check=True,
            capture_output=True,
        )

    git("init", "-q")
    git("add", "-A")
    git("commit", "-q", "-m", "x")
    git("tag", "v1.2.3")
    assert main(["build", "release-notes"]) == 0
    assert (examples / "build" / "release-notes-v1_2_3.docx").is_file()


@needs_dot
def test_the_guide_shows_the_word_features(examples, pandoc_runtime):
    assert main(["build", "guide", "-f", "docx"]) == 0
    document = Document(examples / "build" / "guide-v1.0.docx")
    styles = {p.style.name for p in document.paragraphs}
    assert {"Heading 4", "Block Text", "Lead Paragraph"} <= styles
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert "Diagram types" in headings
    xml = document.element.xml
    assert "m:oMath" in xml, "maths becomes Word equations"
    assert "footnoteReference" in xml
    assert 'w:orient="landscape"' in xml, "a landscape section"
    assert "SEQ" in xml, "captions are numbered with fields"
    with zipfile.ZipFile(examples / "build" / "guide-v1.0.docx") as archive:
        core = archive.read("docProps/core.xml").decode("utf-8")
    assert "Integrating the SN-100" in core and "Technical documentation" in core and "IoT, sensor, MQTT" in core


@needs_dot
def test_the_deck_uses_the_layouts_notes_and_footer(examples, pandoc_runtime):
    assert main(["build", "deck"]) == 0
    presentation = Presentation(examples / "build" / "deck.pptx")
    slides = presentation.slides
    assert len(slides) >= 15
    layouts = {slide.slide_layout.name for slide in slides}
    assert {"Title Slide", "Section Header", "Two Content", "Title and Content"} <= layouts
    assert sum(1 for slide in slides if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()) >= 3
    xml = "".join(slide.part.blob.decode("utf-8") for slide in slides)
    assert "CAPTION" not in xml and "[[" not in xml and "{{" not in xml
    assert "Sensor Node" in xml and "v1.0" in xml, "the footer pattern is applied"
    assert "<a:tbl>" in xml, "a table"
    titles = {slide.shapes.title.text: slide.slide_layout.name for slide in slides if slide.shapes.title is not None}
    assert titles["One column"] == "Title and Content"
    assert titles["Two columns of text"] in ("Two Content", "Comparison")
    assert titles["Two columns: text and picture"] in ("Two Content", "Comparison")
    assert any(shape.shape_type == 13 for slide in slides for shape in slide.shapes), "a picture"
    master = "".join(m.part.blob.decode("utf-8") for m in presentation.slide_masters)
    assert "{{VERSION}}" not in master


@needs_dot
def test_a_released_deck_uses_the_base_template(examples, pandoc_runtime):
    slides = examples / "docs" / "deck" / "slides.md"
    slides.write_text(slides.read_text(encoding="utf-8").replace("status: Draft", "status: Released"), encoding="utf-8")
    assert main(["build", "deck"]) == 0
    master = "".join(m.part.blob.decode("utf-8") for m in Presentation(examples / "build" / "deck.pptx").slide_masters)
    assert "DRAFT" not in master and "Version 1.0" in master


@needs_dot
def test_a_draft_html_guide_uses_the_draft_stylesheet(examples, pandoc_runtime):
    assert main(["build", "guide", "-f", "html"]) == 0
    html = (examples / "build" / "guide.html").read_text(encoding="utf-8")
    assert "DRAFT" in html and "#c62828" in html


@needs_dot
def test_a_draft_deck_uses_the_draft_template(examples, pandoc_runtime):
    assert main(["build", "deck"]) == 0  # slides.md says status: Draft
    master = "".join(m.part.blob.decode("utf-8") for m in Presentation(examples / "build" / "deck.pptx").slide_masters)
    assert "DRAFT" in master


@needs_xelatex
def test_the_datasheet_builds_as_pdf_with_its_own_latex_template(examples, pandoc_runtime):
    assert main(["build", "datasheet", "-f", "pdf"]) == 0
    assert (examples / "build" / "datasheet.pdf").read_bytes().startswith(b"%PDF")


def test_the_word_sample_imports_and_builds_again(examples, pandoc_runtime, tmp_path):
    out = tmp_path / "imported"
    assert main(["import", "word/design_review.docx", "-o", str(out)]) == 0
    text = (out / "design_review.md").read_text(encoding="utf-8")
    assert (
        "title: Design review" in text
        and "- Confirm the sampling interval" in text
        and "| Sampling interval | Firmware |" in text
    )
    assert "[[CAPTION:Figure]] The agreed layout" in text
    assert main(["docx", str(out / "design_review.md")]) == 0
    with zipfile.ZipFile(out / "design_review.docx") as archive:
        assert any(n.startswith("word/media/") for n in archive.namelist())


@pytest.mark.skipif(
    not os.environ.get("PY_DOC_STRUCTURIZR_CLI"), reason="set PY_DOC_STRUCTURIZR_CLI to a structurizr-cli to run it"
)
def test_the_real_structurizr_cli_reads_the_dsl_into_the_same_model(tmp_path):
    """#96: the DSL through the real structurizr-cli gives the model the hand-made JSON describes."""
    from py_doc.c4.model import load_workspace
    from py_doc.c4.structurizr import workspace_json

    shutil.copytree(EXAMPLES / "architecture", tmp_path / "architecture", ignore=shutil.ignore_patterns(".structurizr"))
    exported = workspace_json(
        tmp_path / "architecture" / "workspace.dsl", tmp_path / "cache", cli=os.environ["PY_DOC_STRUCTURIZR_CLI"]
    )
    from_cli = load_workspace(exported)
    by_hand = load_workspace(EXAMPLES / "architecture" / "workspace.json")
    assert {e.name for e in from_cli.elements.values()} == {e.name for e in by_hand.elements.values()}

    def described(ws):
        return {
            (ws.elements[r.source_id].name, ws.elements[r.destination_id].name, r.description) for r in ws.relationships
        }

    # Structurizr adds implied relationships and those of deployment instances, so the CLI's model has more
    assert described(by_hand) <= described(from_cli)
    assert {v.key for v in from_cli.views} == {v.key for v in by_hand.views}
