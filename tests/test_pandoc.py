"""The pandoc engine. Pandoc runs for real: the fake container engine executes the tool on the host."""

from __future__ import annotations

import shutil
import struct
import zlib
from pathlib import Path

import pytest
from docx import Document
from fake_engine import install
from pptx import Presentation

from py_doc.build import build_document
from py_doc.config import load_config
from py_doc.errors import ConversionError, ToolFailed
from py_doc.pandoc import source
from py_doc.runtime import Runtime

needs_pandoc = pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc is not installed")
needs_xelatex = pytest.mark.skipif(shutil.which("xelatex") is None, reason="xelatex is not installed")


def png(path: Path) -> Path:
    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    rows = b"".join(b"\x00" + b"\xff\x00\x00" * 4 for _ in range(4))
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 4, 4, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b"")
    )
    return path


# --- joining the sources ---------------------------------------------------------------------------


class TestSource:
    def test_front_matter_is_split_off(self):
        front, body = source.split_front_matter("---\ntitle: T\n---\n\n# H\n")
        assert (front, body) == ("title: T", "# H\n")
        assert source.split_front_matter("# H\n") == (None, "# H\n")
        assert source.split_front_matter("---\nno end\n") == (None, "---\nno end\n")

    def test_headings_move_but_code_does_not(self):
        text = "# A\n\n```bash\n# a comment\n```\n\n## B\n"
        out = source.shift_headings(text, 1, "x.md")
        assert "## A" in out and "### B" in out and "# a comment" in out and "## a comment" not in out

    def test_headings_are_kept_within_six_levels(self, caplog):
        assert source.shift_headings("##### A\n", 3, "x.md") == "###### A\n"
        assert "would be level 8" in caplog.text

    def test_retitle_replaces_the_first_heading_or_adds_one(self):
        assert source.retitle("# Old\n\nText\n", "New") == "# New\n\nText\n"
        assert source.retitle("Text\n", "New") == "# New\n\nText\n"
        assert source.retitle("```\n# not\n```\n\n## Real\n", "New").splitlines()[4] == "## New"

    def test_sources_are_joined_with_their_options(self, tmp_path):
        (tmp_path / "a.md").write_text("---\ntitle: Guide\nauthor: Ann\n---\n# One\n", encoding="utf-8")
        (tmp_path / "b.md").write_text("# Two\n[[IF: x]]\nshown {{v}}\n[[ENDIF]]\n", encoding="utf-8")
        config = load_config(
            _write(
                tmp_path,
                "documents:\n  g:\n    sources:\n      - a.md\n      - {file: b.md, heading_offset: 1, page_break: true}\n",
            )
        )
        joined = source.join_sources(config.documents["g"].sources, {"v": "1"}, {"x": True}, "g")
        assert joined.metadata.title == "Guide" and joined.metadata.author == "Ann"
        assert joined.markdown == "# One\n\n::: {.pagebreak}\n:::\n\n## Two\nshown 1\n"
        assert joined.first_heading == "One"

    def test_a_document_with_no_sources_is_an_error(self):
        with pytest.raises(ConversionError, match="no sources"):
            source.join_sources([], {}, {}, "g")


def _write(folder: Path, text: str) -> Path:
    path = folder / "py-doc.yml"
    path.write_text(text, encoding="utf-8")
    return path


# --- building through the container runner ---------------------------------------------------------


@pytest.fixture
def host(tmp_path, monkeypatch):
    """A runtime whose container is the host: pandoc and xelatex run for real."""
    engine = install(tmp_path, monkeypatch, mode="host")
    return engine, (lambda: Runtime("tools:test", docker=engine.docker, platform_name="linux"))


GUIDE = """---
title: Guide
author: Ann
version: "1.2"
---
# Overview

Some *text* and a table.

| a | b |
|---|---|
| 1 | 2 |

```python
# keep me
print("hi")
```

> [!NOTE]
> A note.

> **Warning:** A warning.

![pic](pic.png)
"""


def project(tmp_path: Path, outputs: str = "", extra: str = ""):
    (tmp_path / "guide.md").write_text(GUIDE, encoding="utf-8")
    png(tmp_path / "pic.png")
    path = _write(tmp_path, f"documents:\n  guide:\n    sources: [guide.md]\n{extra}    outputs:\n{outputs}")
    config = load_config(path)
    return config, config.documents["guide"]


@needs_pandoc
class TestHtml:
    def test_a_self_contained_page_with_boxes_and_the_picture(self, tmp_path, host):
        _, runtime = host
        config, doc = project(tmp_path, "      html: {}\n")
        (out,) = build_document(config, doc, ["html"], runtime=runtime)
        html = out.read_text(encoding="utf-8")
        assert out == tmp_path / "build" / "guide.html"
        assert "<title>Guide</title>" in html and '<h1 id="overview">Overview</h1>' in html
        assert 'class="admonition note"' in html and 'class="admonition warning"' in html
        assert "data:image/png;base64," in html  # embedded, so the file can be moved
        assert "keep me" in html
        assert "<table>" in html

    def test_variables_and_conditions_apply(self, tmp_path, host):
        _, runtime = host
        config, doc = project(tmp_path, "      html: {}\n", "    variables: {who: Ann}\n")
        (tmp_path / "guide.md").write_text("# T\n\nBy {{who}}\n\n[[IF: x]]\nsecret\n[[ENDIF]]\n", encoding="utf-8")
        (out,) = build_document(config, doc, ["html"], conditions={"x": False}, runtime=runtime)
        html = out.read_text(encoding="utf-8")
        assert "By Ann" in html and "secret" not in html


@needs_pandoc
class TestDocx:
    def test_pandoc_word_output(self, tmp_path, host):
        _, runtime = host
        config, doc = project(tmp_path, "      docx: {}\n")
        (out,) = build_document(config, doc, ["docx"], runtime=runtime)
        document = Document(str(out))
        styles = [(p.style.name, p.text) for p in document.paragraphs]
        assert ("Heading 1", "Overview") in styles
        assert any(style == "Div.note" or style.startswith("Div.") for style, _ in styles)
        assert len(document.tables) == 1
        assert len(document.inline_shapes) == 1
        assert document.core_properties.title == "Guide" and document.core_properties.author == "Ann"
        assert any(p.text.startswith("# keep me") for p in document.paragraphs)

    def test_the_template_is_filled_and_receives_the_body(self, tmp_path, host):
        _, runtime = host
        template = Document()
        template.add_paragraph("Cover: Replace-with-Title v Replace-with-Version")
        template.add_paragraph("[[INSERTCONTENT]]")
        template.add_paragraph("Back page")
        template.save(tmp_path / "tpl.docx")
        config, doc = project(tmp_path, "      docx: {template: tpl.docx}\n")
        (out,) = build_document(config, doc, ["docx"], runtime=runtime)
        texts = [p.text for p in Document(str(out)).paragraphs]
        assert texts[0] == "Cover: Guide v 1.2"
        assert texts.index("Overview") < texts.index("Back page")
        assert "[[INSERTCONTENT]]" not in texts

    def test_a_missing_template_is_an_error_unless_allowed(self, tmp_path, host):
        _, runtime = host
        config, doc = project(tmp_path, "      docx: {template: gone.docx}\n")
        with pytest.raises(ConversionError, match="template was not found"):
            build_document(config, doc, ["docx"], runtime=runtime)
        config, doc = project(tmp_path, "      docx: {template: gone.docx, allow_missing_template: true}\n")
        (out,) = build_document(config, doc, ["docx"], runtime=runtime)
        assert out.is_file()


@needs_pandoc
def test_pandoc_slides(tmp_path, host):
    _, runtime = host
    config, doc = project(tmp_path, "      pptx: {}\n")
    (out,) = build_document(config, doc, ["pptx"], runtime=runtime)
    assert len(Presentation(str(out)).slides) >= 1


@needs_pandoc
@needs_xelatex
def test_pdf_through_xelatex(tmp_path, host):
    _, runtime = host
    config, doc = project(tmp_path, "      pdf: {}\n")
    (out,) = build_document(config, doc, ["pdf"], runtime=runtime)
    assert out.read_bytes().startswith(b"%PDF") and out.stat().st_size > 5000


@needs_pandoc
def test_a_pandoc_failure_carries_pandoc_s_message(tmp_path, host):
    _, runtime = host
    (tmp_path / "bad.tex").write_text("\\begin{document} $body$ $\\undefined$", encoding="utf-8")
    config, doc = project(tmp_path, "      pdf: {template: bad.tex}\n")
    with pytest.raises(ToolFailed, match="pandoc"):
        build_document(config, doc, ["pdf"], runtime=runtime)


@needs_pandoc
def test_nothing_is_left_in_the_working_folder(tmp_path, host):
    _, runtime = host
    config, doc = project(tmp_path, "      html: {}\n")
    before = {p.name for p in tmp_path.iterdir()}
    build_document(config, doc, ["html"], runtime=runtime)
    assert {p.name for p in tmp_path.iterdir()} - before == {"build"}


# --- the container command ---------------------------------------------------------------------


def test_the_command_mounts_resources_read_only_and_translates_paths(tmp_path, monkeypatch):
    engine = install(tmp_path, monkeypatch)
    runtime = Runtime("tools:test", docker=engine.docker, platform_name="linux")
    config, doc = project(tmp_path, "      html: {}\n")
    with pytest.raises(ConversionError, match="wrote no html file"):
        build_document(config, doc, ["html"], runtime=lambda: runtime)
    (run,) = [r for r in engine.runs() if "--version" not in r]  # the other run asks pandoc for its version
    assert "--syntax-highlighting" in run and "--highlight-style" not in run
    mounts = [word.replace("\\", "/") for word in run if word.startswith("type=bind")]
    assert any("pandoc/resources" in m and m.endswith("readonly") for m in mounts)
    assert any(f"source={tmp_path.as_posix()}," in m and m.endswith("readonly") for m in mounts)
    assert any("py-doc-" in m and not m.endswith("readonly") for m in mounts)
    resource_path = next(w for w in run if w.startswith("--resource-path=")).replace("\\", "/")
    assert tmp_path.as_posix() in resource_path
    assert run.count("--lua-filter") == 3 and "--embed-resources" in run


def test_pdf_and_html_need_a_runtime(tmp_path):
    config, doc = project(tmp_path, "      html: {}\n")
    with pytest.raises(ConversionError, match="needs the tool container"):
        build_document(config, doc, ["html"])


# --- the py-doc build command ----------------------------------------------------------------------


@needs_pandoc
class TestCommand:
    def test_build_makes_html_through_the_runtime(self, tmp_path, host, monkeypatch, capsys):
        from py_doc import app

        _, runtime = host
        project(tmp_path, "      html: {}\n")
        monkeypatch.chdir(tmp_path)
        monkeypatch.setattr(app, "connect", lambda settings: runtime())
        assert app.main(["build", "-f", "html"]) == 0
        assert (tmp_path / "build" / "guide.html").is_file()

    def test_a_missing_container_engine_is_exit_status_2_with_the_reason(self, tmp_path, monkeypatch, capsys):
        from py_doc import app
        from py_doc.errors import EnvironmentProblem

        project(tmp_path, "      html: {}\n")
        monkeypatch.chdir(tmp_path)

        def refuse(settings):
            raise EnvironmentProblem("Docker is not running")

        monkeypatch.setattr(app, "connect", refuse)
        assert app.main(["build", "-f", "html"]) == 2
        assert "guide: Docker is not running" in capsys.readouterr().err


def test_latex_template_defines_task_list_boxes():
    """pandoc writes \\boxtimes and \\square for `- [x]` items, so the PDF template must load amssymb."""
    template = (Path(source.__file__).parent / "resources" / "latex_template.tex").read_text(encoding="utf-8")
    assert "\\usepackage{amssymb}" in template


def test_pdf_html_and_pptx_turn_caption_markers_into_captions():
    from py_doc.pandoc.engine import FILTERS

    assert "markers.lua" in FILTERS["pdf"]
    assert "markers.lua" in FILTERS["html"]
    assert "markers.lua" in FILTERS["pptx"]
    assert "markers.lua" not in FILTERS["docx"]  # Word numbers the marker itself


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc is not installed")
def test_marker_filter_makes_captions_and_drops_word_only_markers(tmp_path):
    import subprocess

    from py_doc.pandoc.engine import INPUT_FORMAT, RESOURCES

    md = tmp_path / "c.md"
    md.write_text(
        "[[STYLE:Lead Paragraph]] A lead\n\n"
        "[[CAPTION:Figure]] The logo\n![Logo](logo.png)\n\n"
        "[[CAPTION:Figure]] A diagram\n\n::: {style=\"text-align:center;\"}\n![dot diagram](d.png)\n:::\n\n"
        "[[CAPTION:Table]] Requirements\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n"
        "[[LANDSCAPE]]\n\n"
        "| C | D |\n|---|---|\n| 3 | 4 |\n\n[[CAPTION:Table]] After the table\n\n"
        "[[PORTRAIT]]\n\n"
        "[[CAPTION:Figure]] The code\n```c\nint x;\n```\n\n"
        "[[CAPTION:Table]] Alone\n",
        encoding="utf-8",
    )
    done = subprocess.run(
        ["pandoc", str(md), f"--from={INPUT_FORMAT}", f"--lua-filter={RESOURCES / 'markers.lua'}", "--to=html5"],
        capture_output=True, text=True, check=True, cwd=tmp_path,
    )  # fmt: skip
    html = done.stdout
    assert "[[" not in html
    assert "A lead" in html
    for caption in ("The logo", "A diagram", "The code"):
        assert f"<figcaption>{caption}</figcaption>" in html or f"<figcaption aria-hidden=\"true\">{caption}</figcaption>" in html
    assert "<caption>Requirements</caption>" in html and "<caption>After the table</caption>" in html
    assert "Alone" in html


@pytest.mark.parametrize(("banner", "option"), [
    ("pandoc 3.8.2.1", "--syntax-highlighting"), ("pandoc 3.8", "--syntax-highlighting"),
    ("pandoc 3.1.3", "--highlight-style"), ("nonsense", "--syntax-highlighting"),
])  # fmt: skip
def test_the_highlight_option_follows_the_pandoc_version(banner, option):
    from py_doc.pandoc.engine import _highlight_option

    class Fake:
        def run(self, argv, **kwargs):
            from py_doc.runtime import ToolResult

            return ToolResult(0, banner + "\nFeatures: x\n", "")

    assert _highlight_option(Fake()) == option
