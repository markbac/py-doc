"""Building documents from py-doc.yml: joining sources, metadata, page breaks and outputs. Pandoc runs for real."""

from __future__ import annotations

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx_helpers import texts
from pptx_helpers import titles

from py_doc.build import build_document
from py_doc.config import load_config
from py_doc.errors import ConversionError


@pytest.fixture
def project(tmp_path, pandoc_runtime):
    (tmp_path / "docs").mkdir()
    (tmp_path / "docs" / "index.md").write_text("---\ntitle: From front matter\nauthor: Ann\n---\n# Intro\n\nHello.\n", encoding="utf-8")
    (tmp_path / "docs" / "part.md").write_text("# Part\n\n## Detail\n\ntext\n", encoding="utf-8")

    def make(yaml_text: str, name: str = "g", formats=None):
        (tmp_path / "py-doc.yml").write_text(yaml_text, encoding="utf-8")
        config = load_config(tmp_path / "py-doc.yml")
        return build_document(config, config.documents[name], formats, runtime=lambda: pandoc_runtime)

    return make


def page_breaks(path) -> int:
    return sum(1 for br in Document(str(path)).element.body.iter(qn("w:br")) if br.get(qn("w:type")) == "page")


def headings(path) -> list[str]:
    return [p.text for p in Document(str(path)).paragraphs if p.style.name.startswith("Heading")]


def test_sources_are_joined_in_order_with_a_page_break_before_marked_ones(project, tmp_path):
    [out] = project("documents:\n  g:\n    sources:\n      - docs/index.md\n      - {file: docs/part.md, page_break: true}\n")
    assert out == tmp_path / "build" / "g.docx"
    assert headings(out) == ["Intro", "Part", "Detail"]
    assert page_breaks(out) == 1


def test_the_pagebreak_directive_breaks_the_page(project, tmp_path):
    (tmp_path / "docs" / "part.md").write_text("# Part\n\none\n\n[[PAGEBREAK]]\n\ntwo\n", encoding="utf-8")
    [out] = project("documents:\n  g:\n    sources: [docs/part.md]\n")
    assert page_breaks(out) == 1
    assert not any("PAGEBREAK" in text for text in texts(out))


def test_page_break_on_the_first_source_adds_nothing(project):
    [out] = project("documents:\n  g:\n    sources:\n      - {file: docs/index.md, page_break: true}\n")
    assert page_breaks(out) == 0


def test_heading_offset_and_title(project):
    [out] = project("documents:\n  g:\n    sources:\n      - {file: docs/part.md, heading_offset: 1, title: Renamed}\n")
    assert headings(out) == ["Renamed", "Detail"]
    assert [p.style.name for p in Document(str(out)).paragraphs if p.style.name.startswith("Heading")] == ["Heading 2", "Heading 3"]


def test_excluded_sources_are_skipped_and_all_excluded_is_an_error(project):
    [out] = project("documents:\n  g:\n    sources:\n      - docs/index.md\n      - {file: docs/missing.md, include: false}\n")
    assert "Hello." in texts(out)
    with pytest.raises(ConversionError, match="no sources"):
        project("documents:\n  g:\n    sources:\n      - {file: docs/index.md, include: false}\n")


def test_missing_source_names_the_document(project):
    with pytest.raises(FileNotFoundError, match="document 'g' not found"):
        project("documents:\n  g:\n    sources: [docs/nope.md]\n")


def test_front_matter_properties_are_used_and_the_configuration_overrides_them(project):
    [out] = project("documents:\n  g:\n    title: Configured\n    sources: [docs/index.md]\n")
    props = Document(str(out)).core_properties
    assert (props.title, props.author) == ("Configured", "Ann")


def test_title_falls_back_to_the_first_heading(project):
    [out] = project("documents:\n  g:\n    sources: [docs/part.md]\n")
    assert Document(str(out)).core_properties.title == "Part"


def test_images_resolve_against_their_own_source_folder(project, tmp_path):
    from PIL import Image

    (tmp_path / "docs" / "sub").mkdir()
    Image.new("RGB", (4, 4)).save(tmp_path / "docs" / "sub" / "pic.png")
    (tmp_path / "docs" / "sub" / "doc.md").write_text("![](pic.png)\n", encoding="utf-8")
    [out] = project("documents:\n  g:\n    sources: [docs/index.md, docs/sub/doc.md]\n")
    assert len(Document(str(out)).inline_shapes) == 1


def test_formats_listed_by_the_document_are_all_built(project, tmp_path):
    written = project("documents:\n  g:\n    sources: [docs/index.md, docs/part.md]\n    outputs: {docx: {}, pptx: {path: deck/g.pptx}}\n")
    assert [p.name for p in written] == ["g.docx", "g.pptx"]
    assert titles(tmp_path / "deck" / "g.pptx")[0] == "From front matter"


def test_formats_can_be_chosen_at_build_time(project, tmp_path):
    [out] = project("documents:\n  g:\n    sources: [docs/index.md]\n    outputs: {docx: {}}\n", formats=["pptx"])
    assert out.suffix == ".pptx"


def test_a_retired_option_is_explained(project):
    with pytest.raises(Exception, match="pandoc is the only engine"):
        project("documents:\n  g:\n    sources: [docs/index.md]\n    outputs: {docx: {engine: native}}\n")


def test_page_break_starts_a_new_slide(project, tmp_path):
    (tmp_path / "docs" / "bare.md").write_text("Just text.\n", encoding="utf-8")
    counts = []
    for flag in ("false", "true"):
        [out] = project(
            f"documents:\n  g:\n    sources:\n      - docs/part.md\n      - {{file: docs/bare.md, page_break: {flag}}}\n",
            formats=["pptx"],
        )
        counts.append(len(titles(out)))
    assert counts[1] > counts[0]
