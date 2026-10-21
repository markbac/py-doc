"""py-doc import: PowerPoint to Markdown, in the form `py-doc slides` builds again."""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from py_doc.app import main
from py_doc.errors import ConversionError
from py_doc.pptx_import import InvalidPptx, import_pptx

# Layouts of python-pptx's default template
TITLE, TITLE_AND_CONTENT, SECTION, TWO_CONTENT, TITLE_ONLY, BLANK = 0, 1, 2, 3, 5, 6


def png(colour: str = "red") -> io.BytesIO:
    data = io.BytesIO()
    Image.new("RGB", (8, 8), colour).save(data, format="PNG")
    data.seek(0)
    return data


def bullets(slide, lines: list[tuple[int, str]]) -> None:
    frame = slide.placeholders[1].text_frame
    frame.text = lines[0][1]
    frame.paragraphs[0].level = lines[0][0]
    for level, text in lines[1:]:
        paragraph = frame.add_paragraph()
        paragraph.text, paragraph.level = text, level


@pytest.fixture
def deck(tmp_path):
    prs = Presentation()
    prs.core_properties.author = "A. Writer"
    prs.core_properties.version = "2.0"
    first = prs.slides.add_slide(prs.slide_layouts[TITLE])
    first.shapes.title.text = "Sensor review"
    first.placeholders[1].text = "What we found"
    section = prs.slides.add_slide(prs.slide_layouts[SECTION])
    section.shapes.title.text = "Findings"
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_AND_CONTENT])
    slide.shapes.title.text = "Readings"
    bullets(slide, [(0, "Temperature"), (1, "Within 0.5 C"), (0, "Humidity")])
    slide.notes_slide.notes_text_frame.text = "Say why the drift matters."
    two = prs.slides.add_slide(prs.slide_layouts[TWO_CONTENT])
    two.shapes.title.text = "Before and after"
    two.placeholders[1].text = "Polled"
    two.placeholders[2].text = "Pushed"
    table = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    table.shapes.title.text = "Registers"
    grid = table.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for (row, col), text in {(0, 0): "Name", (0, 1): "Address", (1, 0): "TEMP", (1, 1): "0x01|0x02"}.items():
        grid.cell(row, col).text = text
    picture = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    picture.shapes.title.text = "Board"
    added = picture.shapes.add_picture(png(), Inches(1), Inches(2), Inches(2), Inches(2))
    added._element.nvPicPr.cNvPr.set("descr", "The sensor board")
    box = picture.shapes.add_textbox(Inches(1), Inches(4.1), Inches(2), Inches(0.4))
    box.text_frame.text = "Figure caption"
    code = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    code.shapes.title.text = "Code"
    frame = code.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2)).text_frame
    frame.text = "int main(void)"
    frame.add_paragraph().text = "{ return 0; }"
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Courier New"
    chart = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    chart.shapes.title.text = "Trend"
    data = CategoryChartData()
    data.categories = ["a", "b"]
    data.add_series("s", (1, 2))
    chart.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(1), Inches(2), Inches(4), Inches(3), data)
    blank = prs.slides.add_slide(prs.slide_layouts[BLANK])
    text = blank.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame
    text.text = "No title here. "
    run = text.paragraphs[0].add_run()
    run.text, run.font.bold = "Bold", True
    path = tmp_path / "Sensor Review.pptx"
    prs.save(str(path))
    return path


def test_a_deck_becomes_markdown_with_front_matter_headings_and_lists(deck, tmp_path):
    result = import_pptx(deck, tmp_path / "out")
    text = result.markdown.read_text(encoding="utf-8")
    assert result.markdown.name == "sensor_review.md"
    assert text.startswith(
        "---\ntitle: Sensor review\nsubtitle: What we found\nauthor: A. Writer\nversion: '2.0'\n---\n"
    )
    assert "# Findings\n" in text and "## Readings\n" in text
    assert "- Temperature\n  - Within 0.5 C\n- Humidity" in text
    assert "::: notes\nSay why the drift matters.\n:::" in text


def test_two_content_slides_become_columns(deck, tmp_path):
    text = import_pptx(deck, tmp_path / "out").markdown.read_text(encoding="utf-8")
    assert "::::: columns\n\n::: column\n- Polled\n:::\n\n::: column\n- Pushed\n:::\n\n:::::" in text


def test_tables_pictures_and_their_captions(deck, tmp_path):
    result = import_pptx(deck, tmp_path / "out")
    text = result.markdown.read_text(encoding="utf-8")
    assert "| Name | Address |\n| --- | --- |\n| TEMP | 0x01\\|0x02 |" in text
    assert "[[CAPTION:Figure]] Figure caption\n\n![The sensor board](sensor_review/media/image_001.png)" in text
    assert result.images == 1 and (result.media / "image_001.png").is_file()


def test_monospace_text_becomes_a_code_block(deck, tmp_path):
    text = import_pptx(deck, tmp_path / "out").markdown.read_text(encoding="utf-8")
    assert "```\nint main(void)\n{ return 0; }\n```" in text


def test_what_cannot_be_imported_is_counted(deck, tmp_path, caplog):
    result = import_pptx(deck, tmp_path / "out")
    assert "1 charts not imported" in result.warnings
    assert "1 charts not imported" in caplog.text


def test_a_slide_without_a_title_is_a_page_break(deck, tmp_path):
    text = import_pptx(deck, tmp_path / "out").markdown.read_text(encoding="utf-8")
    assert "::: {.pagebreak}\n:::\n\nNo title here. **Bold**" in text


def test_bold_labels_become_callouts(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY])
    slide.shapes.title.text = "Care"
    frame = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2)).text_frame
    frame.text = "WARNING"
    frame.paragraphs[0].runs[0].font.bold = True
    frame.add_paragraph().text = "Do not open the case."
    frame.add_paragraph().text = "​"
    path = tmp_path / "care.pptx"
    prs.save(str(path))
    text = import_pptx(path, tmp_path).markdown.read_text(encoding="utf-8")
    assert "> [!WARNING]\n> Do not open the case." in text and "​" not in text


def test_an_existing_file_is_kept_unless_forced(deck, tmp_path):
    import_pptx(deck, tmp_path)
    with pytest.raises(ConversionError, match="already exists"):
        import_pptx(deck, tmp_path)
    import_pptx(deck, tmp_path, force=True)


def test_a_file_that_is_not_a_deck_is_refused(tmp_path):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(InvalidPptx, match="not a PowerPoint deck"):
        import_pptx(bad)
    with pytest.raises(InvalidPptx, match="not found"):
        import_pptx(tmp_path / "missing.pptx")


def test_the_command_picks_the_importer_by_file_type(deck, tmp_path, capsys):
    assert main(["import", str(deck), "-o", str(tmp_path / "cli"), "--print-output-paths"]) == 0
    assert "sensor_review.md" in capsys.readouterr().out
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"x")
    assert main(["import", str(bad)]) == 2


def test_a_built_deck_comes_back_as_the_same_slides(convert_slides, tmp_path):
    """Markdown -> slides -> import -> slides gives the same slide titles."""
    from pptx_helpers import titles

    source = (
        "---\ntitle: Round trip\n---\n\n# Part one\n\n## First\n\n- a\n  - b\n\n## Second\n\n"
        "| x | y |\n|---|---|\n| 1 | 2 |\n\n# Part two\n\n## Third\n\nSome **bold** text.\n"
    )
    built = convert_slides(source)
    first = titles(built)
    imported = import_pptx(built, tmp_path / "back").markdown
    again = convert_slides(imported.read_text(encoding="utf-8"))
    assert titles(again) == first
    assert Path(imported).read_text(encoding="utf-8").count("## ") == 3
