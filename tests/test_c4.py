"""Structurizr support: model, discovery, CLI wrapper, pages, drawing, commands and the build source (#96 to #100, #116)."""

from __future__ import annotations

import json
import os
import shutil
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import pytest

from py_doc.errors import ConversionError

from py_doc.app import main
from py_doc.c4 import stack as stack_module
from py_doc.c4.check import check_model
from py_doc.c4.generate import generate
from py_doc.c4.model import WorkspaceError, load_workspace, workspace_from_dict
from py_doc.c4.render import RenderError, render_views
from py_doc.c4.stack import Stack
from py_doc.c4.structurizr import StructurizrCli, find_workspaces, includes, workspace_hash, workspace_json
from py_doc.config import ConfigError, load_config
from py_doc.errors import EnvironmentProblem
from py_doc.lint.config import LintConfig
from py_doc.runtime import ToolResult

FIXTURE = Path(__file__).resolve().parent / "fixtures" / "c4" / "workspace.json"
GOLDEN = Path(__file__).parent / "golden" / "c4"
VIEWS = ["Components", "Containers", "Live", "SignIn", "SystemContext"]


@pytest.fixture(autouse=True)
def in_a_temporary_folder(tmp_path, monkeypatch):
    """The .py-doc cache goes under the working folder, so keep it out of the repository."""
    monkeypatch.chdir(tmp_path)


@pytest.fixture
def ws():
    return load_workspace(FIXTURE)


def png_bytes() -> bytes:
    import io

    from PIL import Image

    data = io.BytesIO()
    Image.new("RGB", (16, 16), "navy").save(data, format="PNG")
    return data.getvalue()


@dataclass
class FakeBox:
    """A container runtime that records calls. ``draw`` decides what the picture tool does."""

    calls: list = field(default_factory=list)
    engine_calls: list = field(default_factory=list)
    failed_views: tuple = ()
    exit_code: int = 0
    skip_files: tuple = ()
    running: bool = False
    export_json: str | None = None
    cli_error: Exception | None = None

    def container_path(self, path):
        return Path(path).resolve().as_posix()

    def engine(self, args, **kwargs):
        self.engine_calls.append(list(args))
        if args[0] == "inspect":
            return ToolResult(0, "true\n", "") if self.running else ToolResult(1, "", "No such object")
        if args[0] == "network" and args[1] == "inspect":
            return ToolResult(1, "", "")
        return ToolResult(0, "", "")

    def run(self, argv, **kwargs):
        argv = [str(a) for a in argv]
        self.calls.append(argv)
        if argv[0] == "structurizr-cli":
            if self.cli_error:
                raise self.cli_error
            if argv[1] == "version":
                return ToolResult(0, "structurizr-cli: 2025.11.09\n", "")
            out = Path(argv[argv.index("-output") + 1])
            fmt = argv[argv.index("-format") + 1]
            if fmt == "json":
                (out / "workspace.json").write_text(self.export_json or FIXTURE.read_text(), encoding="utf-8")
            else:
                ext = {"mermaid": "mmd", "plantuml": "puml"}.get(fmt.split("/")[0], "txt")
                for key in VIEWS:
                    (out / f"structurizr-{key}.{ext}").write_text(f"graph TD\n  {key}\n", encoding="utf-8")
            return ToolResult(0, "", "")
        if argv[0] == "node":
            out = Path(argv[argv.index("--output") + 1])
            fmt = argv[3]
            for key in VIEWS:
                if key not in self.skip_files and key not in self.failed_views:
                    (out / f"{key}.{fmt}").write_bytes(png_bytes())
            tail = f"Failed views: {', '.join(self.failed_views)}" if self.failed_views else ""
            return ToolResult(self.exit_code or (2 if self.failed_views else 0), "", tail)
        raise AssertionError(argv)


@pytest.fixture
def box(monkeypatch):
    monkeypatch.setattr(stack_module.Stack, "_wait", lambda self: None)
    return FakeBox()


@pytest.fixture
def project(tmp_path):
    folder = tmp_path / "arch"
    folder.mkdir()
    (folder / "workspace.dsl").write_text('workspace "Big Bank" {\n  !include model.dsl\n}\n', encoding="utf-8")
    (folder / "model.dsl").write_text("model {}\n", encoding="utf-8")
    return folder


# --- the model (#96) -----------------------------------------------------------------------------


def test_the_model_matches_the_workspace(ws):
    assert ws.name == "Big Bank plc"
    assert len(ws.elements) == 11 and len(ws.relationships) == 5
    assert [v.key for v in ws.views] == VIEWS
    api = ws.element("5")
    assert api.kind == "container" and api.technology == "Java and Spring MVC"
    assert api.perspectives[0].name == "Security" and api.properties == {"owner": "Platform team"}
    assert [c.name for c in api.children] == ["Sign In Controller"]
    instance = ws.element("22")
    assert instance.name == "API Application" and instance.element_id == "5"
    assert ws.element("20").environment == "Live"
    assert [s.relationship_id for s in ws.views[3].steps] == ["12", "13"]
    assert ws.decisions[0].title.startswith("1.")


def test_a_missing_field_is_named():
    with pytest.raises(WorkspaceError, match=r"model.softwareSystems\[0\].containers\[1\] has no field 'name'"):
        workspace_from_dict(
            {"model": {"softwareSystems": [{"id": "1", "name": "S", "containers": [{"id": "2", "name": "A"}, {"id": "3"}]}]}}
        )
    with pytest.raises(WorkspaceError, match="no field 'model'"):
        workspace_from_dict({})


def test_unreadable_json_is_an_error(tmp_path):
    (tmp_path / "w.json").write_text("{nope")
    with pytest.raises(WorkspaceError, match="not JSON"):
        load_workspace(tmp_path / "w.json")


# --- discovery and the CLI wrapper ----------------------------------------------------------------


def test_discovery_is_sorted_and_skips_hidden_vendor_and_fragments(project, tmp_path):
    (tmp_path / "node_modules" / "x").mkdir(parents=True)
    (tmp_path / "node_modules" / "x" / "workspace.dsl").write_text("workspace {}\n")
    (tmp_path / ".hidden").mkdir()
    (tmp_path / ".hidden" / "workspace.dsl").write_text("workspace {}\n")
    (tmp_path / "other").mkdir()
    (tmp_path / "other" / "workspace.json").write_text('{"model": {}}')
    (tmp_path / "other2").mkdir()
    (tmp_path / "other2" / "workspace.json").write_text('{"unrelated": 1}')
    (project / "workspace.json").write_text('{"model": {}}')  # the DSL's own export: not a second workspace
    found = [p.relative_to(tmp_path).as_posix() for p in find_workspaces(tmp_path)]
    assert found == ["arch/workspace.dsl", "other/workspace.json"]


def test_the_hash_follows_every_included_file(project):
    first = workspace_hash(project / "workspace.dsl")
    assert [p.name for p in includes(project / "workspace.dsl")] == ["workspace.dsl", "model.dsl"]
    (project / "model.dsl").write_text("model { person p }\n")
    assert workspace_hash(project / "workspace.dsl") != first


def test_json_is_exported_once_and_cached(project, box, tmp_path):
    cache = tmp_path / "cache"
    first = workspace_json(project / "workspace.dsl", cache, lambda: box)
    again = workspace_json(project / "workspace.dsl", cache, lambda: box)
    assert first == again and load_workspace(first).name == "Big Bank plc"
    assert sum(c[0] == "structurizr-cli" and c[1] == "export" for c in box.calls) == 1
    (project / "model.dsl").write_text("model { person p }\n")
    workspace_json(project / "workspace.dsl", cache, lambda: box)
    assert sum(c[1] == "export" for c in box.calls if c[0] == "structurizr-cli") == 2
    assert len(list(cache.glob("workspace-*.json"))) == 1
    assert workspace_json(FIXTURE, cache, lambda: box) == FIXTURE  # a JSON file needs no Java


def test_a_cli_failure_shows_its_output_and_writes_nothing(project, box, tmp_path):
    from py_doc.errors import ToolFailed

    box.cli_error = ToolFailed("structurizr-cli", 1, "line 3: unexpected token")
    with pytest.raises(ToolFailed, match="unexpected token"):
        workspace_json(project / "workspace.dsl", tmp_path / "cache", lambda: box)
    assert not (tmp_path / "cache").exists() or not list((tmp_path / "cache").glob("*.json"))


def test_a_missing_host_cli_says_what_to_install(project, tmp_path):
    with pytest.raises(EnvironmentProblem, match="needs Java"):
        StructurizrCli(None, "/nonexistent/structurizr-cli").export(project / "workspace.dsl", "json", tmp_path / "o")


# --- the pages (#97) -------------------------------------------------------------------------------


def pages(ws, **kwargs):
    images = {v.key: f"images/{v.key}.png" for v in ws.views}
    return generate(ws, "workspace.json", images=images, **kwargs)


def test_pages_match_the_golden_files(ws):
    made = pages(ws)
    if os.environ.get("UPDATE_GOLDEN") == "1":
        GOLDEN.mkdir(parents=True, exist_ok=True)
        for name, text in made.items():
            (GOLDEN / name).write_text(text, encoding="utf-8")
    assert sorted(made) == sorted(p.name for p in GOLDEN.glob("*.md"))
    for name, text in made.items():
        assert text == (GOLDEN / name).read_text(encoding="utf-8"), name


def test_output_is_stable_and_has_no_times(ws):
    assert pages(ws) == pages(load_workspace(FIXTURE))
    assert all("20" not in line or "2025" not in line for text in pages(ws).values() for line in text.splitlines())


def test_levels_and_systems_limit_the_pages(ws):
    assert sorted(pages(ws, levels=["context", "relationships"])) == ["01_context.md", "04_relationships.md"]
    assert "Internet Banking System" not in pages(ws, systems=["Mainframe Banking System"])["01_context.md"].split("## Software systems")[1].split("|")[3]
    with pytest.raises(ValueError, match="Unknown level"):
        pages(ws, levels=["nope"])


def test_templates_can_be_replaced(ws):
    assert "Landscape of Big Bank plc." in pages(ws, templates={"context": "Landscape of {workspace}."})["01_context.md"]


def test_generated_pages_pass_the_linter(ws, tmp_path, capsys):
    for name, text in pages(ws).items():
        (tmp_path / name).write_text(text, encoding="utf-8")
    assert main(["lint", str(tmp_path), "--engine", "builtin"]) == 0


# --- drawing the views (#98, #116) ----------------------------------------------------------------


def draw(ws, box, tmp_path, **kwargs):
    return render_views(ws, FIXTURE, tmp_path / "out", box, **kwargs)


def test_every_view_is_drawn_as_png_and_svg(ws, box, tmp_path):
    files = draw(ws, box, tmp_path)
    assert sorted(p.name for p in files) == sorted(f"{k}.{f}" for k in VIEWS for f in ("png", "svg"))
    assert any(c[0] == "run" for c in box.engine_calls) and box.engine_calls[-2][0] == "rm"  # started and stopped Lite
    run = [c for c in box.engine_calls if c[0] == "run"][0]
    assert run[-2:] == ["structurizr/structurizr", "local"]
    assert "STRUCTURIZR_AUTOSAVEINTERVAL=0" in run and "STRUCTURIZR_EDITABLE=false" in run


def test_a_view_that_fails_names_the_view_and_fails(ws, box, tmp_path):
    box.failed_views = ("Containers",)
    with pytest.raises(RenderError, match=r"Containers\.png: the browser could not draw it"):
        draw(ws, box, tmp_path)
    assert box.engine_calls[-2][0] == "rm"  # Lite is stopped even after a failure


def test_a_missing_picture_is_an_error_not_a_placeholder(ws, box, tmp_path):
    box.skip_files = ("Live",)
    with pytest.raises(RenderError, match=r"Live\.svg: no picture was written"):
        draw(ws, box, tmp_path)
    assert not (tmp_path / "out" / "Live.png").exists()


def test_an_empty_view_is_an_error(box, tmp_path):
    data = json.loads(FIXTURE.read_text())
    data["views"]["containerViews"][0]["elements"] = []
    with pytest.raises(RenderError, match="empty.*Containers"):
        draw(workspace_from_dict(data), box, tmp_path)
    assert box.calls == []


def test_a_fatal_browser_error_fails_with_its_output(ws, box, tmp_path):
    box.exit_code = 1
    with pytest.raises(RenderError, match="exit status 1"):
        draw(ws, box, tmp_path)


def test_unknown_view_and_format(ws, box, tmp_path):
    with pytest.raises(RenderError, match="no view 'Nope'"):
        draw(ws, box, tmp_path, keys=["Nope"])
    with pytest.raises(RenderError, match="Cannot draw 'gif'"):
        draw(ws, box, tmp_path, formats=("gif",))


def test_a_missing_engine_is_an_environment_problem(ws, tmp_path):
    class Down(FakeBox):
        def engine(self, args, **kwargs):
            raise EnvironmentProblem("The Docker engine is not running.")

    with pytest.raises(EnvironmentProblem):
        draw(ws, Down(), tmp_path)


def test_stale_pictures_are_removed_first(ws, box, tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    (out / "Live.png").write_bytes(b"old")
    box.skip_files = ("Live",)
    with pytest.raises(RenderError):
        draw(ws, box, tmp_path)
    assert not (out / "Live.png").exists()


def test_a_running_server_is_used_and_left_running(ws, box, tmp_path):
    box.running = True
    draw(ws, box, tmp_path)
    assert not any(c[0] in ("run", "rm") for c in box.engine_calls)


def test_local_gets_the_workspace_folder(project, box):
    (project / "workspace.dsl").write_text("workspace {}\n")
    stack = Stack(box, project / "workspace.dsl")
    stack.up()
    run = [c for c in box.engine_calls if c[0] == "run"][0]
    assert f"type=bind,source={project.resolve().as_posix()},target=/usr/local/structurizr" in run


def test_a_workspace_with_another_name_is_staged_as_workspace_dsl(project, box):
    (project / "model2.dsl").write_text("workspace { !include more.dsl }\n")
    (project / "more.dsl").write_text("// included\n")
    (project / "workspace.dsl").write_text("workspace { /* another model */ }\n")
    stack = Stack(box, project / "model2.dsl")
    stack.up()
    run = [c for c in box.engine_calls if c[0] == "run"][0]
    source = next(a for a in run if a.startswith("type=bind")).split("source=")[1].split(",target")[0]
    staged = Path(source)
    assert staged != project.resolve() and (staged / "more.dsl").is_file()
    assert (staged / "workspace.dsl").read_text().startswith("workspace { !include more.dsl")
    stack.down()
    assert not staged.exists()


# --- the commands -----------------------------------------------------------------------------------


@pytest.fixture
def patched(monkeypatch, box):
    monkeypatch.setattr("py_doc.c4.cli.connect", lambda *a, **k: box)
    return box


def test_check_reports_the_model_and_exits_by_severity(patched, tmp_path, capsys):
    assert main(["c4", "check", str(FIXTURE), "--strict"]) == 0  # a complete model is clean
    assert "11 element(s), 5 relationship(s), 5 view(s), 0 error(s), 0 warning(s)" in capsys.readouterr().out
    data = json.loads(FIXTURE.read_text())
    data["model"]["people"][0].pop("description")
    (tmp_path / "w.json").write_text(json.dumps(data))
    assert main(["c4", "check", str(tmp_path / "w.json")]) == 0
    assert "C4001" in capsys.readouterr().out
    assert main(["c4", "check", str(tmp_path / "w.json"), "--strict"]) == 1
    assert main(["c4", "check", str(tmp_path / "nowhere")]) == 1


def test_model_rules(ws):
    data = json.loads(FIXTURE.read_text())
    data["model"]["softwareSystems"][1]["containers"][2].pop("technology")
    data["model"]["softwareSystems"][1]["containers"][2].pop("description")
    data["model"]["softwareSystems"][1]["containers"][0]["relationships"][0].pop("description")
    issues = check_model(workspace_from_dict(data), Path("w.json"))
    assert {(i.rule, i.severity) for i in issues} >= {("C4001", "warning"), ("C4002", "warning"), ("C4003", "warning")}
    off = LintConfig()
    off.disabled_rules = frozenset({"C4001", "C4002", "C4003"})
    assert check_model(workspace_from_dict(data), Path("w.json"), off) == []


def test_lint_rules_can_name_the_model_rules(tmp_path):
    (tmp_path / "py-doc.yml").write_text("lint:\n  rules:\n    C4001: off\n    C4004: error\n", encoding="utf-8")
    config = load_config(tmp_path / "py-doc.yml")
    assert "C4001" in config.lint.disabled_rules and config.lint.severities["C4004"] == "error"


def test_export_dry_run_writes_nothing(patched, tmp_path, capsys):
    out = tmp_path / "out"
    assert main(["c4", "export", str(FIXTURE), "-o", str(out), "-f", "png", "markdown", "--dry-run"]) == 0
    listed = capsys.readouterr().out
    assert "Containers.png" in listed and "01_context.md" in listed
    assert not out.exists() and patched.calls == [] and patched.engine_calls == []


def test_export_writes_pictures_and_markdown(patched, tmp_path, capsys):
    out = tmp_path / "out"
    assert main(["c4", "export", str(FIXTURE), "-o", str(out), "-f", "png", "markdown"]) == 0
    assert (out / "SystemContext.png").read_bytes().startswith(b"\x89PNG")
    context = (out / "markdown" / "01_context.md").read_text()
    assert "![System Context for Internet Banking](../SystemContext.png)" in context


def test_export_failure_is_nonzero_and_names_the_view(patched, tmp_path, capsys):
    patched.failed_views = ("Live",)
    assert main(["c4", "export", str(FIXTURE), "-o", str(tmp_path / "o"), "-f", "svg"]) == 1
    err = capsys.readouterr().err
    assert "Live.svg" in err and "1 of 1 workspace(s) failed" in err


def test_text_formats_come_from_the_cli_and_never_from_a_placeholder(patched, project, tmp_path, capsys):
    out = tmp_path / "o"
    assert main(["c4", "export", str(project / "workspace.dsl"), "-o", str(out), "-f", "mermaid", "c4plantuml"]) == 0
    assert (out / "mermaid" / "structurizr-Containers.mmd").read_text().startswith("graph TD")
    assert any("plantuml/c4plantuml" in c for c in patched.calls)
    assert main(["c4", "export", str(FIXTURE), "-o", str(out), "-f", "mermaid"]) == 1  # JSON has no DSL to export
    assert "needs the .dsl file" in capsys.readouterr().err


def test_a_missing_engine_exits_2(project, tmp_path, monkeypatch, capsys):
    def down(*args, **kwargs):
        raise EnvironmentProblem("The Docker engine is not running.")

    monkeypatch.setattr("py_doc.c4.cli.connect", down)
    assert main(["c4", "export", str(FIXTURE), "-o", str(tmp_path / "o"), "-f", "png"]) == 2


def test_up_status_and_down(patched, project, capsys):
    assert main(["c4", "status", str(project / "workspace.dsl")]) == 0
    assert capsys.readouterr().out.strip() == "absent"
    assert main(["c4", "up", str(project / "workspace.dsl")]) == 0
    assert "started" in capsys.readouterr().out
    assert main(["c4", "down", str(project / "workspace.dsl")]) == 0
    assert patched.engine_calls[-2][:2] == ["rm", "-f"]


def test_dry_run_creates_no_output_folder_even_for_a_missing_input(patched, tmp_path):
    out = tmp_path / "o"
    assert main(["c4", "export", str(tmp_path / "gone"), "-o", str(out), "--dry-run"]) == 1
    assert not out.exists()


# --- the structurizr source of py-doc.yml (#99) ---------------------------------------------------


def config_with(tmp_path, body: str):
    (tmp_path / "py-doc.yml").write_text(body, encoding="utf-8")
    return load_config(tmp_path / "py-doc.yml")


def test_the_source_is_read_with_its_options(tmp_path):
    shutil.copy(FIXTURE, tmp_path / "workspace.json")
    config = config_with(
        tmp_path,
        "documents:\n  arch:\n    sources:\n      - file: intro.md\n"
        "      - structurizr: workspace.json\n        levels: [context, containers]\n        systems: [Internet Banking System]\n"
        "        heading_offset: 1\n        diagrams: {engine: source, format: plantuml}\n"
        "c4:\n  cli: /opt/structurizr-cli\n  port: 8081\n",
    )
    source = config.documents["arch"].sources[1]
    assert source.structurizr.levels == ["context", "containers"] and source.heading_offset == 1
    assert (source.structurizr.engine, source.structurizr.format) == ("source", "plantuml")
    assert config.c4.cli == "/opt/structurizr-cli" and config.c4.port == 8081
    assert config.lint.documents["arch"] == [tmp_path / "intro.md"]


@pytest.mark.parametrize(
    ("body", "message"),
    [
        ("      - structurizr: w.json\n        levels: [cake]\n", r"documents.arch.sources\[0\].levels has unknown level 'cake'"),
        ("      - structurizr: w.json\n        diagrams: {engine: magic}\n", "diagrams.engine must be browser or source"),
        ("      - structurizr: w.json\n        diagrams: {format: gif}\n", "diagrams.format must be one of png, svg"),
        ("      - structurizr: w.json\n        page_brek: true\n", "page_brek"),
    ],
)
def test_source_errors_name_their_place(tmp_path, body, message):
    with pytest.raises(ConfigError, match=message):
        config_with(tmp_path, "documents:\n  arch:\n    sources:\n" + body)


def test_expand_makes_pages_and_reuses_them_when_nothing_changed(tmp_path, box):
    from py_doc.c4.source import expand

    shutil.copy(FIXTURE, tmp_path / "workspace.json")
    config = config_with(
        tmp_path,
        "documents:\n  arch:\n    sources:\n      - file: intro.md\n      - structurizr: workspace.json\n        heading_offset: 1\n"
        "        levels: [context, relationships]\n",
    )
    (tmp_path / "intro.md").write_text("# Intro\n", encoding="utf-8")
    sources = config.documents["arch"].sources
    done = expand(sources, "arch", tmp_path, lambda: box, config.c4)
    assert [s.file.name for s in done] == ["intro.md", "01_context.md", "04_relationships.md"]
    assert all(s.heading_offset == 1 for s in done[1:]) and all(s.structurizr is None for s in done)
    assert (tmp_path / ".py-doc/c4/arch/01_workspace/images/SystemContext.png").exists()
    runs = sum(c[0] == "node" for c in box.calls)
    expand(sources, "arch", tmp_path, lambda: box, config.c4)
    assert sum(c[0] == "node" for c in box.calls) == runs  # unchanged: nothing is drawn again
    config.documents["arch"].sources[1].structurizr.levels = ["context"]
    expand(sources, "arch", tmp_path, lambda: box, config.c4)
    assert sum(c[0] == "node" for c in box.calls) > runs


def test_the_source_engine_embeds_labelled_diagram_source(tmp_path, project, box):
    from py_doc.c4.source import expand

    config = config_with(
        tmp_path,
        f"documents:\n  arch:\n    sources:\n      - structurizr: {project / 'workspace.dsl'}\n        levels: [context]\n"
        "        diagrams: {engine: source, format: mermaid}\n",
    )
    (page,) = expand(config.documents["arch"].sources, "arch", tmp_path, lambda: box, config.c4)
    text = page.file.read_text(encoding="utf-8")
    assert "Diagram source (mermaid) of the view 'System Context for Internet Banking'" in text and "```mermaid\ngraph TD" in text
    assert not any(c[0] == "node" for c in box.calls)


def test_a_broken_view_stops_the_build_source(tmp_path, box):
    from py_doc.c4.source import expand

    shutil.copy(FIXTURE, tmp_path / "workspace.json")
    config = config_with(tmp_path, "documents:\n  arch:\n    sources:\n      - structurizr: workspace.json\n")
    box.failed_views = ("SignIn",)
    with pytest.raises(RenderError, match=r"SignIn\.png"):
        expand(config.documents["arch"].sources, "arch", tmp_path, lambda: box, config.c4)


# --- the old py-c4-exporter names (#100) -------------------------------------------------------------


def test_the_old_name_delegates_and_says_where_it_went(patched, tmp_path, capsys):
    from pyc4exporter.cli import main as old

    code = old([str(FIXTURE), "-o", str(tmp_path / "o"), "-f", "png"])
    assert code == 0
    assert "py-c4-exporter is now part of py-doc. `py-doc c4 export` does the same" in capsys.readouterr().err


def test_a_build_joins_generated_chapters_pictures_and_hand_written_text(tmp_path, box, pandoc_runtime):
    """py-doc build: hand-written chapter, generated chapters with pictures, one Word file."""
    from docx import Document

    from py_doc.build import build_document

    class Both:
        """Pandoc runs for real on the host. Structurizr and the browser are the fakes above."""

        def run(self, argv, **kwargs):
            if str(argv[0]) in ("node", "structurizr-cli"):
                return box.run(argv, **kwargs)
            return pandoc_runtime.run(argv, **kwargs)

        def __getattr__(self, name):
            return getattr(box if name in ("engine",) else pandoc_runtime, name)

    shutil.copy(FIXTURE, tmp_path / "workspace.json")
    (tmp_path / "intro.md").write_text("# Introduction\n\nWhy this exists.\n", encoding="utf-8")
    config = config_with(
        tmp_path,
        "documents:\n  architecture:\n    title: Architecture\n    sources:\n      - file: intro.md\n"
        "      - structurizr: workspace.json\n        levels: [context, relationships]\n        heading_offset: 0\n"
        "    outputs:\n      docx: {}\n",
    )
    (docx,) = build_document(config, config.documents["architecture"], runtime=lambda: Both(), report=None)
    document = Document(docx)
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Why this exists." in text and "Software systems" in text and "Figure" in text
    with zipfile.ZipFile(docx) as archive:
        assert len([n for n in archive.namelist() if n.startswith("word/media/")]) == 1  # the only view the levels show
    assert len(document.tables) == 3


# --- the lock file, the cache root and the model rules in py-doc lint ---------------------------------


DIGEST = "structurizr/structurizr@sha256:" + "a" * 64


def test_lock_pins_lite_and_builds_use_it(tmp_path, box):
    from py_doc.c4.lock import LOCK, lock_image, read_lock, resolve_image

    box.engine = lambda args, **kw: ToolResult(0, DIGEST + "\n" if args[0] == "image" else "", "")
    assert lock_image(box, tmp_path) == DIGEST
    assert read_lock(tmp_path) == {"structurizr/structurizr": DIGEST}
    assert resolve_image(None, tmp_path) == DIGEST
    assert resolve_image("structurizr/structurizr:2025", tmp_path) == "structurizr/structurizr:2025"  # py-doc.yml wins
    (tmp_path / LOCK).write_text("structurizr/structurizr nonsense\n")
    with pytest.raises(Exception, match=r"c4-images.lock:1: expected"):
        read_lock(tmp_path)


def test_an_unpinned_image_warns(tmp_path, caplog):
    from py_doc.c4.lock import resolve_image

    assert resolve_image(None, tmp_path) == "structurizr/structurizr"
    assert "py-doc c4 lock" in caplog.text


def test_the_lock_command(patched, tmp_path, capsys):
    patched.engine = lambda args, **kw: ToolResult(0, DIGEST + "\n" if args[0] == "image" else "", "")
    assert main(["c4", "lock"]) == 0
    assert DIGEST in capsys.readouterr().out and (tmp_path / "c4-images.lock").exists()


def test_the_cache_goes_to_the_project_folder_not_the_working_folder(patched, project, tmp_path, monkeypatch):
    (tmp_path / "py-doc.yml").write_text("documents: {}\n", encoding="utf-8")
    sub = tmp_path / "sub"
    sub.mkdir()
    monkeypatch.chdir(sub)
    assert main(["c4", "check", str(project / "workspace.dsl")]) == 0
    assert list((tmp_path / ".py-doc" / "c4" / "json").glob("workspace-*.json")) and not (sub / ".py-doc").exists()


def test_lint_runs_the_model_rules_for_the_documents_workspaces(patched, tmp_path, capsys):
    data = json.loads(FIXTURE.read_text())
    data["model"]["people"][0].pop("description")
    (tmp_path / "workspace.json").write_text(json.dumps(data))
    (tmp_path / "a.md").write_text("# A\n\nText.\n", encoding="utf-8")
    (tmp_path / "py-doc.yml").write_text(
        "documents:\n  d:\n    sources:\n      - file: a.md\n      - structurizr: workspace.json\n"
        "lint:\n  rules:\n    C4004: off\n",
        encoding="utf-8",
    )
    assert main(["lint", str(tmp_path), "--engine", "builtin"]) == 0
    out = capsys.readouterr().out
    assert "C4001" in out and "workspace.json:1:" in out and "C4004" not in out
    assert main(["lint", str(tmp_path), "--engine", "builtin", "--strict"]) == 1


# --- waiting for Structurizr Local --------------------------------------------------------------------


class _Waiting:
    """A runtime whose probe fails a set number of times, and whose container may stop."""

    def __init__(self, failures=0, running=True, logs="x"):
        self.failures, self.running, self.logs, self.calls = failures, running, logs, []

    def engine(self, args, **kwargs):
        self.calls.append(args[0])
        if args[0] == "inspect":
            return ToolResult(0, "true\n" if self.running else "false\n", "")
        if args[0] == "logs":
            return ToolResult(0, self.logs, "")
        return ToolResult(0, "", "")

    def run(self, argv, **kwargs):
        assert argv[3] == "http://localhost:8080/workspace/diagrams"
        assert argv[:2] == ["node", "-e"] and kwargs["network"] == "container:py-doc-c4-local"  # probed as localhost, the only name Local answers to
        self.calls.append("probe")
        self.failures -= 1
        return ToolResult(1 if self.failures >= 0 else 0, "", "")


@pytest.fixture
def quick(monkeypatch):
    from py_doc.c4 import stack as stack_module

    monkeypatch.setattr(stack_module.time, "sleep", lambda s: None)
    clock = iter(range(0, 10_000, 40))  # each look at the clock is 40 seconds later
    monkeypatch.setattr(stack_module.time, "monotonic", lambda: next(clock))
    return stack_module


def test_the_wait_ends_when_the_probe_answers(quick, tmp_path, caplog):
    runtime = _Waiting(failures=2)
    with caplog.at_level("INFO"):
        Stack(runtime, tmp_path / "workspace.json")._wait()
    assert runtime.calls.count("probe") == 3
    assert "ready after" in caplog.text and "Still waiting" in caplog.text


def test_the_wait_fails_at_once_when_the_container_stops(quick, tmp_path):
    runtime = _Waiting(running=False, logs="\x1b[31mERROR\x1b[0m port in use")
    with pytest.raises(ConversionError, match="stopped while starting.\nERROR port in use"):
        Stack(runtime, tmp_path / "workspace.json")._wait()
    assert "probe" not in runtime.calls and runtime.calls[-2:] == ["rm", "network"]


def test_the_wait_gives_up_after_the_limit_with_clean_logs(quick, tmp_path):
    runtime = _Waiting(failures=10_000, logs="\x1b[34mINFO\x1b[0;39m Themes:")
    with pytest.raises(ConversionError, match=r"did not answer within 300 seconds.\nINFO Themes:"):
        Stack(runtime, tmp_path / "workspace.json")._wait()
