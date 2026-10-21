"""Document control in Word output: placeholders, marker, status templates, file names, versions, properties."""

from __future__ import annotations

import logging
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fake_engine import install

from py_doc.build import build_document
from py_doc.config import ConfigError, load_config
from py_doc.errors import ConversionError
from py_doc.naming import check_name_pattern, expand_name, git_version, safe_version, status_template
from py_doc.runtime import Runtime

needs_pandoc = pytest.mark.skipif(shutil.which("pandoc") is None, reason="pandoc is not installed")
needs_git = pytest.mark.skipif(shutil.which("git") is None, reason="git is not installed")


# --- naming ---------------------------------------------------------------------------------------------


class TestNaming:
    def test_version_safe(self):
        assert safe_version("1.2.3") == "1_2_3" and safe_version(None) == "x_x_x" and safe_version(" ") == "x_x_x"
        assert safe_version("v1.0-rc1") == "v1_0-rc1"

    def test_pattern_expansion(self):
        assert expand_name("{name}-{version_safe}", "guide", "1.2") == "guide-1_2"
        assert expand_name("{name}_v{version}", "guide", "1.2") == "guide_v1.2"
        assert expand_name("{name}-{version_safe}", "guide", None) == "guide-x_x_x"

    def test_a_pattern_needing_a_version_without_one_is_an_error(self):
        with pytest.raises(ConversionError, match=r"uses \{version\}, but the document has no version"):
            expand_name("{name}-{version}", "guide", None)

    @pytest.mark.parametrize(
        ("pattern", "message"), [("{nmae}", "unknown field {nmae}"), ("a{b", "stray|brace"), ("a}", "brace")]
    )
    def test_bad_patterns(self, pattern, message):
        with pytest.raises(ValueError, match=message):
            check_name_pattern(pattern, "x")

    def test_status_template_rule(self, tmp_path, caplog):
        base = tmp_path / "T.docx"
        base.write_bytes(b"")
        (tmp_path / "T_Draft.docx").write_bytes(b"")
        (tmp_path / "T_approved.docx").write_bytes(b"")
        with caplog.at_level(logging.INFO):
            assert status_template(base, "Draft") == tmp_path / "T_Draft.docx"
            assert status_template(base, "Approved") == tmp_path / "T_approved.docx"
            assert "Using the template for status 'Draft'" in caplog.text
            caplog.clear()
            assert status_template(base, "Obsolete") == base
            assert "No template for status 'Obsolete'" in caplog.text and "so " in caplog.text
        assert status_template(base, None) == base and status_template(None, "Draft") is None


@needs_git
class TestGitVersion:
    def git(self, folder, *args):
        subprocess.run(
            ["git", "-C", str(folder), "-c", "user.name=t", "-c", "user.email=t@t", *args],
            check=True,
            capture_output=True,
        )

    def test_describes_the_repository(self, tmp_path):
        self.git(tmp_path, "init", "-q")
        (tmp_path / "a").write_text("x")
        self.git(tmp_path, "add", "a")
        self.git(tmp_path, "commit", "-q", "-m", "one")
        self.git(tmp_path, "tag", "v1.4")
        assert git_version(tmp_path) == "v1.4"
        (tmp_path / "a").write_text("y")
        self.git(tmp_path, "commit", "-qam", "two")
        assert git_version(tmp_path).startswith("v1.4-1-g")

    def test_outside_a_repository_is_an_error_not_a_made_up_version(self, tmp_path, monkeypatch):
        monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(tmp_path.parent))
        with pytest.raises(ConversionError, match="needs a git repository"):
            git_version(tmp_path)

    def test_a_repository_without_a_commit_is_an_error(self, tmp_path):
        self.git(tmp_path, "init", "-q")
        with pytest.raises(ConversionError, match="needs a git repository with a commit"):
            git_version(tmp_path)


# --- configuration -----------------------------------------------------------------------------------


def write_config(folder: Path, text: str):
    (folder / "py-doc.yml").write_text(text, encoding="utf-8")
    return load_config(folder / "py-doc.yml")


class TestConfig:
    def test_the_default_is_the_exact_path(self, tmp_path):
        config = write_config(tmp_path, "documents:\n  g: {sources: [a.md], outputs: {docx: {path: out/x.docx}}}\n")
        assert config.output_path(config.documents["g"], "docx", "1.0") == tmp_path / "out" / "x.docx"

    def test_a_name_pattern_applies_to_the_default_and_to_a_given_path(self, tmp_path):
        config = write_config(
            tmp_path,
            "outputs:\n  docx: {name: '{name}-{version_safe}'}\ndocuments:\n  g: {sources: [a.md]}\n"
            "  h: {sources: [a.md], outputs: {docx: {path: out/x.docx}}}\n",
        )
        assert config.output_path(config.documents["g"], "docx", "1.2") == tmp_path / "build" / "g-1_2.docx"
        assert config.output_path(config.documents["h"], "docx", "1.2") == tmp_path / "out" / "x-1_2.docx"

    def test_a_bad_pattern_is_a_config_error_naming_the_place(self, tmp_path):
        with pytest.raises(ConfigError, match=r"outputs\.docx\.name uses unknown field \{nmae\}"):
            write_config(tmp_path, "outputs:\n  docx: {name: '{nmae}'}\n")


# --- the finished Word file --------------------------------------------------------------------------


@pytest.fixture
def host(tmp_path, monkeypatch):
    engine = install(tmp_path, monkeypatch, mode="host")
    return lambda: Runtime("tools:test", docker=engine.docker, platform_name="linux")


def make_template(path: Path, body_text: str = "Cover") -> Path:
    document = Document()
    document.add_paragraph(f"{body_text}: Replace-with-Title / Replace-with-Author")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Number"
    table.rows[0].cells[1].text = "Replace-with-Document-Number"
    document.add_paragraph("[[INSERTCONTENT]]")
    document.add_paragraph("Back page")
    section = document.sections[0]
    section.header.paragraphs[0].text = "Replace-with-Classification | Replace-with-Title"
    section.footer.paragraphs[0].text = "v Replace-with-Version on Replace-with-Date"
    document.save(path)
    return path


def build(tmp_path, host, front: str, outputs: str = "template: tpl.docx", extra_files=None):
    (tmp_path / "g.md").write_text(f"---\n{front}\n---\n# Body\n\nText.\n", encoding="utf-8")
    config = write_config(tmp_path, f"documents:\n  g:\n    sources: [g.md]\n    outputs:\n      docx: {{{outputs}}}\n")
    return build_document(config, config.documents["g"], runtime=host)


def all_text(document) -> str:
    parts = [p.text for p in document.paragraphs]
    parts += [c.text for t in document.tables for r in t.rows for c in r.cells]
    for section in document.sections:
        parts += [p.text for p in section.header.paragraphs] + [p.text for p in section.footer.paragraphs]
    return "\n".join(parts)


FRONT = "title: Guide\nauthor: Ann\nversion: '1.2'\ndate: 2026-10-12\ndocument_number: GD-1\ndescription: About it\nclassification: Internal"


@needs_pandoc
class TestTemplate:
    def test_every_placeholder_is_filled_in_body_table_header_and_footer(self, tmp_path, host, caplog):
        make_template(tmp_path / "tpl.docx")
        with caplog.at_level(logging.WARNING):
            (out,) = build(tmp_path, host, FRONT)
        document = Document(str(out))
        text = all_text(document)
        assert "Replace-with" not in text and "[[INSERTCONTENT]]" not in text
        assert "Cover: Guide / Ann" in text and "GD-1" in text
        assert "Internal | Guide" in text and "v 1.2 on 2026-10-12" in text
        assert not [r for r in caplog.records if r.levelno >= logging.WARNING and "Table style" not in r.getMessage()]
        texts = [p.text for p in document.paragraphs]
        assert texts.index("Body") < texts.index("Back page")

    def test_a_placeholder_without_a_value_or_a_name_is_reported(self, tmp_path, host, caplog):
        template = make_template(tmp_path / "tpl.docx")
        document = Document(str(template))
        document.add_paragraph("Replace-with-Mystery")
        document.save(template)
        with caplog.at_level(logging.WARNING):
            (out,) = build(tmp_path, host, "title: Guide")
        assert "The template has Replace-with-Author, but the document has no author" in caplog.text
        assert "still contains Replace-with-Mystery" in caplog.text
        assert "Replace-with-Author" not in all_text(Document(str(out)))

    def test_the_prefix_and_marker_can_be_changed(self, tmp_path, host):
        document = Document()
        document.add_paragraph("<<Title>> by <<Author>>")
        document.add_paragraph("{{BODY}}")
        document.add_paragraph("end")
        document.save(tmp_path / "tpl.docx")
        (out,) = build(tmp_path, host, FRONT, "template: tpl.docx, placeholder_prefix: '<<', marker: '{{BODY}}'")
        # the prefix is the start of the name only, so the closing brackets stay
        assert "{{BODY}}" not in all_text(Document(str(out)))

    def test_without_a_marker_the_body_goes_at_the_end(self, tmp_path, host):
        document = Document()
        document.add_paragraph("Cover Replace-with-Title")
        document.save(tmp_path / "tpl.docx")
        (out,) = build(tmp_path, host, FRONT)
        texts = [p.text for p in Document(str(out)).paragraphs]
        assert texts[0] == "Cover Guide" and texts.index("Body") > 0

    def test_status_chooses_the_template_and_falls_back(self, tmp_path, host, caplog):
        make_template(tmp_path / "tpl.docx", "Base")
        make_template(tmp_path / "tpl_Draft.docx", "Draft")
        with caplog.at_level(logging.INFO):
            (out,) = build(tmp_path, host, FRONT + "\nstatus: Draft")
            assert all_text(Document(str(out))).startswith("Draft:")
            (out,) = build(tmp_path, host, FRONT + "\nstatus: Final")
        assert all_text(Document(str(out))).startswith("Base:")
        assert "No template for status 'Final'" in caplog.text

    def test_the_version_is_in_the_file_name_when_asked(self, tmp_path, host):
        make_template(tmp_path / "tpl.docx")
        (out,) = build(tmp_path, host, FRONT, "template: tpl.docx, name: '{name}-{version_safe}'")
        assert out == tmp_path / "build" / "g-1_2.docx"

    @needs_git
    def test_version_git_uses_the_repository(self, tmp_path, host):
        subprocess.run(["git", "-C", str(tmp_path), "init", "-q"], check=True)
        (tmp_path / "x").write_text("x")
        subprocess.run(["git", "-C", str(tmp_path), "add", "x"], check=True)
        subprocess.run(
            ["git", "-C", str(tmp_path), "-c", "user.name=t", "-c", "user.email=t@t", "commit", "-qm", "c"], check=True
        )
        subprocess.run(["git", "-C", str(tmp_path), "tag", "v7"], check=True)
        make_template(tmp_path / "tpl.docx")
        (out,) = build(
            tmp_path, host, FRONT.replace("'1.2'", "git"), "template: tpl.docx, name: '{name}-{version_safe}'"
        )
        assert out.name == "g-v7.docx" and "v v7 on" in all_text(Document(str(out)))

    def test_core_properties_are_written_in_the_right_namespaces(self, tmp_path, host):
        make_template(tmp_path / "tpl.docx")
        (out,) = build(tmp_path, host, FRONT + "\nkeywords: [a, b]\ncategory: Guides\nsubject: Things")
        with zipfile.ZipFile(out) as archive:
            core = archive.read("docProps/core.xml").decode("utf-8")
        assert "<dc:title>Guide</dc:title>" in core and "<dc:creator>Ann</dc:creator>" in core
        assert "<cp:keywords>a, b</cp:keywords>" in core and "<cp:category>Guides</cp:category>" in core
        assert "<dc:subject>Things</dc:subject>" in core
        assert "Doc GD-1 | Version 1.2 | About it" in core
        assert 'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"' in core
        assert core.count("<cp:keywords>") == 1
        # python-docx reads "cp:" properties wrongly once docxcompose is imported (see docx_post), so the
        # file is checked as XML above and by title and author here
        reloaded = Document(str(out)).core_properties
        assert (reloaded.title, reloaded.author) == ("Guide", "Ann")

    def test_a_value_longer_than_word_allows_is_cut_with_a_warning(self, tmp_path, host, caplog):
        make_template(tmp_path / "tpl.docx")
        with caplog.at_level(logging.WARNING):
            (out,) = build(tmp_path, host, FRONT + "\ncategory: " + "c" * 300)
        assert "The document property 'category' is longer than Word allows (255 characters)" in caplog.text
        with zipfile.ZipFile(out) as archive:
            core = archive.read("docProps/core.xml").decode("utf-8")
        assert "<cp:category>" + "c" * 254 + "…</cp:category>" in core
