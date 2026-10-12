"""Finish a PowerPoint file written by pandoc: embedded fonts, version token, footer, long-slide warnings."""

from __future__ import annotations

import copy
import logging
import re
import zipfile
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from py_doc.directives import Fence
from py_doc.errors import ConversionError
from py_doc.markdown import Metadata

logger = logging.getLogger(__name__)

LINES_PER_SLIDE = 16  # beyond this a slide is probably too full to read
_NOTES = re.compile(r"^<!--\s*notes?:\s*(.*?)\s*-->\s*$", re.IGNORECASE | re.DOTALL)
_COMMENT_START = re.compile(r"^<!--\s*notes?:", re.IGNORECASE)


def notes_to_divs(markdown: str) -> str:
    """Turn ``<!-- notes: text -->`` comments into pandoc speaker notes, leaving code fences alone."""
    out: list[str] = []
    fence = Fence()
    pending: list[str] | None = None
    for line in markdown.split("\n"):
        if pending is not None:
            pending.append(line)
            if "-->" in line:
                out += _note("\n".join(pending))
                pending = None
            continue
        if not fence.inside(line) and _COMMENT_START.match(line):
            if "-->" in line:
                out += _note(line)
            else:
                pending = [line]
            continue
        out.append(line)
    if pending is not None:
        out += pending
    return "\n".join(out)


def _note(comment: str) -> list[str]:
    match = _NOTES.match(comment.strip())
    text = match[1] if match else comment
    return ["", "::: notes", text, ":::", ""]


def ensure_fntdata_content_type(path: Path) -> None:
    """Register .fntdata in ``[Content_Types].xml`` when a template embeds fonts.

    Pandoc copies the embedded fonts of a reference document but not always their content type, and
    python-pptx then fails to open the file.
    """
    name = "[Content_Types].xml"
    with zipfile.ZipFile(path) as zin:
        if not any(n.lower().endswith(".fntdata") for n in zin.namelist()):
            return
        types = zin.read(name).decode("utf-8")
        if 'Extension="fntdata"' in types:
            return
        types = types.replace("</Types>", '<Default Extension="fntdata" ContentType="application/x-fontdata"/></Types>', 1)
        temporary = path.with_name(path.name + ".tmp")
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, types.encode("utf-8") if item.filename == name else zin.read(item.filename))
    temporary.replace(path)
    logger.info("Registered the .fntdata content type for embedded fonts")


def _footer_text(pattern: str, metadata: Metadata) -> str:
    values = {"title": metadata.title or "", "version": metadata.version or "", "author": metadata.author or ""}
    return re.sub(r"\{(\w+)\}", lambda m: str(values.get(m[1], m[0])), pattern)


def _add_footer(slide, text: str) -> bool:
    """Put ``text`` in the slide's footer, copying the footer placeholder from its layout. False if the layout has none."""
    for shape in slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.type == PP_PLACEHOLDER.FOOTER:
            shape.text_frame.text = text
            return True
    for shape in slide.slide_layout.placeholders:
        if shape.placeholder_format.type == PP_PLACEHOLDER.FOOTER:
            element = copy.deepcopy(shape._element)
            slide.shapes._spTree.append(element)
            for placed in slide.placeholders:
                if placed._element is element:
                    placed.text_frame.text = text
            return True
    return False


def _lines(slide) -> int:
    count = 0
    for shape in slide.shapes:
        if shape.has_text_frame and not (shape.is_placeholder and shape.placeholder_format.type in (
            PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.FOOTER,
            PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE,
        )):  # fmt: skip
            count += sum(1 for p in shape.text_frame.paragraphs if p.text.strip())
    return count


def _fill_version(shapes, version: str) -> None:
    for shape in shapes:
        if shape.has_text_frame and "{{VERSION}}" in shape.text_frame.text:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.text = run.text.replace("{{VERSION}}", version)


def finish_pptx(path: Path, metadata: Metadata, options: dict[str, Any]) -> None:
    """Rewrite ``path`` in place."""
    ensure_fntdata_content_type(path)
    presentation = Presentation(str(path))
    if not len(presentation.slides):
        raise ConversionError(
            "The deck has no slides. The template probably lacks the layouts pandoc uses "
            "(Title Slide, Title and Content, Section Header, Two Content, Comparison, Content with Caption, Blank). "
            "Start from the output of `pandoc -o custom-reference.pptx --print-default-data-file reference.pptx`"
        )
    version = metadata.version or ""
    footer = options.get("footer")
    text = _footer_text(footer, metadata) if footer else None
    without_footer = 0
    for master in presentation.slide_masters:  # the token usually sits on the master or a layout
        _fill_version(master.shapes, version)
        for layout in master.slide_layouts:
            _fill_version(layout.shapes, version)
    for number, slide in enumerate(presentation.slides, start=1):
        _fill_version(slide.shapes, version)
        if text and not _add_footer(slide, text):
            without_footer += 1
        lines = _lines(slide)
        if lines > LINES_PER_SLIDE:
            logger.warning(f"Slide {number} has {lines} lines, which is likely too many to read. Split it with another heading")
    if without_footer:
        logger.warning(f"The layout of {without_footer} slide(s) has no footer placeholder, so they have no footer")
    presentation.save(str(path))
