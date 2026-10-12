"""Finish a Word file written by pandoc: merge it into the template, fill placeholders, tidy the layout."""

from __future__ import annotations

import logging
import warnings
from pathlib import Path
from typing import Any

from docx import Document

from py_doc.markdown import Metadata

from . import docx_post as post

logger = logging.getLogger(__name__)


def _composer():
    # docxcompose imports pkg_resources, which warns on every import and is gone from setuptools 81.
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        from docxcompose.composer import Composer
    return Composer


def finish_docx(body_path: Path, template: Path | None, metadata: Metadata, options: dict[str, Any]) -> None:
    """Rewrite ``body_path`` in place. Without a template pandoc's own output is only tidied."""
    extra = metadata.extra
    text = lambda value: value or ""  # noqa: E731
    if template is not None:
        base = Document(str(template))
        composer = _composer()(base)
        post.replace_placeholders(
            base,
            {
                "Replace-with-Title": text(metadata.title),
                "Replace-with-Author": text(metadata.author),
                "Replace-with-Version": text(metadata.version),
                "Replace-with-Document-Number": text(metadata.document_id),
                "Replace-with-Date": text(metadata.date),
                "Replace-with-Description": text(extra.get("description")),
                "Replace-with-Classification": text(extra.get("classification")),
            },
        )
        index = post.find_and_remove_marker(base)
        body = Document(str(body_path))
        if index is not None:
            composer.insert(index, body)
        else:
            composer.append(body)
        post.mark_for_toc_update(base)
        composer.save(str(body_path))

    document = Document(str(body_path))
    post.clean_trailing_pagebreak(document)
    post.apply_orientation_markers(document)
    post.apply_layout_fixes(document)
    post.rescale_landscape_images(document)
    post.normalize_body_bold(document)
    post.apply_explicit_styles(document)
    post.format_all_tables(document)
    post.update_core_properties(
        document,
        title=text(metadata.title),
        author=text(metadata.author),
        description=text(extra.get("description")),
        version=text(metadata.version),
        document_number=text(metadata.document_id),
        subject=text(extra.get("subject")),
        keywords=text(metadata.keywords),
        category=text(extra.get("category")),
    )
    post.mark_for_toc_update(document)
    document.save(str(body_path))
