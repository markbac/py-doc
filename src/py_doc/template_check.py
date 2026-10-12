"""Check Word and PowerPoint templates for what the pandoc pipeline needs.

``check_template`` returns findings at three levels. FAIL means the template will give a wrong result
(a style of the wrong type, bold leaking into the body text, a deck that comes out empty). WARN means a
degraded result (a style pandoc will have to take from its own defaults). INFO is shown on request.
"""

from __future__ import annotations

import logging
import warnings
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from .errors import ConversionError
from .files import template_source
from .pandoc.engine import _OFFICE
from .pandoc.reference import CALLOUTS, PREFIX

logger = logging.getLogger(__name__)

WORD_SUFFIXES = (".docx", ".dotx")
POWERPOINT_SUFFIXES = (".pptx", ".potx")
LEVELS = ("FAIL", "WARN", "INFO")

# Styles pandoc writes. A template that lacks one gets pandoc's own default for it.
PANDOC_PARAGRAPH_STYLES = (
    "Title", "Subtitle", "Author", "Date", "Abstract", "Body Text", "First Paragraph", "Compact", "Block Text",
    "Source Code", "Image Caption", "Table Caption", "Footnote Text", "Heading 1", "Heading 2", "Heading 3",
)  # fmt: skip
PANDOC_CHARACTER_STYLES = ("Verbatim Char", "Hyperlink", "Footnote Reference")
# Layouts pandoc looks up by name; the first two are needed for any deck.
REQUIRED_LAYOUTS = ("Title Slide", "Title and Content")
OTHER_LAYOUTS = ("Section Header", "Two Content", "Comparison", "Content with Caption", "Blank")
MAX_FONTS = 2


@dataclass(frozen=True)
class Finding:
    level: str  # FAIL, WARN or INFO
    check: str
    message: str


@dataclass
class TemplateReport:
    path: Path
    findings: list[Finding] = field(default_factory=list)

    def add(self, level: str, check: str, message: str) -> None:
        self.findings.append(Finding(level, check, message))

    @property
    def failures(self) -> list[Finding]:
        return [f for f in self.findings if f.level == "FAIL"]

    @property
    def warnings(self) -> list[Finding]:
        return [f for f in self.findings if f.level == "WARN"]


def kind_of(path: Path) -> str:
    """``word`` or ``powerpoint`` for the file's extension. Raises ConversionError for anything else."""
    suffix = path.suffix.lower()
    if suffix in WORD_SUFFIXES:
        return "word"
    if suffix in POWERPOINT_SUFFIXES:
        return "powerpoint"
    raise ConversionError(f"'{path.name}' is not a template: expected {', '.join(WORD_SUFFIXES + POWERPOINT_SUFFIXES)}")


def _opened(path: Path, kind: str):
    """The document or deck for ``path``. A .dotx or .potx is opened as the document type."""
    if not path.is_file():
        raise FileNotFoundError(f"Template not found: {path}")
    if not zipfile.is_zipfile(path):
        raise ConversionError(f"'{path}' is not a valid {path.suffix} file")
    if kind == "word":
        source = template_source(
            path, ".dotx", f"{_OFFICE}.wordprocessingml.template.main+xml".encode(),
            f"{_OFFICE}.wordprocessingml.document.main+xml".encode(),
        )  # fmt: skip
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            return Document(source)
    source = template_source(
        path, ".potx", f"{_OFFICE}.presentationml.template.main+xml".encode(),
        f"{_OFFICE}.presentationml.presentation.main+xml".encode(),
    )  # fmt: skip
    return Presentation(source)


# --- Word ------------------------------------------------------------------------------------------


_TYPES = {
    WD_STYLE_TYPE.PARAGRAPH: "paragraph", WD_STYLE_TYPE.CHARACTER: "character",
    WD_STYLE_TYPE.TABLE: "table", WD_STYLE_TYPE.LIST: "list",
}  # fmt: skip


def styles_of(document) -> dict[str, tuple[str, object]]:
    """Style name -> (type name, style) for a Word document."""
    return {style.name: (_TYPES.get(style.type, "other"), style) for style in document.styles}


def _bold(style) -> bool:
    try:
        return bool(style.font.bold)
    except Exception:  # a style type without fonts
        return False


def _language(style) -> str | None:
    rpr = getattr(style.element, "rPr", None)
    lang = rpr.find(qn("w:lang")) if rpr is not None else None
    return lang.get(qn("w:val")) if lang is not None else None


def _check_word(document, report: TemplateReport, prefix: str = "Replace-with-", marker_text: str = "[[INSERTCONTENT]]") -> None:
    styles = styles_of(document)
    counts = {kind: sum(1 for k, _ in styles.values() if k == kind) for kind in ("paragraph", "character", "table")}
    report.add(
        "INFO", "Inventory",
        f"{len(styles)} styles: {counts['paragraph']} paragraph, {counts['character']} character, {counts['table']} table",
    )  # fmt: skip

    missing = [PREFIX + kind for kind in CALLOUTS if PREFIX + kind not in styles]
    for name in (PREFIX + kind for kind in CALLOUTS):
        if name in styles and styles[name][0] != "paragraph":
            report.add("FAIL", "Callout styles", f"{name} must be a paragraph style, not a {styles[name][0]} style")
    if missing:
        report.add(
            "WARN", "Callout styles",
            f"Missing {', '.join(missing)}: callouts of those kinds are plain paragraphs",
        )  # fmt: skip

    for name in PANDOC_PARAGRAPH_STYLES + PANDOC_CHARACTER_STYLES:
        wanted = "character" if name in PANDOC_CHARACTER_STYLES else "paragraph"
        if name in styles and styles[name][0] != wanted:
            report.add("FAIL", "Pandoc styles", f"{name} must be a {wanted} style, not a {styles[name][0]} style")
    absent = [n for n in PANDOC_PARAGRAPH_STYLES + PANDOC_CHARACTER_STYLES if n not in styles]
    if absent:
        report.add("WARN", "Pandoc styles", f"Not in the template, so pandoc's defaults are used: {', '.join(absent)}")

    if "Caption" not in styles:
        report.add("WARN", "Caption style", "No 'Caption' style: [[CAPTION:Figure]] lines fall back to the default look")
    elif styles["Caption"][0] != "paragraph":
        report.add("FAIL", "Caption style", f"Caption must be a paragraph style, not a {styles['Caption'][0]} style")
    elif _bold(styles["Caption"][1]):
        report.add("WARN", "Caption style", "Caption is bold")

    if "Table Grid" not in styles:
        report.add("WARN", "Table style", "No 'Table Grid' table style: tables use the Word default")
    elif styles["Table Grid"][0] != "table":
        report.add("FAIL", "Table style", f"'Table Grid' exists but is a {styles['Table Grid'][0]} style, not a table style")

    for name in ("Source Code", "Verbatim Char"):
        if name in styles and not styles[name][1].font.no_proof:
            report.add("INFO", "Code spelling", f"{name} has spell-check on")

    soft = [n for n, (kind, s) in styles.items() if kind == "paragraph" and n.lower().startswith("heading")]
    loose = [n for n in soft if not (styles[n][1].paragraph_format.keep_with_next and styles[n][1].paragraph_format.keep_together)]
    if not soft:
        report.add("WARN", "Headings", "No heading styles")
    elif loose:
        report.add("INFO", "Headings", f"'Keep with next' or 'Keep lines together' is off for: {', '.join(loose)}")

    for name in ("Normal", "Body Text"):
        if name in styles and _bold(styles[name][1]):
            report.add("FAIL", "Bold leak", f"'{name}' is bold, so the body text will be bold")
    for name in soft:
        following = getattr(styles[name][1], "next_paragraph_style", None)
        if following is not None and following.name != name and _bold(following):
            report.add("FAIL", "Bold leak", f"The style after {name} is '{following.name}', which is bold")
    if "Normal" in styles:
        normal = styles["Normal"][1]
        base = normal.base_style
        if base is not None and base.name.lower().startswith("heading"):
            report.add("FAIL", "Normal style", f"Normal is based on '{base.name}' and inherits its look")
        language = _language(normal)
        if language == "en-US":
            report.add("WARN", "Language", "Normal is en-US. Change it to en-GB if the documents are British English")
        elif language:
            report.add("INFO", "Language", f"Normal language: {language}")

    from .pandoc.docx_post import find_placeholders

    found = sorted(find_placeholders(document, prefix))
    marker = any(marker_text in p.text for p in document.paragraphs)
    report.add("INFO", "Placeholders", f"{prefix}* placeholders: {', '.join(found) if found else 'none'}")
    report.add(
        "INFO" if marker else "WARN", "Content marker",
        f"{marker_text} found" if marker else f"No {marker_text} line, so the document goes at the end of the template",
    )  # fmt: skip


# --- PowerPoint ------------------------------------------------------------------------------------


def _check_deck(deck, report: TemplateReport) -> None:
    layouts = {layout.name: layout for layout in deck.slide_layouts}
    report.add("INFO", "Inventory", f"{len(layouts)} layouts: {', '.join(layouts)}")
    for name in REQUIRED_LAYOUTS:
        if name not in layouts:
            report.add("FAIL", "Layouts", f"No layout named '{name}': pandoc cannot place slides without it")
    absent = [n for n in OTHER_LAYOUTS if n not in layouts]
    if absent:
        report.add("WARN", "Layouts", f"No layout named {', '.join(absent)}")
    without_footer = [
        n for n, layout in layouts.items()
        if not any(p.placeholder_format.type == PP_PLACEHOLDER.FOOTER for p in layout.placeholders)
    ]  # fmt: skip
    if layouts and len(without_footer) == len(layouts):
        report.add("WARN", "Footer", "No layout has a footer placeholder, so the 'footer' option has no effect")
    elif without_footer:
        report.add("INFO", "Footer", f"No footer placeholder in: {', '.join(without_footer)}")
    if not any("Slide Number" in shape.name for master in deck.slide_masters for shape in master.shapes):
        report.add("INFO", "Slide numbers", "The slide master has no slide number placeholder")
    if len(deck.slides):
        _audit_slides(deck, report)


def _audit_slides(deck, report: TemplateReport) -> None:
    """Consistency of a finished deck: fonts, empty placeholders."""
    fonts: dict[str, list[int]] = {}
    for number, slide in enumerate(deck.slides, start=1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        fonts.setdefault(run.font.name or "theme font", []).append(number)
        for shape in slide.placeholders:
            if shape.has_text_frame and not shape.text_frame.text.strip():
                if shape.placeholder_format.type in (
                    PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT,
                ):  # fmt: skip
                    report.add("WARN", "Empty placeholder", f"Slide {number}: '{shape.name}' is empty")
    if len(fonts) > MAX_FONTS:
        listed = "; ".join(f"{font} (slide {slides[0]})" for font, slides in fonts.items())
        report.add("WARN", "Fonts", f"{len(fonts)} fonts are used, which looks pasted-in: {listed}")


# --- public ----------------------------------------------------------------------------------------


def check_template(path: Path, prefix: str = "Replace-with-", marker: str = "[[INSERTCONTENT]]") -> TemplateReport:
    """Check one template. Raises FileNotFoundError, or ConversionError for a file that is not a template.

    ``prefix`` and ``marker`` are the ``placeholder_prefix`` and ``marker`` options of the Word output, for a
    template that does not use the defaults."""
    path = Path(path)
    kind = kind_of(path)
    report = TemplateReport(path)
    opened = _opened(path, kind)
    if kind == "word":
        _check_word(opened, report, prefix, marker)
    else:
        _check_deck(opened, report)
    return report


def inventory(path: Path) -> list[tuple[str, ...]]:
    """Rows describing what the template holds. Word: (type, style, bold, next). PowerPoint: (layout, placeholders)."""
    path = Path(path)
    kind = kind_of(path)
    opened = _opened(path, kind)
    if kind == "word":
        rows = []
        for name, (type_name, style) in styles_of(opened).items():
            following = getattr(style, "next_paragraph_style", None) if type_name == "paragraph" else None
            rows.append((type_name, name, "bold" if _bold(style) else "", following.name if following else ""))
        return sorted(rows)
    return [
        (layout.name, ", ".join(sorted({str(p.placeholder_format.type).split(".")[-1].split(" ")[0] for p in layout.placeholders})))
        for layout in opened.slide_layouts
    ]


def _definition(style) -> tuple[str, ...]:
    """What a style is, in a form two templates can be compared by."""
    font, fmt = getattr(style, "font", None), getattr(style, "paragraph_format", None)
    if font is None:  # a list style
        return ()
    size = font.size.pt if font.size else None
    based = style.base_style.name if getattr(style, "base_style", None) is not None else ""
    following = getattr(style, "next_paragraph_style", None)
    parts = {
        "based on": based,
        "next": following.name if following is not None else "",
        "font": font.name,
        "size": size,
        "bold": font.bold,
        "italic": font.italic,
        "colour": str(font.color.rgb) if font.color is not None and font.color.type is not None and font.color.rgb else None,
        "keep with next": fmt.keep_with_next if fmt is not None else None,
        "keep together": fmt.keep_together if fmt is not None else None,
        "space before": fmt.space_before.pt if fmt is not None and fmt.space_before else None,
        "space after": fmt.space_after.pt if fmt is not None and fmt.space_after else None,
        "alignment": str(fmt.alignment) if fmt is not None and fmt.alignment is not None else None,
    }
    return tuple(f"{key}: {value}" for key, value in parts.items() if value not in (None, ""))


def compare(first: Path, second: Path) -> list[tuple[str, str, str]]:
    """Rows (change, name, detail) between two Word templates, or two decks (by layout name)."""
    kinds = kind_of(Path(first)), kind_of(Path(second))
    if kinds[0] != kinds[1]:
        raise ConversionError(f"Cannot compare a {kinds[0]} template with a {kinds[1]} one")
    a, b = _opened(Path(first), kinds[0]), _opened(Path(second), kinds[1])
    rows: list[tuple[str, str, str]] = []
    if kinds[0] == "word":
        left, right = styles_of(a), styles_of(b)
        for name in sorted(left.keys() | right.keys()):
            if name not in right:
                rows.append(("only in first", name, left[name][0]))
            elif name not in left:
                rows.append(("only in second", name, right[name][0]))
            elif left[name][0] != right[name][0]:
                rows.append(("type differs", name, f"{left[name][0]} / {right[name][0]}"))
            else:
                one, two = _definition(left[name][1]), _definition(right[name][1])
                if one != two:
                    changed = sorted(set(one) ^ set(two))
                    rows.append(("differs", name, "; ".join(changed)))
        return rows
    left = {layout.name: sorted(str(p.placeholder_format.type) for p in layout.placeholders) for layout in a.slide_layouts}
    right = {layout.name: sorted(str(p.placeholder_format.type) for p in layout.placeholders) for layout in b.slide_layouts}
    for name in sorted(left.keys() | right.keys()):
        if name not in right:
            rows.append(("only in first", name, ""))
        elif name not in left:
            rows.append(("only in second", name, ""))
        elif left[name] != right[name]:
            rows.append(("differs", name, f"{len(left[name])} / {len(right[name])} placeholders"))
    return rows


def usage(path: Path) -> list[tuple[str, int]]:
    """How often each paragraph and table style is used in a Word file, most used first."""
    path = Path(path)
    if kind_of(path) != "word":
        raise ConversionError(f"Style usage is only counted for Word files: {path}")
    document = _opened(path, "word")
    counts: dict[str, int] = {}

    def paragraphs(container):
        for paragraph in container.paragraphs:
            counts[paragraph.style.name] = counts.get(paragraph.style.name, 0) + 1
        for table in getattr(container, "tables", []):
            counts[table.style.name if table.style is not None else "(none)"] = (
                counts.get(table.style.name if table.style is not None else "(none)", 0) + 1
            )
            for row in table.rows:
                for cell in row.cells:
                    paragraphs(cell)

    paragraphs(document)
    return sorted(counts.items(), key=lambda item: (-item[1], item[0]))


_logged: set[tuple[Path, float]] = set()


def log_problems(path: Path, skip: tuple[str, ...] = ()) -> TemplateReport | None:
    """Check ``path`` for a build and log what is wrong, once per file. Never raises: a build goes on."""
    try:
        stamp = (path.resolve(), path.stat().st_mtime)
        report = check_template(path)
    except (OSError, ConversionError, ValueError, KeyError) as exc:
        logger.debug(f"Template check skipped for {path}: {exc}")
        return None
    if stamp not in _logged:
        _logged.add(stamp)
        for finding in report.failures:
            if finding.check not in skip:
                logger.warning(f"Template {path.name}: {finding.check}: {finding.message}")
        softer = [f for f in report.warnings if f.check not in skip]
        if softer:
            logger.info(f"Template {path.name} has {len(softer)} warning(s). `py-doc template check {path}` lists them")
    return report
