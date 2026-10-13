"""Turn a PowerPoint deck into Markdown that ``py-doc slides`` builds again: ``py-doc import deck.pptx``.

The Markdown follows the conventions the slides converter reads (``--split-level 2``):

- the title slide becomes the front matter (title, subtitle), with the file's own properties;
- a Section Header slide becomes a ``#`` heading, any other slide a ``##`` heading, and a slide without a
  title a page break;
- text placeholders become bullet lists, nested by their level, and text boxes become paragraphs;
  monospace paragraphs become a code block, and a bold NOTE, TIP, ... paragraph starts a callout;
- the two sides of a Two Content or Comparison slide become a ``::::: columns`` block;
- tables become pipe tables, pictures go to ``<name>/media`` as PNG files, and speaker notes become a
  ``::: notes`` block.

What cannot be carried over (charts, SmartArt, video, equations, other embedded objects) is counted in one
warning each, never dropped quietly.
"""

from __future__ import annotations

import hashlib
import logging
import re
import zipfile
from collections.abc import Callable
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.exc import PackageNotFoundError

from .docx_import import (
    BITMAP_TYPES,
    MONOSPACE,
    VECTOR_TYPES,
    LIBRARY_DEFAULTS,
    ImportResult,
    _custom_properties,
    _escape,
    _line_start,
    inkscape_png,
    pillow_png,
    snake_case_name,
)
from .errors import ConversionError, EnvironmentProblem
from .runtime import Runtime, connect

logger = logging.getLogger(__name__)

TITLES = (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
SUBTITLE = PP_PLACEHOLDER.SUBTITLE
# Placeholders that repeat on every slide and are filled from the deck's properties, not its text
FURNITURE = (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.HEADER)
NOT_IMPORTED = {
    MSO_SHAPE_TYPE.CHART: "charts",
    MSO_SHAPE_TYPE.MEDIA: "videos and sounds",
    MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT: "embedded objects",
    MSO_SHAPE_TYPE.LINKED_OLE_OBJECT: "embedded objects",
    MSO_SHAPE_TYPE.DIAGRAM: "SmartArt diagrams",
}
CALLOUTS = ("NOTE", "TIP", "IMPORTANT", "WARNING", "CAUTION")
CAPTION_GAP = 1_500_000  # EMU, about 4 cm: how far below a picture or table its caption can sit
LINE_BREAK = "\x0b"  # python-pptx's text for <a:br/>
ZERO_WIDTH = "​"  # pandoc puts one between two callouts
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_A14 = "{http://schemas.microsoft.com/office/drawing/2010/main}"
_MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
_FILE_NAME = re.compile(r"^\S+\.(png|jpe?g|gif|svg|bmp|emf|wmf|tiff?)$", re.IGNORECASE)


class InvalidPptx(ConversionError):
    """The input is missing or is not a PowerPoint deck."""


class _Deck:
    def __init__(self, presentation, name: str, runtime: Callable[[], Runtime] | None) -> None:
        self.prs = presentation
        self.name = name
        self._runtime = runtime
        self._connected: Runtime | None = None
        self.media: dict[str, bytes] = {}
        self._by_hash: dict[str, str] = {}
        self.counts: dict[str, int] = {"pictures not imported": 0}
        self.problems: list[str] = []
        self.title: str | None = None
        self.subtitle: str | None = None
        self.chunks: list[str] = []

    # -- slides ---------------------------------------------------------------------------------

    def run(self) -> None:
        for number, slide in enumerate(self.prs.slides, start=1):
            if number == 1 and self._title_slide(slide):
                continue
            self._slide(slide)

    def _title_slide(self, slide) -> bool:
        """Take the title and subtitle of a first slide that has a centred title, as the front matter does."""
        kinds = {shape.placeholder_format.type: shape for shape in slide.placeholders}
        if PP_PLACEHOLDER.CENTER_TITLE not in kinds:
            return False
        self.title = _plain(kinds[PP_PLACEHOLDER.CENTER_TITLE].text_frame.text) or None
        if SUBTITLE in kinds:  # pandoc puts the author and the date on later lines of the subtitle
            lines = kinds[SUBTITLE].text_frame.text.replace(LINE_BREAK, "\n").strip().splitlines()
            self.subtitle = _plain(lines[0]) if lines else None
        others = [s for s in _shapes(slide) if not (s.is_placeholder and s.placeholder_format.type in (*TITLES, SUBTITLE, *FURNITURE))]
        if others:  # more than a title slide: keep its content as a slide of its own
            self._slide(slide, skip_title=True)
        return True

    def _slide(self, slide, skip_title: bool = False) -> None:
        title = None
        for shape in slide.placeholders:
            if shape.placeholder_format.type in TITLES:
                title = _plain(shape.text_frame.text)
        layout = slide.slide_layout.name.lower()
        if skip_title:
            pass
        elif title:
            self.chunks.append(f"{'#' if 'section' in layout else '##'} {_escape(title)}")
        else:
            self.chunks.append("::: {.pagebreak}\n:::")
        bodies = [s for s in _reading_order(_shapes(slide)) if not _is_title_or_furniture(s)]
        if "two" in layout or "comparison" in layout:
            middle = self.prs.slide_width // 2
            left = [s for s in bodies if (s.left or 0) + (s.width or 0) // 2 < middle]
            right = [s for s in bodies if s not in left]
            if left and right:
                self._columns([left, right])
                bodies = []
        self._shapes_to(bodies, self.chunks)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            text = "\n\n".join(_runs(p) for p in notes.paragraphs if p.text.strip()) if notes is not None else ""
            if text:
                self.chunks.append(f"::: notes\n{text}\n:::")

    def _columns(self, columns: list[list]) -> None:
        parts = ["::::: columns"]
        for shapes in columns:
            inner: list[str] = []
            self._shapes_to(shapes, inner)
            parts.append("::: column\n" + "\n\n".join(inner) + "\n:::")
        parts.append(":::::")
        self.chunks.append("\n\n".join(parts))

    # -- shapes ---------------------------------------------------------------------------------

    def _shapes_to(self, shapes, out: list[str]) -> None:
        shapes = list(shapes)
        index = 0
        while index < len(shapes):
            shape, after = shapes[index], shapes[index + 1] if index + 1 < len(shapes) else None
            label = "Table" if getattr(shape, "has_table", False) else "Figure" if _is_picture(shape) else None
            if label and after is not None and _is_caption(shape, after):
                inner: list[str] = []
                self._shape(shape, inner)
                if inner:
                    out.append(f"[[CAPTION:{label}]] {_runs(after.text_frame.paragraphs[0])}\n\n{inner[0]}")
                index += 2
                continue
            self._shape(shape, out)
            index += 1

    def _shape(self, shape, out: list[str]) -> None:
        kind = shape.shape_type
        if kind == MSO_SHAPE_TYPE.GROUP:
            self._shapes_to(_reading_order(_shapes(shape)), out)
        elif kind in NOT_IMPORTED or getattr(shape, "has_chart", False):
            label = NOT_IMPORTED.get(kind, "charts")
            self.counts[label] = self.counts.get(label, 0) + 1
        elif getattr(shape, "has_table", False):
            out.append(self._table(shape.table))
        elif kind == MSO_SHAPE_TYPE.PICTURE or (shape.is_placeholder and hasattr(shape, "image")):
            picture = self._picture(shape)
            if picture:
                out.append(picture)
        elif shape.has_text_frame and shape.text_frame.text.replace(ZERO_WIDTH, "").strip():
            text = self._text(shape)
            if text:
                out.append(text)

    def _text(self, shape) -> str:
        """Bullets, numbered items, paragraphs, code and callouts, in the order they come."""
        bulleted = shape.is_placeholder and _is_content(shape)
        blocks: list[str] = []
        items: list[str] = []
        code: list[str] = []
        callout: list[str] = []

        def flush() -> None:
            if items:
                blocks.append("\n".join(items))
                items.clear()
            if code:
                text = "\n".join(code)
                fence = "````" if "```" in text else "```"
                blocks.append(f"{fence}\n{text}\n{fence}")
                code.clear()
            if callout:
                blocks.append("\n>\n".join(callout))
                callout.clear()

        for paragraph in shape.text_frame.paragraphs:
            if _is_code(paragraph):
                if items or callout:
                    flush()
                code.extend(paragraph.text.replace(LINE_BREAK, "\n").split("\n"))
                continue
            equations = _equations(paragraph)
            if equations:
                self.counts["equations"] = self.counts.get("equations", 0) + equations
            text = " ".join(filter(None, [_runs(paragraph), *["(equation not imported)"] * equations]))
            if not text.replace(ZERO_WIDTH, "").strip():
                continue
            label = _callout_label(paragraph)
            marker = _marker(paragraph, bulleted)
            if label:
                flush()
                callout.append(f"> [!{label}]")
            elif callout and not marker:
                if len(callout) == 1:
                    callout[0] += f"\n> {_line_start(text)}"
                else:
                    callout.append(f"> {_line_start(text)}")
            elif marker:
                if code or callout:
                    flush()
                items.append(f"{'  ' * paragraph.level}{marker} {text}")
            else:
                flush()
                blocks.append(_line_start(text))
        flush()
        return "\n\n".join(blocks)

    def _table(self, table) -> str:
        rows = [[_cell(cell) for cell in row.cells] for row in table.rows]
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
        lines += ["| " + " | ".join(row) + " |" for row in rows[1:]]
        return "\n".join(lines)

    def _picture(self, shape) -> str | None:
        try:
            image = shape.image
        except (AttributeError, ValueError, KeyError):
            return None
        kind = (image.content_type or "").lower()
        digest = hashlib.sha256(image.blob).hexdigest()
        file = self._by_hash.get(digest)
        if file is None:
            try:
                if kind == "image/png":
                    data = image.blob
                elif kind in BITMAP_TYPES:
                    data = pillow_png(image.blob)
                elif kind in VECTOR_TYPES:
                    if self._connected is None:
                        self._connected = self._runtime() if self._runtime is not None else connect()
                    data = inkscape_png(self._connected, image.blob, kind)
                else:
                    raise ConversionError(f"the type {kind or 'unknown'} is not supported")
            except EnvironmentProblem:
                raise
            except (ConversionError, OSError, ValueError) as exc:
                first = str(exc).strip().splitlines()[0] if str(exc).strip() else "it could not be converted"
                self.counts["pictures not imported"] += 1
                self.problems.append(f"Picture not imported, {shape.name} ({kind or 'unknown type'}): {first}")
                logger.warning("Picture not imported, %s (%s): %s", shape.name, kind or "unknown type", first)
                return None
            file = f"image_{len(self.media) + 1:03d}.png"
            self.media[file] = data
            self._by_hash[digest] = file
        described = (shape._element.xpath("./*[1]/p:cNvPr/@descr") or [""])[0]
        if _FILE_NAME.search(described):  # pandoc stores the picture's path there, which says nothing
            described = ""
        alt = re.sub(r"([\[\]])", r"\\\1", described.replace("\n", " "))
        return f"![{alt}]({self.name}/media/{file})"


# --- shapes and text --------------------------------------------------------------------------------


def _shapes(container) -> list:
    """The shapes of a slide or a group, and those pandoc wraps in ``mc:AlternateContent`` (a shape with maths)."""
    from pptx.shapes.shapetree import SlideShapeFactory

    found = list(container.shapes)
    for choice in container.shapes._spTree.iterchildren(f"{_MC}AlternateContent"):
        chosen = choice.find(f"{_MC}Choice")
        if chosen is not None:
            found += [SlideShapeFactory(element, container.shapes) for element in chosen]
    return found


def _is_picture(shape) -> bool:
    return shape.shape_type == MSO_SHAPE_TYPE.PICTURE or (shape.is_placeholder and hasattr(shape, "image"))


def _is_caption(shape, below) -> bool:
    """``below`` is one short line of text right under ``shape``, as pandoc writes a caption."""
    if not below.has_text_frame or below.has_text_frame and len([p for p in below.text_frame.paragraphs if p.text.strip()]) != 1:
        return False
    if _marker(below.text_frame.paragraphs[0], False) or _is_code(below.text_frame.paragraphs[0]):
        return False
    bottom = (shape.top or 0) + (shape.height or 0)
    overlap = min(shape.left + shape.width, below.left + below.width) - max(shape.left, below.left)
    gap = (below.top or 0) - bottom  # pandoc puts the caption at the foot of the content area
    return -CAPTION_GAP < gap < CAPTION_GAP and overlap > 0 and len(below.text_frame.text) < 200


def _reading_order(shapes) -> list:
    """Shapes from top to bottom, then left to right, as a reader takes them in."""
    return sorted(shapes, key=lambda s: ((s.top or 0) // 100_000, s.left or 0))


def _is_title_or_furniture(shape) -> bool:
    return shape.is_placeholder and shape.placeholder_format.type in (*TITLES, *FURNITURE)


def _is_content(shape) -> bool:
    """A body or content placeholder, whose paragraphs are bullets unless they say otherwise."""
    return shape.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT)


def _is_code(paragraph) -> bool:
    """Every run of the paragraph is in a monospace font: a line of a code block."""
    runs = [run for run in paragraph.runs if run.text]
    return bool(runs) and all((run.font.name or "").lower() in MONOSPACE for run in runs)


def _callout_label(paragraph) -> str | None:
    """The kind of callout a paragraph that is only a bold NOTE, TIP, ... starts, as pandoc writes one."""
    text = paragraph.text.strip().rstrip(":").upper()
    runs = [run for run in paragraph.runs if run.text.strip()]
    if text in CALLOUTS and runs and all(run.font.bold for run in runs):
        return text
    return None


def _equations(paragraph) -> int:
    return sum(1 for element in paragraph._p.iter(f"{_A14}m"))


def _marker(paragraph, bulleted: bool) -> str | None:
    """``-`` or ``1.`` for a list item, None for a plain paragraph."""
    props = paragraph._p.pPr
    if props is not None:
        if props.find(f"{_A}buNone") is not None:
            return None
        if props.find(f"{_A}buAutoNum") is not None:
            return "1."
        if props.find(f"{_A}buChar") is not None:
            return "-"
    return "-" if bulleted else None


def _runs(paragraph) -> str:
    """The paragraph's text with bold, italic, code and links as Markdown."""
    pieces: list[str] = []
    for run in paragraph.runs:
        text = run.text.replace(LINE_BREAK, " ").replace(ZERO_WIDTH, "")
        if not text:
            continue
        font = run.font
        if (font.name or "").lower() in MONOSPACE:
            piece = f"`{text}`" if "`" not in text else f"`` {text} ``"
        else:
            core = _escape(text.strip())
            if core and font.bold and font.italic:
                core = f"***{core}***"
            elif core and font.bold:
                core = f"**{core}**"
            elif core and font.italic:
                core = f"*{core}*"
            if core and font._rPr is not None and font._rPr.get("strike") in ("sngStrike", "dblStrike"):
                core = f"~~{core}~~"
            lead, trail = text[: len(text) - len(text.lstrip())], text[len(text.rstrip()) :]
            piece = f"{lead}{core}{trail}"
        address = run.hyperlink.address if run.hyperlink is not None else None
        pieces.append(f"[{piece}]({address})" if address else piece)
    return "".join(pieces).strip()


def _plain(text: str) -> str:
    return " ".join(text.replace(LINE_BREAK, " ").split())


def _cell(cell) -> str:
    text = " ".join(_runs(p) for p in cell.text_frame.paragraphs if p.text.strip())
    return text.replace("|", "\\|")


def _front_matter(deck: _Deck, path: Path) -> list[str]:
    core = deck.prs.core_properties
    data: dict[str, str] = {}

    def put(key: str, value) -> None:
        text = str(value).strip() if value is not None else ""
        if text:
            data[key] = text

    put("title", deck.title or core.title)
    put("subtitle", deck.subtitle)
    put("author", core.author)
    put("subject", core.subject)
    put("keywords", core.keywords)
    put("category", core.category)
    put("description", None if LIBRARY_DEFAULTS.match(core.comments or "") else core.comments)
    put("version", core.version)
    put("status", core.content_status)
    for key, value in _custom_properties(path).items():
        data[key] = value
    if not data:
        return []
    dumped = yaml.safe_dump(data, allow_unicode=True, sort_keys=False, default_flow_style=False, width=1000)
    return ["---", dumped.rstrip(), "---", ""]


def import_pptx(
    source: Path, out_dir: Path | None = None, *, force: bool = False, runtime: Callable[[], Runtime] | None = None
) -> ImportResult:
    """Write ``<out_dir>/<name>.md`` and ``<out_dir>/<name>/media/*.png`` for the deck ``source``.

    Raises :class:`InvalidPptx` if the file is missing or is not a deck, and :class:`ConversionError` if the
    Markdown file exists and ``force`` is not set.
    """
    source = Path(source)
    if not source.is_file():
        raise InvalidPptx(f"Input file not found: {source}")
    try:
        presentation = Presentation(str(source))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise InvalidPptx(f"'{source}' is not a PowerPoint deck (.pptx)") from exc

    name = snake_case_name(source)
    out_dir = Path(out_dir) if out_dir else source.resolve().parent
    markdown = out_dir / f"{name}.md"
    media = out_dir / name / "media"
    if markdown.exists() and not force:
        raise ConversionError(f"Output file already exists: {markdown} (use --force to replace it)")

    deck = _Deck(presentation, name, runtime)
    deck.run()
    front = _front_matter(deck, source)
    text = ("\n".join(front) + "\n" if front else "") + "\n\n".join(deck.chunks)
    for what, count in deck.counts.items():
        if count and what != "pictures not imported":
            message = f"{count} {what} not imported"
            deck.problems.append(message)
            logger.warning("%s: %s", source.name, message)

    try:
        out_dir.mkdir(parents=True, exist_ok=True)
        if deck.media:
            media.mkdir(parents=True, exist_ok=True)
            for file, data in deck.media.items():
                (media / file).write_bytes(data)
        markdown.write_text(text.rstrip() + "\n", encoding="utf-8")
    except OSError as exc:
        raise ConversionError(f"Cannot write '{markdown}': {exc.strerror or exc}") from exc
    logger.info("Imported %s to %s (%d picture(s))", source.name, markdown, len(deck.media))
    return ImportResult(markdown, media, len(deck.media), deck.problems)
