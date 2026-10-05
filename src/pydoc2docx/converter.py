import re
from pathlib import Path
from typing import Optional

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .pylogkit import setup_logging

logger = setup_logging(name="Doc2Docx", to_console=True, to_file=False)

# CommonMark ATX heading: one to six '#' followed by whitespace or end of line.
HEADING_RE = re.compile(r"^(#{1,6})(?:\s+(.*))?$")

# CommonMark ordered list item: one to nine digits, '.', then whitespace or end of line.
ORDERED_ITEM_RE = re.compile(r"^\d{1,9}\.(?:\s+(.*))?$")

# Characters that cannot appear in an XML 1.0 document (python-docx rejects them).
XML_INVALID_RE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


class ConversionError(Exception):
    """The input cannot be converted (for example, it is not valid UTF-8)."""


class DocxConverter:
    """
    Converts Markdown technical documentation into Word (.docx) documents based on createdocs specification.
    """

    def __init__(
        self,
        template_path: Optional[Path] = None,
        allow_missing_template: bool = False,
        keep_template_body: bool = False,
    ):
        """
        Args:
            template_path: Optional reference Word (.docx) template.
            allow_missing_template: If False (the default), a template path that is not
                an existing file raises FileNotFoundError. If True, a warning is logged
                and a default blank document is used instead.
            keep_template_body: If False (the default), the template's body content
                (cover page, placeholder text) is removed and only its styles, page
                setup of the final section, headers, footers and properties are used.
                If True, the converted content is appended after the template body.
        """
        self.template_path = Path(template_path) if template_path else None
        self.allow_missing_template = allow_missing_template
        self.keep_template_body = keep_template_body

    @staticmethod
    def _clear_body(doc):
        """Remove all body content except the final section properties."""
        body = doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)

    @staticmethod
    def _read_markdown(md_path: Path) -> list:
        """Return the source as a list of lines, safe to hand to python-docx."""
        try:
            # utf-8-sig drops a byte order mark; universal newlines turn CRLF and CR into LF.
            content = md_path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ConversionError(f"Cannot read '{md_path}': it is not valid UTF-8 ({exc.reason})") from exc

        content, removed = XML_INVALID_RE.subn("", content)
        if removed:
            logger.warning(
                f"Removed {removed} control character(s) that Word documents cannot hold from '{md_path.name}'"
            )

        # Split on LF only: str.splitlines() also splits on characters such as U+2028 and U+0085,
        # which are not line breaks in Markdown.
        lines = content.split("\n")
        if lines and lines[-1] == "":
            lines.pop()
        return lines

    @staticmethod
    def _add_code_block(doc, code_lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.4)
        run = p.add_run("\n".join(code_lines))
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 41, 59)

    def convert_file(self, md_path: Path, output_path: Path) -> Path:
        md_path = Path(md_path).resolve()
        output_path = Path(output_path).resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)

        if not md_path.exists():
            raise FileNotFoundError(f"Markdown file not found: {md_path}")

        logger.info(f"Converting Markdown '{md_path.name}' -> Word '{output_path.name}'...")

        if self.template_path and self.template_path.is_file():
            doc = Document(str(self.template_path))
            logger.info(f"Loaded reference Word template: {self.template_path.name}")
            if not self.keep_template_body:
                self._clear_body(doc)
        elif self.template_path:
            if not self.allow_missing_template:
                raise FileNotFoundError(f"Word template not found: {self.template_path}")
            logger.warning(f"Word template not found, using a default blank document: {self.template_path}")
            doc = Document()
        else:
            doc = Document()

        lines = self._read_markdown(md_path)

        in_code_block = False
        code_lines = []

        for line in lines:
            stripped = line.strip()

            # Handle Code Blocks
            if stripped.startswith("```"):
                if in_code_block:
                    # End of code block
                    self._add_code_block(doc, code_lines)
                    code_lines = []
                    in_code_block = False
                else:
                    in_code_block = True
                    code_lines = []
                continue

            if in_code_block:
                code_lines.append(line)
                continue

            # Headings
            if stripped.startswith("#"):
                match = HEADING_RE.match(stripped)
                if match:
                    text = (match.group(2) or "").strip()
                    if text:
                        doc.add_heading(text, level=min(len(match.group(1)), 4))
                    continue

            # Bullet points
            if stripped.startswith("- ") or stripped.startswith("* "):
                doc.add_paragraph(stripped[2:], style='List Bullet')
                continue

            # Numbered list
            num_match = ORDERED_ITEM_RE.match(stripped)
            if num_match:
                doc.add_paragraph(num_match.group(1) or "", style='List Number')
                continue

            # Paragraph
            if stripped:
                doc.add_paragraph(stripped)

        if in_code_block:
            logger.warning(
                f"Unclosed code fence in '{md_path.name}': keeping the remaining lines as a code block"
            )
            self._add_code_block(doc, code_lines)

        doc.save(str(output_path))
        logger.info(f"Successfully generated Word document: {output_path}")
        return output_path
